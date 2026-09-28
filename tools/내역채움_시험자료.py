# -*- coding: utf-8 -*-
"""
📥 내역채움 시험용 «가상» 공내역서 (2026-09-28) — K-건설맵이 지어낸 것(실제 공사·회사 아님)
  python3 tools/내역채움_시험자료.py  →  tools/시험자료/내역채움_공내역.xlsx
  · 시트 «표지»(내역 아님) · «내역서»(품명·규격·단위·수량·재료비 단가/금액·노무비 단가/금액·합계) · «일위대가»
  · 수량이 빈 줄(넣을 자리) · 이미 수량이 있는 줄 · 수량이 수식인 줄 · 금액 = 수량×단가 수식
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
wb = openpyxl.Workbook()
ws0 = wb.active; ws0.title = '표지'
ws0['A1'] = '가상 공내역서 (시험용 — 실제 공사 아님)'
ws = wb.create_sheet('내역서')
thin = Side(style='thin'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ws['A1'] = '1. 철근콘크리트공사 (가상)'; ws['A1'].font = Font(bold=True, size=14)
hd = ['품      명', '규      격', '단위', '수량', '재  료  비', None, '노  무  비', None, '합      계', None]
for j, h in enumerate(hd): ws.cell(3, j + 1, h)
for j, h in enumerate([None, None, None, None, '단  가', '금  액', '단  가', '금  액', '단  가', '금  액']): ws.cell(4, j + 1, h)
ws.merge_cells('E3:F3'); ws.merge_cells('G3:H3'); ws.merge_cells('I3:J3')
rows = [
  ('레미콘', '25-24-150', 'M3', None, 72000, 0),
  ('레미콘', '25-30-15', 'M3', None, 76000, 0),          # 도면에 30MPa 없음 → 못 찾음
  ('레미콘 펌프차 타설', '슬럼프 15cm', 'M3', None, 0, 9000),
  ('철근', 'SD400 HD10', 'TON', None, 820000, 0),
  ('철근', 'HD13', 'TON', None, 815000, 0),
  ('철근', 'HD19', 'TON', 1.5, 810000, 0),                # 이미 수량 있음 (도면과 다름)
  ('철근', 'SHD22', 'TON', None, 830000, 0),
  ('철근', 'HD25', 'TON', None, 830000, 0),              # 도면에 D25 없음
  ('철근가공조립', '공장가공 및 현장조립', 'TON', None, 0, 450000),
  ('합판거푸집', '3회', 'M2', None, 9000, 21000),
  ('먹매김', '구조', 'M2', 99, 0, 1200),                  # 골조 아님 — 건드리지 않음
  ('방수턱 거푸집', '', 'M2', None, 5000, 15000),         # 골조 거푸집이 아님
  ('콘크리트 양생', '', 'M3', '=D5', 0, 800),             # 수량이 수식
]
r0 = 5
for k, (nm, sp, u, q, ms, ls) in enumerate(rows):
  r = r0 + k
  ws.cell(r, 1, nm); ws.cell(r, 2, sp); ws.cell(r, 3, u)
  if q is not None: ws.cell(r, 4, q)
  ws.cell(r, 5, ms); ws.cell(r, 6, f'=ROUND(D{r}*E{r},0)')
  ws.cell(r, 7, ls); ws.cell(r, 8, f'=ROUND(D{r}*G{r},0)')
  ws.cell(r, 9, f'=E{r}+G{r}'); ws.cell(r, 10, f'=F{r}+H{r}')
  for c in range(1, 11):
    ws.cell(r, c).border = bd
  ws.cell(r, 4).number_format = '#,##0.000'
  ws.cell(r, 4).fill = PatternFill('solid', fgColor='FFDDEBF7')
rE = r0 + len(rows)
ws.cell(rE, 1, '합      계'); ws.cell(rE, 10, f'=SUM(J{r0}:J{rE-1})')
ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 20
ws2 = wb.create_sheet('일위대가')
ws2['A1'] = '품명'; ws2['B1'] = '규격'; ws2['C1'] = '단위'; ws2['D1'] = '수량'
ws2['A2'] = '레미콘'; ws2['B2'] = '25-24-150'; ws2['C2'] = 'M3'; ws2['D2'] = 1
wb.save('tools/시험자료/내역채움_공내역.xlsx')
print('ok')
