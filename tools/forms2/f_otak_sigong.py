# -*- coding: utf-8 -*-
"""오탁방지막 시공계획서 — 현장에서 쓰던 계획서의 틀(표지 · 목차 · Ⅰ 시공계획 1~7 · Ⅱ 품질관리 · Ⅲ 환경·안전관리계획) 그대로 새로 만든 것.

■ 다른 현장의 공사명·사람 이름은 모두 뺐습니다. 공법 설명(재료 기준 · 구조 · 부자재 · 보관 · 시공 방법 · 품질 · 안전)은 남기고 글을 다듬었습니다.
■ 원본에서 표를 그림으로 붙여 두었던 «막체 성능 기준» · «부속재료» 는 같은 내용을 칸으로 옮겨 고쳐 쓸 수 있게 했습니다.
■ 수식: 구간 연장 → 오탁방지막 SPAN 수(막 한 장 길이, 노란 칸) · 앵커 수((연장 ÷ 앵커 간격 + 1) × 줄 수, 노란 칸) · 합계 ·
        공정표 막대(착수·완료일) · 인원·장비 누계 = 일투입 × 투입일수 · 공사명이 개요 글에 저절로 들어감.
   원본 표(250m → 13 SPAN · 52개, 1,520m → 76 SPAN · 306개)가 20m · 10m · 2줄로 꼭 맞습니다.
■ 원본의 해상작업 중지 기준 «파고 1.0m/sec» 는 단위가 틀려 «파고 1.0m · 유속 1.0m/s» 로 고치고 노란 칸(현장 기준으로 고쳐 씀)으로 두었습니다."""
import datetime

from fb import F, I
from openpyxl.formatting.rule import Rule
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as CL

import docw as d

SLUG = "otak-sigong"
TITLE = "오탁방지막 시공계획서"
WHERE = "forms"
PREV = [(3, "1. 공사개요 — 빈 서식"), (13, "공사개요 — 작성 예시 (SPAN · 앵커 수 자동)"), (14, "예정공정표 — 작성 예시")]
SHORT = "오탁방지막 시공계획서 — 개요·공정표·조직도·흐름도·공종별 시공계획·인원장비 투입·문제점과 대책·품질·안전까지. 구간 연장을 넣으면 SPAN·앵커 수, 투입 누계가 저절로 나옵니다."
NOTE_REPLACE = {
    "🏗 **현장에서 실제로 쓰던 서식입니다** — 내용만 지우고 틀은 원본 그대로 두었습니다.":
        "🏗 **현장에서 쓰던 계획서의 틀**(차례·표)을 그대로 두고 새로 만들었습니다 — 다른 현장의 이름은 모두 빼고, 공법 설명은 다듬어 남겼습니다. 구간 연장 → SPAN·앵커 수, 투입 누계가 자동입니다.",
    "칸 이름·차례·결재란을 바꾸지 않았습니다. 채워 넣기만 하시면 됩니다.":
        "노란 칸 = 기준값(현장에 맞게 고쳐 씀) · 하늘색 칸 = 자동 계산. 뒤쪽 주황 탭 시트는 가상의 현장으로 채운 «작성 예시» 입니다.",
}
MULTI = True

D = datetime.date
W = [7.6] * 12   # 리브레오피스 칸 너비(글자×7+5px)로도 A4 에 배율 1 — 쪽 나눔 어림이 엑셀·리브레 모두 맞게
PC = 12
WON = '#,##0;[Red]-#,##0;""'
BAR = PatternFill("solid", fgColor="2F5597", bgColor="2F5597")
NB = 24   # 공정표 눈금 — 반달 × 12달

SEC = [("공사개요", "1. 공사개요"), ("예정공정표", "2. 예정공정표"), ("조직도", "3. 조직도"), ("흐름도", "4. 시공 흐름도"),
       ("시공계획", "5. 공종별 시공계획"), ("투입계획", "6. 인원 · 장비 · 자재 투입계획"), ("문제점", "7. 오탁방지막 문제점 및 대책 검토안")]
EX_GU = [("1구간", 250), ("2구간", 250), ("3구간", 1_520)]
EX_GJ = [("오탁방지막 제작", D(2026, 4, 1), D(2026, 4, 20), ""), ("닻(앵커) 제작", D(2026, 4, 1), D(2026, 4, 20), ""),
         ("오탁방지막 설치", D(2026, 4, 21), D(2026, 5, 15), ""), ("닻(앵커) 설치", D(2026, 4, 21), D(2026, 5, 15), ""),
         ("유지보수(월 1회 점검)", D(2026, 5, 16), D(2027, 2, 28), "월 1회"), ("오탁방지막 철수", D(2027, 3, 1), D(2027, 3, 15), "")]
ORG = [("공사 과장", "김철수"), ("작업반 소장", "이반장"), ("작업반장", "박반장"), ("잠수사", "최잠수"), ("잠수사", "정잠수"), ("특별인부", "6명")]
MAN = [("관리", "인", 1, 40, "현장 기술지도 및 안전관리"), ("반장", "인", 1, 40, "현장관리 및 작업관리"),
       ("잠수사", "인", 2, 30, "오탁방지막 설치 및 닻 설치"), ("특별인부", "인", 6, 40, "오탁방지막 설치 및 닻 설치")]
EQ = [("바지(PONTOON)", "5m×5m", "대", 1, 25, ""), ("예선", "150HP", "대", 1, 25, ""), ("작업선", "모터보트", "대", 1, 40, ""),
      ("크레인", "25톤", "대", 1, 10, "닻 상차"), ("카고트럭", "11톤", "대", 1, 8, "자재 운반")]
PROB = [("해상 가시설 설치", "가시설(가이드 프레임) 항타 때 작업 바지선의 앵커와 오탁방지막 앵커가 겹치고, 앵커 줄이 길어짐 · 자재를 나를 해상장비 항로가 없음",
         "① 오탁방지막 설치 위치를 현장 여건에 맞게 다시 잡음 ② 공정이 바뀌면 준설 면적이 늘어나는지 검토 ③ 해상장비 항로·수심 확인"),
        ("현장타설 말뚝 · 기초", "가시설 위로 대형 크레인·천공장비를 올리고 작업 바지(SEP)를 고정할 때 앵커 줄이 길어짐 · 콘크리트 타설 장비(펌프카·레미콘 운반 바지)의 항로",
         "① 자재 운반·해상장비 항로를 검토해 수심 항로 확보 ② 설치 위치 재조정 ③ 해체·이동 때도 같은 방법"),
        ("굴착 · 가시설 보강", "굴착과 함께 버팀보를 설치하는 동안 굴착 토사를 나를 해상장비 항로가 없음",
         "① 토사 운반 항로 확보 ② 설치 위치 재조정 ③ 준설 면적 증가 검토"),
        ("교각 · 상부공(강교 거치)", "대형 부재를 운반 바지에서 인양해 거치할 때 윈치·와이어로 바지를 앞뒤로 움직이므로 앵커 줄이 더 길어짐",
         "① 거치 순서에 맞춰 오탁방지막을 옮겨 달거나 여닫는 구간을 둠 ② 항로·수심 확인 ③ 설치 위치 재조정")]
PERF = ["오탁방지막", "25형", "700 이상", "635 이상", "10 ~ 30", "635 이상", "α×10⁻²~⁻⁴\n(α: 1~9)", "1.0 이상", "250 이상", "Poly-\nester", "±0.2 이하"]
PARTS = [("부체", "Polystyrene", "φ300×835×1/2", "비중 0.016 이상으로 부력 유지"),
         ("부체 커버", "Polystyrene", "φ310×840×1/2", "부체 보호용 — 외부 충격에 견디고 내한성·내구성이 좋은 재질"),
         ("부체 밴드", "SUS", "30×480", "부체 결합용 — 부식되지 않는 스테인리스"),
         ("볼트·너트", "Brass", "5/16\"×1 1/4\"", "부체 연결용 — 부식되지 않는 황동"),
         ("보강 벨트", "Polyester", "60mm", "단위 제품(20m) 사이 원단 이음부·강도 보강부에 부착"),
         ("샤클", "SS400 Zn도금", "φ14 · φ20", "제품 사이 연결용 — 조류에 풀리지 않게 스플릿 핀을 끼우는 S.B 타입, 아연도금"),
         ("보강 금구", "SS400 Zn도금", "60×150", "보강 벨트 위 샤클 연결부 보강용 — 부식 방지 아연도금"),
         ("구멍쇠", "Brass", "φ20", "제품 이음부와 하부 체인 연결용 — 부식으로 막체가 상하지 않게"),
         ("링크", "Brass", "φ7", "체인 연결용"),
         ("체인", "Steel", "φ16", "막체가 주름지지 않게 누르는 추 — 제품 폭에 따라 고름")]


def _pg(bk, ex, lg, name, landscape=False):
    return bk.page(lg, W, ex=ex, fit_height=1, landscape=landscape, sheetname=name if not ex else "예시-" + name,
                   margins=(0.45, 0.45, 0.5, 0.5))


def cover(bk, ex):
    p = _pg(bk, ex, "표지", "표지")
    p.gap(120)
    p.row([(PC, "오 탁 방 지 막   시 공 계 획 서", {"kind": "title", "size": 24})], h=60)
    p.gap(80)
    L = {"kind": "label"}
    p.row([(2, "", {"kind": "free"}), (2, "공 사 명", L), (6, I("공사명", ex="가나항 물양장 정비공사", align="center")), (2, "", {"kind": "free"})], h=30)
    p.row([(2, "", {"kind": "free"}), (2, "제 출 일", L), (6, I("제출일", ex=D(2026, 3, 25), fmt="date", align="center")), (2, "", {"kind": "free"})], h=30)
    p.gap(200)
    p.row([(PC, I("시공사", ex="예시건설(주)", blank="(회 사 명)", align="center", size=16, bold=True), {"border": False})], h=40)
    d.done(p)
    p.after_note(["공사명·제출일·회사를 여기 한 번만 적으면 뒤 장(공사개요 등)에 저절로 들어갑니다."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 공사)."])
    return p


def toc(bk, ex):
    p = _pg(bk, ex, "목차", "목차")
    p.gap(30)
    p.row([(PC, "목          차", {"kind": "title", "size": 20})], h=50)
    p.gap(30)
    lines = [("Ⅰ.", "시 공 계 획", True)] + [("", t, False) for _, t in SEC] + [("Ⅱ.", "품 질 관 리", True), ("Ⅲ.", "환경 · 안전관리계획", True)]
    for a, t, big in lines:
        p.row([(2, "", {"kind": "free"}), (1, a, {"kind": "free", "bold": True, "size": 12}),
               (7, t if big else "    " + t, {"kind": "free", "bold": big, "size": 12 if big else 11}), (2, "", {"kind": "free"})], h=30)
    d.done(p)
    return p


def gaeyo(bk, ex):
    p = _pg(bk, ex, "공사개요", "1.공사개요")
    p.row([(PC, "Ⅰ.  시 공 계 획", {"kind": "title", "size": 16})], h=36)
    d.h1(p, "1. 공사개요")
    p.row([(PC, F("개요글", '="    본 공사는 "&IF({표지!공사명}="","○○ 공사","「"&{표지!공사명}&"」")&" 기간 중 준설과 기초공사로 생기는 오탁수와 부유물질이 '
                           '퍼져 바다(하천)가 오탁되는 것을 막기 위하여 오탁방지막을 설치하는 공사입니다."', align="left", ink=False),
            {"kind": "free", "valign": "top", "size": 10.5})], h=48)
    p.gap(6)
    L = {"kind": "label"}
    p.row([(3, "막 한 장(1 SPAN) 길이", L), (2, I("막길이", ex=20, blank=20, fmt='0"m"', rate=True, align="center")),
           (3, "앵커 간격 · 줄 수", L), (2, I("앵커간격", ex=10, blank=10, fmt='0"m 간격"', rate=True, align="center")),
           (2, I("앵커줄", ex=2, blank=2, fmt='0"줄"', rate=True, align="center"))], h=24)
    p.gap(4)
    H = {"kind": "head"}
    p.row([(2, "구  분", H), (3, "연  장", H), (2, "오탁방지막", H), (2, "닻가지형 앵커", H), (3, "비  고", H)], h=24)
    r0 = p.r
    for i in range(1, 6):
        g = EX_GU[i - 1] if i <= len(EX_GU) else (None, None)
        p.row([(2, I(f"구간#{i}", ex=g[0], align="center")), (3, I(f"연장#{i}", ex=g[1], fmt='#,##0.##"m";;""', align="center")),
               (2, F(f"SPAN#{i}", f'=IF(OR(N({{연장#{i}}})=0,N({{막길이}})=0),"",ROUNDUP({{연장#{i}}}/{{막길이}},0))', fmt='#,##0" SPAN";;""', align="center")),
               (2, F(f"앵커#{i}", f'=IF(OR(N({{연장#{i}}})=0,N({{앵커간격}})=0),"",(ROUNDUP({{연장#{i}}}/{{앵커간격}},0)+1)*N({{앵커줄}}))',
                     fmt='#,##0" EA";;""', align="center")),
               (3, I(f"구간비고#{i}", size=9))], h=22)
    p.k["연장[]"] = f"C{r0}:C{p.r - 1}"
    p.k["SPAN[]"] = f"F{r0}:F{p.r - 1}"
    p.k["앵커[]"] = f"H{r0}:H{p.r - 1}"
    S = {"kind": "sum"}
    p.row([(2, "합  계", S), (3, F("합연장", '=IF(SUM({연장[]})=0,"",SUM({연장[]}))', fmt='#,##0.##"m";;""', align="center", bold=True)),
           (2, F("합SPAN", '=IF(SUM({SPAN[]})=0,"",SUM({SPAN[]}))', fmt='#,##0" SPAN";;""', align="center", bold=True)),
           (2, F("합앵커", '=IF(SUM({앵커[]})=0,"",SUM({앵커[]}))', fmt='#,##0" EA";;""', align="center", bold=True)), (3, "", S)], h=24)
    p.row([(PC, "※ SPAN = 연장 ÷ 막 한 장 길이(올림) · 앵커 = (연장 ÷ 앵커 간격(올림) + 1) × 줄 수 — 노란 칸은 설계에 맞게 고쳐 씁니다.",
            {"kind": "note"})], h=16)
    d.para(p, "※ 오랫동안 설치해 두어 오탁방지막의 평균 수명(2~3년)을 넘기면 다시 설치합니다.", indent=1)
    d.h1(p, "2. 예정공정표 (별첨)")
    d.h1(p, "3. 조직도 (별첨)")
    d.done(p)
    return p


def _gantt(bk, ex):
    Wg = [18, 9, 9] + [2.4] * NB + [9]
    p = bk.page("예정공정표", Wg, ex=ex, landscape=True, fit_height=1, sheetname="2.예정공정표" if not ex else "예시-2.예정공정표",
                margins=(0.35, 0.35, 0.45, 0.45))
    n = len(Wg)
    L = {"kind": "label"}
    H = {"kind": "head"}
    p.row([(n, "2.  예 정 공 정 표", {"kind": "title", "size": 16})], h=34)
    p.row([(1, "착 공 일", L), (2, I("착공일", ex=D(2026, 4, 1), fmt="date", align="center")),
           (NB + 1, "※ 착공일을 적으면 달 눈금이, 공종마다 착수·완료일을 적으면 막대가 그려집니다.", {"kind": "note"})], h=22)
    C0 = 4
    p.row([(1, "공    종", {"kind": "head", "rs": 2}), (1, "착 수", {"kind": "head", "rs": 2}), (1, "완 료", {"kind": "head", "rs": 2})]
          + [(2, F(f"월#{m + 1}", f'=IF(N({{착공일}})=0,"",EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m}))', fmt='yy"."m;;""',
                   align="center", size=8), H) for m in range(NB // 2)]
          + [(1, "비 고", {"kind": "head", "rs": 2})], h=20)
    p.row([(1, "상" if k % 2 == 0 else "하", {"kind": "head", "size": 8}) for k in range(NB)], h=15, start=C0)
    rs_, re_ = p.r, p.r + 1
    for rr, kind in ((rs_, "s"), (re_, "e")):
        for k in range(NB):
            m, half = k // 2, k % 2
            base = f"EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m})"
            f = (f"=IF(N({{착공일}})=0,\"\",{base}{'+15' if half else ''})" if kind == "s" else
                 (f"=IF(N({{착공일}})=0,\"\",{base}+14)" if half == 0 else
                  f"=IF(N({{착공일}})=0,\"\",EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m + 1})-1)"))
            p.put(rr, C0 + k, C0 + k, F(f"{kind}{k + 1}", f, fmt="yyyy-mm-dd"), kind="free", size=7)
        p.ws.row_dimensions[rr].height = 12
        p.ws.row_dimensions[rr].hidden = True
    p.r = re_ + 1
    r0 = p.r
    for i in range(1, 9):
        g = EX_GJ[i - 1] if i <= len(EX_GJ) else (None, None, None, None)
        p.row([(1, I(f"공종#{i}", ex=g[0], size=9.5)), (1, I(f"착수#{i}", ex=g[1], fmt="ymd", size=9, align="center")),
               (1, I(f"완료#{i}", ex=g[2], fmt="ymd", size=9, align="center"))]
              + [(1, "", {"kind": "text"}) for _ in range(NB)] + [(1, I(f"공정비고#{i}", ex=g[3] or None, size=8.5))], h=26)
    r1 = p.r - 1
    p.ws.conditional_formatting.add(
        f"{CL(C0)}{r0}:{CL(C0 + NB - 1)}{r1}",
        Rule(type="expression", dxf=DifferentialStyle(fill=BAR),
             formula=[f"AND(N($B{r0})>0,N($C{r0})>=N($B{r0}),$B{r0}<={CL(C0)}${re_},$C{r0}>={CL(C0)}${rs_})"]))
    p.row([(n, "※ 오랫동안 설치해 두어 오탁방지막의 평균 수명(2~3년)을 넘기면 다시 설치합니다.", {"kind": "note"})], h=16)
    d.done(p)
    return p


def org(bk, ex):
    p = _pg(bk, ex, "조직도", "3.조직도")
    d.h1(p, "3.  조 직 도")
    p.gap(10)
    H = {"kind": "head"}
    p.row([(1, "", {"kind": "free"}), (3, "직  책", H), (3, "성  명", H), (4, "맡은 일", H), (1, "", {"kind": "free"})], h=24)
    jobs = ["공사 총괄 · 감독(감리) 협의", "작업 지휘 · 안전", "작업 배치 · 자재", "수중 설치 · 점검", "수중 설치 · 점검", "육상 조립 · 운반 보조"]
    for i, (job, name) in enumerate(ORG, 1):
        p.row([(1, "", {"kind": "free"}), (3, job, {"kind": "label"}), (3, I(f"성명#{i}", ex=name, align="center")),
               (4, jobs[i - 1], {"kind": "text", "size": 9.5}), (1, "", {"kind": "free"})], h=26)
    p.gap(14)
    d.flow(p, [("공사 과장", None), ("작업반 소장", None), ("작업반장", None), ("잠수사 · 특별인부", None)], span=4, gap_h=12)
    d.done(p)
    return p


def heureum(bk, ex):
    p = _pg(bk, ex, "흐름도", "4.흐름도")
    d.h1(p, "4.  시공 흐름도")
    p.gap(8)
    d.flow(p, [("육상 조립 (40 ~ 50 SPAN)", "설치 위치 측량"), ("바지 선적 및 해상 운반", None),
               ("단계별 닻가지형 앵커 거치", "해상 안전시설 설치"), ("월 1회 점검", None), ("오탁방지막 보수", None),
               ("오탁방지막 철수", None)], span=6)
    d.done(p)
    return p


def sigong(bk, ex):
    p = _pg(bk, ex, "시공계획", "5.공종별시공계획")
    p.ws.page_setup.fitToHeight = 0
    d.h1(p, "5.  공종별 시공계획")
    d.h2(p, "5.1 준비공")
    d.para(p, "공사 착수 전 설치 구간의 수심·조류·항로와 주변 어장·시설을 조사하고, 오탁방지막 설치 위치와 방법을 정해 "
              "감독(감리)의 승인을 받은 뒤 작업합니다.")
    d.h2(p, "5.2 재료 반입")
    d.para(p, "1) 품질 기준", indent=1)
    d.para(p, "오탁방지막은 물속과 햇빛에 노출된 상태에서도 내구성이 강하고 투수성·여과성이 좋아 물속의 흙탕물이 퍼지는 것을 막을 수 있는 재료를 씁니다.")
    d.para(p, "2) 오탁방지막 구조 및 형상", indent=1)
    d.items(p, ["막체는 운반과 설치가 쉬워야 하고, 이음부는 보강 벨트를 두 겹으로 댑니다.",
                "조류·파랑에도 막체가 말려 올라가거나 뒤틀리지 않고 주름 없는 평면 모양을 유지해 오탁이 퍼지는 것을 막습니다.",
                "막체와 막체, 막체와 부자재(부체·체인·샤클)는 단단히 연결합니다.",
                "앵커(닻)는 조류·조위·파랑에 막체가 밀리지 않게 붙잡을 수 있는 크기로 하고, 앵커와 막체를 잇는 와이어로프는 이 힘을 견디는 강도로 합니다."],
            marks=["①", "②", "③", "④"], lead=2)
    d.keep(p, 120)
    d.para(p, "3) 막체의 재질 및 성능 — 직포 토목섬유 성능 기준", indent=1)
    heads = [("종류", 1), ("형별", 1), ("중량\n(g/㎡)", 1), ("인장강도\n(kg/in)", 1), ("인장신도\n(%)", 1), ("봉합강도\n(kg/in)", 1),
             ("투수계수\n(cm/sec)", 1), ("비중", 1), ("인열강도\n(kg)", 1), ("재질", 1), ("수축률\n(%)", 1)]
    d.table(p, heads, [PERF], h=40, head_h=34, size=8.5)
    p.row([(PC, "※ 시험성적서를 감독(감리)에 내고 승인을 받은 뒤 씁니다.", {"kind": "note"})], h=16)
    d.keep(p, 330)
    d.para(p, "4) 부속재료", indent=1)
    d.para(p, "부속재료는 바닷물에 부식되지 않는 재질로 하고, 조류에 풀리거나 떨어져 나가지 않는 구조로 만듭니다.")
    d.table(p, [("부자재", 2), ("재질", 2), ("규격", 2), ("용도", 6)], [[(a, {"kind": "label"}), b, c, e] for a, b, c, e in PARTS], h=22, size=9)
    d.para(p, "5) 오탁방지막 보관", indent=1)
    d.para(p, "막체는 햇빛(자외선)을 오래 받으면 강도가 떨어질 수 있으므로 창고에 보관하고, 밖에 둘 때는 천막 등으로 덮어 바람·비·햇빛을 막으며, "
              "시공 직전에 꺼내 씁니다.")
    d.h2(p, "5.3 시공 방법")
    d.items(p, ["측량 — 설계도면에 따라 기점·꼭지점·종점 등 설치 위치를 측량한 뒤 부표를 띄워 위치를 표시합니다.",
                "자재가 들어오면 육상에서 40 ~ 50 SPAN 씩 완전히 조립해 두고, 바다로 옮기기 쉽게 배치합니다.",
                "해상의 기점·꼭지점·종점과 각 위치에 잠수사를 투입해 정해진 앵커(닻)를 심습니다.",
                "기점에 앵커가 설치되고 오탁방지막이 약 30 SPAN 씩 조립되어 준비가 끝나면, 예선(150HP)을 대어 조립된 막을 바다로 끌어냅니다.",
                "끌어낸 오탁방지막은 P.P 로프로 가설치합니다.",
                "앵커는 육상에서 크레인(25톤)으로 바지(5m×5m)에 실어 나릅니다.",
                "3)~5)와 같은 방법으로 가설치를 마치면 설계도면과 같게 10m 간격으로 앵커를 설치하고, 와이어와 부속자재를 단단히 체결합니다."],
            lead=1)
    d.done(p)
    return p


def tuip(bk, ex):
    p = _pg(bk, ex, "투입계획", "6.투입계획")
    d.h1(p, "6.  인원 · 장비 · 자재 투입계획")
    d.h2(p, "6.1 작업인원 투입계획")
    H = {"kind": "head"}
    p.row([(2, "직  종", H), (1, "단위", H), (1, "일투입", H), (2, "투입일수", H), (2, "누  계", H), (4, "작 업 내 용", H)], h=24)
    r0 = p.r
    for i in range(1, 7):
        m = MAN[i - 1] if i <= len(MAN) else (None, "인", None, None, None)
        p.row([(2, I(f"직종#{i}", ex=m[0], blank=(MAN[i - 1][0] if i <= len(MAN) else None), align="center")),
               (1, I(f"인단위#{i}", ex=m[1], blank=m[1] if i <= len(MAN) else None, align="center")),
               (1, I(f"인일#{i}", ex=m[2], fmt='#,##0;;""', align="center")), (2, I(f"인일수#{i}", ex=m[3], fmt='#,##0"일";;""', align="center")),
               (2, F(f"인누계#{i}", f'=IF(OR(N({{인일#{i}}})=0,N({{인일수#{i}}})=0),"",{{인일#{i}}}*{{인일수#{i}}})', fmt='#,##0"인";;""', align="center")),
               (4, I(f"인내용#{i}", ex=m[4], blank=(MAN[i - 1][4] if i <= len(MAN) else None), size=9))], h=24)
    p.k["인누계[]"] = f"G{r0}:G{p.r - 1}"
    S = {"kind": "sum"}
    p.row([(6, "계", S), (2, F("합인", '=IF(SUM({인누계[]})=0,"",SUM({인누계[]}))', fmt='#,##0"인";;""', align="center", bold=True)), (4, "", S)], h=24)
    d.h2(p, "6.2 장비 투입계획")
    p.row([(2, "장 비 명", H), (2, "규  격", H), (1, "단위", H), (1, "일투입", H), (2, "투입일수", H), (2, "누  계", H), (2, "비  고", H)], h=24)
    r0 = p.r
    for i in range(1, 7):
        e = EQ[i - 1] if i <= len(EQ) else (None, None, "대", None, None, None)
        bl = EQ[i - 1] if i <= len(EQ) else (None,) * 6
        p.row([(2, I(f"장비#{i}", ex=e[0], blank=bl[0], align="center")), (2, I(f"규격#{i}", ex=e[1], blank=bl[1], align="center")),
               (1, I(f"장단위#{i}", ex=e[2], blank=bl[2], align="center")), (1, I(f"장일#{i}", ex=e[3], fmt='#,##0;;""', align="center")),
               (2, I(f"장일수#{i}", ex=e[4], fmt='#,##0"일";;""', align="center")),
               (2, F(f"장누계#{i}", f'=IF(OR(N({{장일#{i}}})=0,N({{장일수#{i}}})=0),"",{{장일#{i}}}*{{장일수#{i}}})', fmt='#,##0"대";;""', align="center")),
               (2, I(f"장비고#{i}", ex=e[5] or None, size=9))], h=24)
    p.k["장누계[]"] = f"I{r0}:I{p.r - 1}"
    p.row([(8, "계", S), (2, F("합장비", '=IF(SUM({장누계[]})=0,"",SUM({장누계[]}))', fmt='#,##0"대";;""', align="center", bold=True)), (2, "", S)], h=24)
    d.h2(p, "6.3 자재")
    d.para(p, "오탁방지막(막체) · 부체 · 체인 · 샤클 · 앵커(닻) · 와이어로프 — 수량은 1. 공사개요의 SPAN 수와 앵커 수를 따릅니다.")
    p.row([(PC, "※ 작업 인원과 장비 투입은 현장 여건과 투입 날짜에 따라 바뀔 수 있습니다. 누계 = 일투입 × 투입일수.", {"kind": "note"})], h=18)
    d.done(p)
    return p


def munje(bk, ex):
    p = _pg(bk, ex, "문제점", "7.문제점및대책", landscape=True)
    d.h1(p, "7.  오탁방지막 설치 문제점 및 대책 검토안")
    H = {"kind": "head"}
    p.row([(2, "구  분", H), (4, "문 제 점", H), (4, "대 책 검 토 안", H), (2, "비  고", H)], h=24)
    for i in range(1, 7):
        x = PROB[i - 1] if i <= len(PROB) else (None, None, None)
        lines = max(2, max((len(t or "") // 34 + 1) for t in x[1:]) if ex and x[1] else 2)
        p.row([(2, I(f"구분#{i}", ex=x[0], align="center")), (4, I(f"문제#{i}", ex=x[1], size=9)), (4, I(f"대책#{i}", ex=x[2], size=9)),
               (2, I(f"문제비고#{i}", ex=("여건이 바뀌면 설치 위치·수심을 다시 검토해 계획서를 고침" if ex and x[0] else None), size=8.5))],
              h=max(40, lines * 14 + 8))
    d.done(p)
    return p


def pumjil(bk, ex):
    p = _pg(bk, ex, "품질안전", "품질·환경·안전")
    d.h1(p, "Ⅱ.  품 질 관 리")
    d.items(p, ["오탁방지막은 물속과 햇빛에 노출된 상태에서도 내구성이 강하고 투수성이 좋으며 오탁 확산을 막을 수 있는 재료로 합니다.",
                "오탁방지막은 성능 기준표(5.2의 3))의 규격을 만족해야 합니다.",
                "오탁방지막은 시험성적서를 감독(감리)에 내고 승인을 받은 뒤 씁니다.",
                "설치할 때 막체가 찢어지거나 상하지 않도록 품질관리에 힘씁니다."])
    d.h1(p, "Ⅲ.  환경 · 안전관리계획")
    d.items(p, ["새로 들어온 근로자는 건강진단과 신규 채용 교육(1시간 이상)을 받은 뒤 작업에 들어갑니다.",
                "매일 작업 전 안전조회로 건강 상태를 확인한 뒤 배치합니다.",
                "앵커(닻)는 무거우므로 운반·설치 때 끼임과 허리 부상에 특히 주의합니다.",
                "안전관리 책임자는 늘 일기예보를 확인하고, 날씨가 나빠지면 자재를 육상의 안전한 곳으로 옮깁니다.",
                "물 위에서 일할 때는 구명조끼를 입습니다.",
                "잠수사는 작업 전에 충분히 준비운동을 합니다.",
                "작업 중 기름 등으로 바다가 오염되지 않도록 특히 주의합니다.",
                "작업 중 생긴 폐기물은 법에 정한 절차대로 처리합니다."])
    p.gap(8)
    L = {"kind": "label"}
    p.row([(3, "해상작업 중지 기준", L), (2, "파고", L), (2, I("파고", ex=1.0, blank=1.0, fmt='0.0"m 이상"', rate=True, align="center")),
           (2, "유속", L), (3, I("유속", ex=1.0, blank=1.0, fmt='0.0"m/s 이상"', rate=True, align="center"))], h=26)
    p.row([(PC, "※ 노란 칸은 현장·발주기관 기준에 맞게 고쳐 씁니다.", {"kind": "note"})], h=16)
    d.done(p)
    p.after_note(["이 장은 빈 서식과 작성 예시가 같습니다(공법 설명)."] if not ex else [""])
    return p


def draw(bk, ex):
    if not ex:
        cover(bk, ex); toc(bk, ex); gaeyo(bk, ex); _gantt(bk, ex); org(bk, ex); heureum(bk, ex); sigong(bk, ex); tuip(bk, ex); munje(bk, ex); pumjil(bk, ex)
    else:
        cover(bk, ex); gaeyo(bk, ex); _gantt(bk, ex); org(bk, ex); tuip(bk, ex); munje(bk, ex)


def expect(pages):
    import math
    g = pages["공사개요"]
    e = {"개요글": f"    본 공사는 「{pages['표지']['공사명']}」 기간 중 준설과 기초공사로 생기는 오탁수와 부유물질이 퍼져 바다(하천)가 오탁되는 것을 막기 위하여 오탁방지막을 설치하는 공사입니다."}
    ts = ta = tl = 0
    for i in range(1, 6):
        L = g.get(f"연장#{i}")
        if L:
            s = math.ceil(L / g["막길이"])
            a = (math.ceil(L / g["앵커간격"]) + 1) * g["앵커줄"]
            e[f"SPAN#{i}"], e[f"앵커#{i}"] = s, a
            ts += s; ta += a; tl += L
        else:
            e[f"SPAN#{i}"] = e[f"앵커#{i}"] = ""
    e.update({"합연장": tl, "합SPAN": ts, "합앵커": ta})
    t = pages["투입계획"]
    x = {}
    si = se = 0
    for i in range(1, 7):
        a, b = t.get(f"인일#{i}"), t.get(f"인일수#{i}")
        x[f"인누계#{i}"] = a * b if a and b else ""
        si += a * b if a and b else 0
        a, b = t.get(f"장일#{i}"), t.get(f"장일수#{i}")
        x[f"장누계#{i}"] = a * b if a and b else ""
        se += a * b if a and b else 0
    x.update({"합인": si, "합장비": se})
    return {"공사개요": e, "투입계획": x}
