# -*- coding: utf-8 -*-
"""건설기계 점검표(31종 + 일일 안전점검표) — 원본(엑셀 32쪽: 기계마다 큰 사진 · No · 점검항목[적합 ○ / 부적합 ×] · 결과 · 참조 사진 ·
차량번호 · 규격 · 확인(점검자 · 운전원) · 점검일 · 조치사항, 끝에 한 달치 «일일 안전점검표»)을 같은 틀로 새로 만든 것 (2026-09-29).

■ 사진: 원본 설명 사진 그대로(소장님 「설명에 필요한 사진은 그대로 쓸 것」). 상표 글자 · 번호판 숫자만 흐리게 가림(13장).
  사진 파일 img/o-geonseolgigye-check/NN-p1.jpg(큰 사진) · NN-r3L.jpg(참조 사진), 자리 · 설명글은 data/o-geonseolgigye-check.json.
■ 점검항목 글은 원본 그대로(줄바꿈도). 바로잡은 곳: 12-⑦ «일치할» → «일치할 것», 16-① «경보과» → «경보와», 19-⑩ 뒤섞인 줄 차례,
  22-④ «있을» → «있을 것», 23-① «블레이더» → «블레이드», 11-⑥ 줄 차례.
■ 수식(기계마다): 결과(○ · × · －)를 고르면 맨 위 «적합 · 부적합 · 해당없음 · 미점검» 건수와 판정(사용 가능 / 조치 후 사용 / 점검 중)이 저절로.
  부적합인데 조치사항이 비었거나, 점검했는데 점검일이 비었으면 알려 줌.
■ 일일 안전점검표: 장비명을 고르면 ➀~➉항 점검항목 글이 저절로 들어오고, 날마다 ○ · △ · × 를 고르면 항목별 · 날짜별 △ · × 건수,
  이달 점검한 날 수, 불량 난 날 · 서명 빠진 날 · 그 달에 없는 날짜 표시를 셈함."""
import datetime
import io
import json
import os

from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image as PImage

import docw as d
from fb import F, I, text_lines

SLUG = "o-geonseolgigye-check"
TITLE = "건설기계 점검표"
WHERE = "orig"
PREV = [(2, "굴착기(백호) — 빈 서식"), (34, "굴착기(백호) — 작성 예시"), (36, "일일 안전점검표 — 작성 예시"), (1, "목차 — 누르면 그 장비로 감")]
PAGES = 36
NSHEETS = 5
MULTI = True
SHORT = ("굴착기 · 덤프트럭 · 크레인 · 고소작업대 등 건설기계 31종 점검표 + 한 달치 일일 안전점검표. 결과를 ○ · × 로 고르면 판정(사용 가능 / 조치 후 사용)이 "
         "저절로 나오고, 일일 점검표는 장비명을 고르면 점검항목이 들어옵니다.")
NOTE = "설명 사진은 원본 그대로 넣었습니다(상표 글자 · 번호판 숫자만 흐리게 가림). 점검항목 글도 원본 그대로이며 오타 몇 곳만 바로잡았습니다."

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img", SLUG)
DATA = json.load(open(os.path.join(HERE, "data", f"{SLUG}.json"), encoding="utf-8"))
D = datetime.date
SHEET = "건설기계 점검표"
DAILY = "일일 안전점검표"
NUM = "①②③④⑤⑥⑦⑧⑨⑩"
NUMD = "➀➁➂➃➄➅➆➇➈➉"
BLANK_D = "20   .   .   ."
BAR = PatternFill("solid", fgColor="2B579A")
WHITE = "FFFFFF"
L = {"kind": "label"}
HD = {"kind": "head"}

# 점검표 열: 왼쪽 사진 · 아래 차량번호/확인(a~h) | No · 점검항목 · 결과 | 참조 사진 두 칸
W = [4, 9, 13, 8, 9, 10, 7, 4, 4.5, 42, 6.5, 13.5, 13.5]
N = len(W)
CA, CI, CJ, CK, CLL, CM = 1, 9, 10, 11, 12, 13
ITEM_BUDGET = 326          # 한 쪽에 들어가는 항목 줄 높이 합(pt) — 머리 · 확인 · 경고 줄을 뺀 나머지
CAP_H = 13

# 작성 예시 — (장비 번호, 결과 {항목: 값}, 조치사항, 차량번호, 규격, 점검자1, 점검자2, 운전원, 점검일)
EXS = [(1, {7: "×"}, "⑦ 웨이트 고정볼트 1개 풀림 → 현장에서 다시 조이고(4/7 10:20) 재점검 결과 적합", "가나 04가 1234", "0.8㎥ (21톤급)",
        "박안전", "이감독", "김운전", D(2026, 4, 7)),
       (6, {}, None, "가나 80바 5678", "6㎥", "박안전", "이감독", "홍길동", D(2026, 4, 8))]


def _read(fn):
    with open(os.path.join(IMG, fn), "rb") as f:
        return f.read()


def _place_px(p, data, c1, r1, x, y, w, h):
    """(c1, r1) 칸 왼쪽 위에서 (x, y) px 떨어진 곳에 w × h px 상자 — 그 안에 비율을 지켜 가운데 맞춤."""
    iw, ih = PImage.open(io.BytesIO(data)).size
    sc = min(w / iw, h / ih)
    ww, hh = iw * sc, ih * sc
    ox, oy = x + (w - ww) / 2, y + (h - hh) / 2
    c, r = c1, r1
    while ox >= p.col_px(c, c):
        ox -= p.col_px(c, c)
        c += 1
    while oy >= p.row_px(r):
        oy -= p.row_px(r)
        r += 1
    img = XLImage(io.BytesIO(data))
    img.width, img.height = ww, hh
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=c - 1, colOff=pixels_to_EMU(int(ox)), row=r - 1, rowOff=pixels_to_EMU(int(oy))),
                               ext=XDRPositiveSize2D(pixels_to_EMU(int(ww)), pixels_to_EMU(int(hh))))
    p.ws.add_image(img)


def _heights(items):
    """항목 줄 높이 — 글 줄 수로 어림하고 남는 높이는 고르게 나눔. 짝수 항목은 설명글 줄(CAP_H)과 합친 높이."""
    for size, lh in ((9, 12.2), (8.5, 11.4), (8, 10.8)):
        need = []
        for i, t in enumerate(items):
            n = text_lines(t, W[CJ - 1] * 10.0 / size) if t else 1
            h = max(26, n * lh + 7)
            need.append(h - (CAP_H if i % 2 else 0))
        if sum(need) <= ITEM_BUDGET:
            break
    extra = (ITEM_BUDGET - sum(need)) / 10
    return size, [h + extra for h in need]


def form(p, o, ex=None):
    """기계 한 장. ex = (결과, 조치, 차량번호, 규격, 점검자1, 점검자2, 운전원, 점검일) 또는 None"""
    j = o["n"]
    res, fix, carno, spec, s1, s2, drv, day = ex if ex else ({}, None, None, None, None, None, None, None)
    items = o["items"]
    size, hs = _heights(items)
    bar = {"fill": BAR, "ink": WHITE, "border": True}
    p.row([(8, "■ " + o["name"], dict(bar, kind="text", size=13, bold=True, indent=1, wrap=False)),
           (2, F(f"요약{j}", "", align="right", size=9.5), dict(bar, kind="text")),
           (1, "판정", dict(bar, kind="ctext", size=10, bold=True)),
           (2, F(f"판정{j}", "", align="center", bold=True, size=11))], h=26)
    top = p.row([(8, "", {"rs": 14}), (1, "No", HD), (1, "점검항목  [ 적합 : ○  /  부적합 : ×  /  해당없음 : － ]", dict(HD, size=9.5)),
                 (1, "결과", HD), (2, "참        조", HD)], h=22)
    rows = []
    refs = o["refs"] + [[None, None]] * (5 - len(o["refs"]))
    for k in range(5):
        a, b = 2 * k, 2 * k + 1
        cells = []
        r1 = p.r
        for i, rs in ((a, 1), (b, 2)):
            t = items[i]
            txt = (t, {"kind": "text", "size": size, "rs": rs}) if t else (I(f"항목{j}#{i + 1}", size=size), {"rs": rs})
            row = [(1, NUM[i], {"kind": "ctext", "size": 11, "rs": rs}), (1,) + txt,
                   (1, I(f"결과{j}#{i + 1}", ex=res.get(i + 1, "○") if (ex and t) else None, dv="○,×,－", align="center", size=12), {"rs": rs})]
            if i == a:
                row += [(1, "", {"rs": 2}), (1, "", {"rs": 2})]
            if k == 4 and i == b:
                row = [(2, "차 량 번 호", dict(L, rs=2)), (2, I(f"차량번호{j}", ex=carno, align="center"), {"rs": 2}),
                       (1, "규  격", dict(L, rs=2)), (3, I(f"규격{j}", ex=spec, align="center", size=9.5), {"rs": 2})] + row
            p.row(row, h=hs[i])
            rows.append(p.r - 1)
        caps = []
        for side in (0, 1):
            rf = refs[k][side]
            caps.append((1, (rf or {}).get("cap") or "", {"kind": "ctext", "size": 8}))
        p.row(caps, h=CAP_H)
        for side, c in ((0, CLL), (1, CM)):
            rf = refs[k][side]
            if rf and rf.get("f"):
                p.place_image(_read(rf["f"]), c, c, r1, r1 + 1, pad=3)
    last = p.r - 1
    # 왼쪽 큰 사진(원본 자리 비율 그대로)
    bw = p.col_px(1, 8) - 12
    bh = sum(p.row_px(r) for r in range(top, rows[8] + 1)) - 12
    for bg in o["big"]:
        x, y, w, h = bg["box"]
        _place_px(p, _read(bg["f"]), 1, top, 6 + x * bw + 2, 6 + y * bh + 2, w * bw - 4, h * bh - 4)
    # 확인 · 조치사항
    p.row([(1, "확\n인", dict(L, rs=2)), (1, "점 검 자", L), (1, I(f"점검자{j}a", ex=s1, align="center")), (1, "(서명) /", {"kind": "ctext", "size": 9}),
           (2, I(f"점검자{j}b", ex=s2, align="center")), (2, "(서명)", {"kind": "ctext", "size": 9}),
           (1, "조치\n사항", dict(L, rs=2, size=9.5)), (4, I(f"조치{j}", ex=fix, size=9), {"rs": 2, "valign": "top"})], h=27)
    p.row([(1, "운 전 원", L), (1, I(f"운전원{j}", ex=drv, align="center")), (1, "(서명)", {"kind": "ctext", "size": 9}),
           (1, "점 검 일", L), (3, I(f"점검일{j}", ex=day, fmt="ymd", blank=BLANK_D, align="center"))], h=27, start=2)
    rg = f"{{결과{j}[]}}"
    tx = f"{{글{j}[]}}"
    p.k[f"결과{j}[]"] = f"{CL(CK)}{rows[0]}:{CL(CK)}{last}"
    p.k[f"글{j}[]"] = f"{CL(CJ)}{rows[0]}:{CL(CJ)}{last}"
    a_, b_, c_ = f'COUNTIF({rg},"○")', f'COUNTIF({rg},"×")', f'COUNTIF({rg},"－")'
    m_ = f"SUMPRODUCT((LEN({tx})>0)*(LEN({rg})=0))"      # 글이 있는데 결과가 빈 항목(COUNTIFS «<>» 는 계산기마다 달라 LEN 으로)
    any_ = f"COUNTA({rg})=0"
    p.calc_keys[f"요약{j}"] = None
    _setf(p, f"요약{j}", f'=IF({any_},"","적합 "&{a_}&" · 부적합 "&{b_}&IF({c_}>0," · 해당없음 "&{c_},"")&IF({m_}>0," · 미점검 "&{m_},""))')
    _setf(p, f"판정{j}", f'=IF({any_},"",IF({b_}>0,"조치 후 사용",IF({m_}>0,"점검 중","사용 가능")))')
    dd = f"{{점검일{j}}}"
    p.row([(N, F(f"경고{j}", f'=IF({any_},"",IF(AND({b_}>0,{{조치{j}}}=""),"⚠ 부적합(×) "&{b_}&"건 — 조치사항 칸에 무엇을 언제 고쳤는지 적으세요.",'
                             f'IF(OR({dd}&""="",{dd}&""="{BLANK_D}"),"⚠ 점검일을 적으세요.","")))', align="left"), {"kind": "warn"})], h=15)
    # 결과 × 빨강 · 판정 색
    ws = p.ws
    ws.conditional_formatting.add(p.k[f"결과{j}[]"], FormulaRule(formula=[f'{CL(CK)}{rows[0]}="×"'], font=Font(color="DC2626", bold=True)))
    pk = p.k[f"판정{j}"]
    ws.conditional_formatting.add(pk, FormulaRule(formula=[f'{pk}="조치 후 사용"'], font=Font(color="DC2626", bold=True)))
    ws.conditional_formatting.add(pk, FormulaRule(formula=[f'{pk}="사용 가능"'], font=Font(color="15803D", bold=True)))
    return rows


def _setf(p, key, f):
    for i, (cell, t) in enumerate(p.formulas):
        if cell.coordinate == p.k[key]:
            p.formulas[i] = (cell, f)
            p.calc_keys[key] = f
            return
    raise KeyError(key)


def sheets(bk, ex):
    p = bk.page(SHEET, W, ex=ex, landscape=True, fit_height=0, sheetname=SHEET if not ex else "예시-" + SHEET, margins=(0.3, 0.3, 0.35, 0.35))
    rows = {}
    todo = [(o, None) for o in DATA] if not ex else [(DATA[n - 1], (r, fx, cn, sp, a, b, dr, dy)) for n, r, fx, cn, sp, a, b, dr, dy in EXS]
    for k, (o, e) in enumerate(todo):
        if k:
            p.__dict__.setdefault("_brk", set()).add(p.r)
        rows[o["n"]] = p.r
        form(p, o, e)
    d.done(p)
    p.after_note(["결과 칸을 누르면 목록(○ 적합 · × 부적합 · － 해당없음)이 나옵니다. 하나라도 × 이면 맨 위 판정이 «조치 후 사용» 으로 바뀝니다.",
                  "빈 항목 줄(⑦~⑩ 등)에는 현장에서 더 볼 항목을 적어 쓰세요 — 적으면 미점검 건수에도 들어갑니다.",
                  "사진은 점검 부위를 알려 주는 참고 사진입니다(원본 그대로, 상표 글자 · 번호판 숫자만 가림). 필요한 장비만 «목차» 에서 눌러 가서 그 쪽만 인쇄하세요."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 장비 · 사람) — 두 가지 장비만 채워 보았습니다."])
    return p, rows


# ── 일일 안전점검표 ─────────────────────────────────────────────────────
DW = [6.5, 9.5, 24] + [3.1] * 31 + [4.2, 4.2]
DPC = len(DW)
HELP0 = DPC + 2              # 도움 표(장비명 + 항목 10개) 시작 열 — 인쇄되지 않고 숨김
DC1, DC31 = 4, 34
EXD = {"파트너사": "예시장비(주)", "장비명": "굴착기(백호)", "등록번호": "가나 04가 1234", "년월": D(2026, 4, 1),
       "빈날": {5, 12, 19, 26, 31}, "△": {(7, 7), (21, 2)}, "서명": "박안전"}


def _short(t):
    """일일 점검표용 — 줄바꿈을 잇고, 적어 넣는 칸(«:      ton» 같은) 줄은 뺌."""
    if not t:
        return ""
    keep = [ln.strip() for ln in t.split("\n") if not ((":" in ln or ";" in ln) and "  " in ln)]
    s = " ".join(keep)
    return s.replace(" ,", ",").strip()


def daily(bk, ex):
    widths = DW + [2, 24] + [30] * 10
    p = bk.page(DAILY, widths, ex=ex, landscape=True, fit_height=1, sheetname=DAILY if not ex else "예시-" + DAILY,
                margins=(0.3, 0.3, 0.4, 0.4), print_cols=DPC)
    e = EXD if ex else {}
    p.row([(DPC, "일 일  안 전 점 검 표", {"kind": "title", "size": 18})], h=34)
    p.row([(DPC, "건설기계 한 대에 한 달 한 장 — 날마다 작업 전에 ➀~➉항을 점검해 ○ · △ · × 를 고르세요(항목 글은 «건설기계 점검표» 의 그 장비 항목).",
            {"kind": "sub", "size": 9})], h=16)
    p.gap(6)
    lg = "양    호 : ○\n수리필요 : △\n불    량 : ×"
    p.row([(2, "파 트 너 사", L), (7, I("파트너사", ex=e.get("파트너사"))), (6, "장  비  명", L),
           (12, I("장비명", ex=e.get("장비명"), align="center")), (9, lg, {"kind": "text", "size": 8.5, "rs": 2, "indent": 4})], h=26)
    p.row([(2, "등 록 번 호", L), (7, I("등록번호", ex=e.get("등록번호"))), (6, "점 검 년 월", L),
           (12, I("년월", ex=e.get("년월"), fmt='yyyy"년" m"월"', align="center"))], h=26)
    p.gap(6)
    hdr = p.row([(1, "날짜\n내용", dict(HD, size=8)), (2, "점 검 항 목  (장비명을 고르면 저절로 들어옴)", dict(HD, size=9))]
                + [(1, dd, dict(HD, size=8)) for dd in range(1, 32)] + [(1, "△\n계", dict(HD, size=8)), (1, "×\n계", dict(HD, size=8))], h=26)
    # 도움 표 — 장비명 · 항목 글(숨김 열)
    ws = p.ws
    for i, o in enumerate(DATA):
        ws.cell(row=hdr + 1 + i, column=HELP0, value=o["name"])
        for k in range(10):
            v = _short(o["items"][k])
            if v:
                ws.cell(row=hdr + 1 + i, column=HELP0 + 1 + k, value=v)
    ws.cell(row=hdr, column=HELP0, value="(도움 표: 장비명)")
    for c in range(HELP0 - 1, HELP0 + 11):
        ws.column_dimensions[CL(c)].hidden = True
    nm_rng = f"${CL(HELP0)}${hdr + 1}:${CL(HELP0)}${hdr + len(DATA)}"
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1="=" + nm_rng, allow_blank=True, showDropDown=False)   # 31가지 — 글 목록(255자) 대신 도움 표 범위
    dv.error = "목록에서 고르세요"
    dv.showErrorMessage = False
    ws.add_data_validation(dv)
    dv.add(p.k["장비명"])
    th = max(text_lines(_short(o["items"][k]), 33.5 * 10 / 7.5) for o in DATA for k in range(10))
    rh = max(30, min(th, 3) * 10.6 + 5)
    grid = []
    for k in range(10):
        it_rng = f"${CL(HELP0 + 1 + k)}${hdr + 1}:${CL(HELP0 + 1 + k)}${hdr + len(DATA)}"
        r = p.r
        cells = [(1, NUMD[k] + "항", {"kind": "label", "size": 9.5}),
                 (2, F(f"항목글{k + 1}", f'=IFERROR(INDEX({it_rng},MATCH({{장비명}},{nm_rng},0))&"","")', size=7.5, align="left"),
                  {"valign": "center"})]
        for dd in range(1, 32):
            v = None
            if ex and dd not in e["빈날"]:
                v = "△" if (dd, k + 1) in e["△"] else "○"
            cells.append((1, I(f"일{k + 1}#{dd}", ex=v, dv="○,△,×", align="center", size=10)))
        cells += [(1, F(f"△{k + 1}", f'=COUNTIF({CL(DC1)}{r}:{CL(DC31)}{r},"△")', fmt="int", align="center", size=9)),
                  (1, F(f"×{k + 1}", f'=COUNTIF({CL(DC1)}{r}:{CL(DC31)}{r},"×")', fmt="int", align="center", size=9))]
        p.row(cells, h=rh)
        grid.append(r)
    g1, g10 = grid[0], grid[-1]
    sr = p.row([(1, "점검자\n확  인", dict(L, size=8)), (2, "(점검한 날 서명)", {"kind": "ctext", "size": 8})]
               + [(1, I(f"서명#{dd}", ex=(e["서명"] if ex and dd not in e["빈날"] else None), size=7, align="center")) for dd in range(1, 32)]
               + [(1, "", {}), (1, "", {})], h=36)
    for dd in range(1, 32):
        c = ws.cell(row=sr, column=DC1 + dd - 1)
        c.alignment = Alignment(horizontal="center", vertical="center", textRotation=255, wrap_text=False)
    ar = p.row([(3, "수리필요 · 불량 건수 (자동)", dict(L, size=8.5))]
               + [(1, F(f"날△×#{dd}", f'=COUNTIF({CL(DC1 + dd - 1)}{g1}:{CL(DC1 + dd - 1)}{g10},"△")+COUNTIF({CL(DC1 + dd - 1)}{g1}:{CL(DC1 + dd - 1)}{g10},"×")',
                           fmt="int", align="center", size=8.5)) for dd in range(1, 32)]
               + [(1, F("△계", f"=SUM({CL(DPC - 1)}{g1}:{CL(DPC - 1)}{g10})", fmt="int", align="center", size=9, bold=True)),
                  (1, F("×계", f"=SUM({CL(DPC)}{g1}:{CL(DPC)}{g10})", fmt="int", align="center", size=9, bold=True))], h=22)
    p.gap(5)
    sm = p.row([(3, "이 달 합 계", L), (DPC - 3, F("합계글", "", align="left", size=10))], h=24)
    w1 = p.row([(DPC, F("경고1", "", align="left"), {"kind": "warn"})], h=15)
    w2 = p.row([(DPC, F("경고2", "", align="left"), {"kind": "warn"})], h=15)
    p.end_print()
    # 숨김 도움 줄: 날짜별 표시 칸 수 · × 수
    h1 = p.r + 1
    p.r = h1
    p.row([(3, "(도움) 표시 칸 수", {"kind": "free", "size": 8})]
          + [(1, F(f"표시#{dd}", f"=COUNTA({CL(DC1 + dd - 1)}{g1}:{CL(DC1 + dd - 1)}{g10})", fmt="int", size=8)) for dd in range(1, 32)], h=14)
    p.row([(3, "(도움) × 수", {"kind": "free", "size": 8})]
          + [(1, F(f"불량#{dd}", f'=COUNTIF({CL(DC1 + dd - 1)}{g1}:{CL(DC1 + dd - 1)}{g10},"×")', fmt="int", size=8)) for dd in range(1, 32)], h=14)
    ws.row_dimensions[h1].hidden = True
    ws.row_dimensions[h1 + 1].hidden = True
    H1 = f"{CL(DC1)}{h1}:{CL(DC31)}{h1}"
    H2 = f"{CL(DC1)}{h1 + 1}:{CL(DC31)}{h1 + 1}"
    HD_ = f"{CL(DC1)}{hdr}:{CL(DC31)}{hdr}"
    SG = f"{CL(DC1)}{sr}:{CL(DC31)}{sr}"
    GR = f"{CL(DC1)}{g1}:{CL(DC31)}{g10}"
    ym = "{년월}"
    n_ = f'COUNTIF({H1},">0")'
    _setf(p, "합계글", f'=IF({n_}=0,"","점검한 날 "&{n_}&"일 · 양호(○) "&COUNTIF({GR},"○")&"칸 · 수리필요(△) "&{{△계}}&"건 · 불량(×) "&{{×계}}&"건")')
    _setf(p, "경고1", f'=IF(COUNTIF({H2},">0")>0,"⚠ 불량(×)이 나온 날이 "&COUNTIF({H2},">0")&"일 있습니다 — 그 장비는 고친 뒤 다시 점검하고 쓰세요.","")')
    nos = f"SUMPRODUCT(({H1}>0)*(LEN({SG})=0))"
    beyond = f"IF(ISNUMBER({ym}),SUMPRODUCT(({H1}>0)*({HD_}>DAY(EOMONTH({ym},0)))),0)"
    _setf(p, "경고2", f'=IF({beyond}>0,"⚠ "&MONTH({ym})&"월은 "&DAY(EOMONTH({ym},0))&"일까지입니다 — 없는 날짜에 표시한 날이 "&{beyond}&"일 있습니다.",'
                     f'IF({nos}>0,"⚠ 점검한 날 중 점검자 확인(서명)이 빈 날이 "&{nos}&"일 있습니다.",""))')
    # 모양: ×/△ 색, 그 달에 없는 날 회색
    ws.conditional_formatting.add(GR, FormulaRule(formula=[f'{CL(DC1)}{g1}="×"'], font=Font(color="DC2626", bold=True)))
    ws.conditional_formatting.add(GR, FormulaRule(formula=[f'{CL(DC1)}{g1}="△"'], font=Font(color="C2410C", bold=True)))
    ymc = "$" + p.k["년월"][0] + "$" + p.k["년월"][1:]
    ws.conditional_formatting.add(f"{CL(DC1)}{hdr}:{CL(DC31)}{ar}",
                                  FormulaRule(formula=[f"AND(ISNUMBER({ymc}),{CL(DC1)}${hdr}>DAY(EOMONTH({ymc},0)))"],
                                              fill=PatternFill("solid", fgColor="D1D5DB", bgColor="D1D5DB")))
    p.r = h1 + 3
    p.after_note(["장비명을 목록에서 고르면 ➀~➉항 점검항목 글이 들어옵니다(«건설기계 점검표» 의 그 장비 항목, 적어 넣는 칸 줄은 뺌).",
                  "날마다 칸을 누르면 목록(○ 양호 · △ 수리필요 · × 불량)이 나옵니다. 점검 년월을 적으면 그 달에 없는 날(31일 등)은 회색이 됩니다."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 장비 · 회사 · 사람) — 2026년 4월(30일까지), 일요일은 쉼."])
    return p


def mokcha(bk, rows):
    Wm = [7, 46, 10, 10, 12]
    p = bk.page("목차", Wm, ex=False, fit_height=1, sheetname="목차", margins=(0.5, 0.5, 0.45, 0.45))
    p.row([(5, "건 설 기 계  점 검 표  목 차", {"kind": "title", "size": 18})], h=36)
    p.row([(5, f"건설기계 {len(DATA)}종 + 일일 안전점검표 — 장비 이름을 누르면 그 점검표로 갑니다.", {"kind": "sub", "size": 9.5})], h=18)
    p.gap(4)
    p.row([(1, "번호", HD), (1, "장 비 명", HD), (1, "점검항목", HD), (1, "참조 사진", HD), (1, "바로가기", HD)], h=24)
    for o in DATA:
        r = p.r
        nref = sum(1 for pr in o["refs"] for x in pr if x and x.get("f"))
        p.row([(1, o["n"], {"kind": "ctext", "size": 9.5}), (1, o["name"], {"kind": "text", "size": 10}),
               (1, f"{sum(1 for t in o['items'] if t)}항목", {"kind": "ctext", "size": 9}), (1, f"{nref}장", {"kind": "ctext", "size": 9}),
               (1, "→ 가기", {"kind": "ctext", "size": 9})], h=19)
        for c in (2, 5):
            cell = p.ws.cell(row=r, column=c)
            cell.hyperlink = f"#'{SHEET}'!A{rows[o['n']]}"
            cell.font = cell.font.copy(color="1D4ED8", underline="single")
    r = p.r
    p.row([(1, "끝", {"kind": "ctext", "size": 9.5}), (1, "일일 안전점검표 (한 달치 · ➀~➉항 × 1~31일)", {"kind": "text", "size": 10}),
           (1, "10항목", {"kind": "ctext", "size": 9}), (1, "—", {"kind": "ctext", "size": 9}), (1, "→ 가기", {"kind": "ctext", "size": 9})], h=19)
    for c in (2, 5):
        cell = p.ws.cell(row=r, column=c)
        cell.hyperlink = f"#'{DAILY}'!A1"
        cell.font = cell.font.copy(color="1D4ED8", underline="single")
    d.done(p)
    p.after_note(["설명 사진은 원본 그대로입니다(상표 글자 · 번호판 숫자만 흐리게 가림). 점검항목 글도 원본 그대로이고, 오타 몇 곳만 바로잡았습니다."])
    return p


def draw(bk, ex):
    p, rows = sheets(bk, ex)
    daily(bk, ex)
    if not ex:
        mokcha(bk, rows)
        wb = bk.wb
        first = ["목차", SHEET, DAILY]
        wb._sheets = [wb[n] for n in first] + [s for s in wb._sheets if s.title not in first]


# ── 검증용 계산(파이썬으로 따로) ─────────────────────────────────────────
def _one(inp, o):
    j = o["n"]
    vals = [inp.get(f"결과{j}#{i}") for i in range(1, 11)]
    texts = [o["items"][i - 1] or inp.get(f"항목{j}#{i}") for i in range(1, 11)]
    a = vals.count("○")
    b = vals.count("×")
    c = vals.count("－")
    m = sum(1 for t, v in zip(texts, vals) if t and not v)
    if not any(vals):
        return {f"요약{j}": "", f"판정{j}": "", f"경고{j}": ""}
    s = f"적합 {a} · 부적합 {b}" + (f" · 해당없음 {c}" if c else "") + (f" · 미점검 {m}" if m else "")
    pj = "조치 후 사용" if b else ("점검 중" if m else "사용 가능")
    dday = inp.get(f"점검일{j}")
    if b and not inp.get(f"조치{j}"):
        w = f"⚠ 부적합(×) {b}건 — 조치사항 칸에 무엇을 언제 고쳤는지 적으세요."
    elif dday in (None, "", BLANK_D):
        w = "⚠ 점검일을 적으세요."
    else:
        w = ""
    return {f"요약{j}": s, f"판정{j}": pj, f"경고{j}": w}


def _eom(dt):
    nx = D(dt.year + (dt.month == 12), dt.month % 12 + 1, 1)
    return (nx - datetime.timedelta(days=1)).day


def _daily(inp):
    out = {}
    nm = inp.get("장비명")
    o = next((x for x in DATA if x["name"] == nm), None)
    for k in range(1, 11):
        out[f"항목글{k}"] = _short(o["items"][k - 1]) if o else ""
        row = [inp.get(f"일{k}#{dd}") for dd in range(1, 32)]
        out[f"△{k}"] = row.count("△")
        out[f"×{k}"] = row.count("×")
    marks = {dd: [inp.get(f"일{k}#{dd}") for k in range(1, 11)] for dd in range(1, 32)}
    for dd in range(1, 32):
        out[f"날△×#{dd}"] = marks[dd].count("△") + marks[dd].count("×")
        out[f"표시#{dd}"] = sum(1 for v in marks[dd] if v)
        out[f"불량#{dd}"] = marks[dd].count("×")
    out["△계"] = sum(out[f"△{k}"] for k in range(1, 11))
    out["×계"] = sum(out[f"×{k}"] for k in range(1, 11))
    days = [dd for dd in range(1, 32) if out[f"표시#{dd}"] > 0]
    ok = sum(v == "○" for dd in range(1, 32) for v in marks[dd])
    out["합계글"] = (f"점검한 날 {len(days)}일 · 양호(○) {ok}칸 · 수리필요(△) {out['△계']}건 · 불량(×) {out['×계']}건") if days else ""
    xd = sum(1 for dd in range(1, 32) if out[f"불량#{dd}"] > 0)
    out["경고1"] = f"⚠ 불량(×)이 나온 날이 {xd}일 있습니다 — 그 장비는 고친 뒤 다시 점검하고 쓰세요." if xd else ""
    ym = inp.get("년월")
    beyond = 0
    if isinstance(ym, (datetime.date, datetime.datetime)):
        e = _eom(ym)
        beyond = sum(1 for dd in days if dd > e)
    nosign = sum(1 for dd in days if not inp.get(f"서명#{dd}"))
    if beyond:
        out["경고2"] = f"⚠ {ym.month}월은 {_eom(ym)}일까지입니다 — 없는 날짜에 표시한 날이 {beyond}일 있습니다."
    elif nosign:
        out["경고2"] = f"⚠ 점검한 날 중 점검자 확인(서명)이 빈 날이 {nosign}일 있습니다."
    else:
        out["경고2"] = ""
    return out


def expect(pages):
    inp = pages[SHEET]
    out = {}
    for n, *_r in EXS:
        out.update(_one(inp, DATA[n - 1]))
    return {SHEET: out, DAILY: _daily(pages[DAILY])}
