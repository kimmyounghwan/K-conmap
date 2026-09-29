# -*- coding: utf-8 -*-
"""
run.py — 서식을 굽고(빈 서식 + 작성 예시) 두 번 검증합니다.

  python3 run.py wonga chakgong …      # 고른 것만
  python3 run.py --all                 # 전부
  옵션: --out DIR (기본 /tmp/g10/f2/out) · --render (PDF·PNG 그림까지)

검증 ①  LibreOffice 로 다시 계산 → 작성 예시의 수식 결과 = 파이썬 expect() 값 (원 단위까지 똑같아야)
검증 ②  모든 시트: 오류값(#REF!·#VALUE!·#DIV/0!·#NAME?·Err:) 0, 바깥 파일 참조 0,
        빈 서식의 수식 칸은 빈칸(0 은 숨김 서식) — 글이 나오는 칸은 목록으로 보여 줌
"""
import importlib
import os
import shutil
import subprocess
import sys
import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from fb import Book  # noqa: E402

ERRS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:", "#ERR")


def build(slug, out):
    m = importlib.import_module("f_" + slug.replace("-", "_"))
    bk = Book(m.TITLE)
    m.draw(bk, False)
    m.draw(bk, True)
    os.makedirs(out, exist_ok=True)
    path = os.path.join(out, f"{m.SLUG}.xlsx")
    bk.save(path)
    return m, bk, path


def recalc(path, out):
    rd = os.path.join(out, "_recalc")
    os.makedirs(rd, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", rd, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return os.path.join(rd, os.path.basename(path))


def render(path, out, dpi=60):
    pd = os.path.join(out, "_pdf")
    os.makedirs(pd, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", pd, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    pdf = os.path.join(pd, os.path.basename(path)[:-5] + ".pdf")
    png = os.path.join(out, "_png", os.path.basename(path)[:-5])
    os.makedirs(os.path.dirname(png), exist_ok=True)
    for f in os.listdir(os.path.dirname(png)):
        if f.startswith(os.path.basename(png) + "-"):
            os.remove(os.path.join(os.path.dirname(png), f))
    subprocess.run(["pdftoppm", "-r", str(dpi), "-png", pdf, png], check=True)
    return pdf


def _eq(a, b):
    if isinstance(b, (int, float)) and not isinstance(b, bool):
        if not isinstance(a, (int, float)):
            return False
        if isinstance(b, int):
            return abs(a - b) < 1e-6
        return abs(a - b) < 1e-9 * max(1, abs(b))
    if isinstance(b, datetime.date):
        if isinstance(a, datetime.datetime):
            a = a.date()
        return a == b
    return (a if a is not None else "") == b


def verify(m, bk, path, out):
    import openpyxl
    rc = recalc(path, out)
    wf = openpyxl.load_workbook(path)                 # 수식
    wv = openpyxl.load_workbook(rc, data_only=True)   # 다시 계산한 값
    bad, info = [], []
    # ① 예시 값 대조
    multi = getattr(m, "MULTI", False)
    ex_pages = {lg: p for (st, lg), p in bk.pages.items() if st == "ex"}
    if multi:
        inp = {lg: p.inputs for lg, p in ex_pages.items()}
        exp = m.expect(inp)
    else:
        (lg, p), = ex_pages.items()
        exp = {lg: m.expect(p.inputs)}
    n_ok = 0
    for lg, kv in exp.items():
        p = ex_pages[lg]
        ws = wv[p.name]
        for k, want in kv.items():
            if k not in p.k:
                bad.append(f"[{p.name}] expect 이름 없음: {k}")
                continue
            got = ws[p.k[k]].value
            if _eq(got, want):
                n_ok += 1
            else:
                bad.append(f"[{p.name}] {k}({p.k[k]}): 엑셀 {got!r} ≠ 파이썬 {want!r}")
    # 수식 칸인데 expect 에 없는 것(검사 빠짐) 알려 주기
    for lg, p in ex_pages.items():
        miss = [k for k in p.calc_keys if k not in exp.get(lg, {})]
        if miss:
            info.append(f"[{p.name}] expect 에 없는 수식 이름 {len(miss)}: {', '.join(miss[:12])}")
    # ② 오류값·바깥 참조·빈 서식 모양
    for ws in wv.worksheets:
        fws = wf[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and any(v.startswith(e) for e in ERRS):
                    bad.append(f"[{ws.title}] {c.coordinate} 오류값 {v}")
                f = fws[c.coordinate].value
                if isinstance(f, str) and f.startswith("="):
                    if "[" in f and "]" in f.split("!")[0]:
                        bad.append(f"[{ws.title}] {c.coordinate} 바깥 파일 참조 {f[:60]}")
    for (st, lg), p in bk.pages.items():
        if st != "blank":
            continue
        ws = wv[p.name]
        shown = []
        for cell, _t in p.formulas:
            v = ws[cell.coordinate].value
            if v in (None, "", 0):
                continue
            shown.append(f"{cell.coordinate}={str(v)[:40]!r}")
        if shown:
            info.append(f"[{p.name}] 빈 서식인데 값이 보이는 수식 칸 {len(shown)}: " + " · ".join(shown[:8]))
    return n_ok, bad, info


def main():
    args = sys.argv[1:]
    out = "/tmp/g10/f2/out"
    if "--out" in args:
        i = args.index("--out")
        out = args[i + 1]
        del args[i:i + 2]
    do_render = "--render" in args
    args = [a for a in args if not a.startswith("--")]
    if not args:
        args = sorted(f[2:-3].replace("_", "-") for f in os.listdir(HERE) if f.startswith("f_") and f.endswith(".py"))
    total_bad = 0
    for slug in args:
        m, bk, path = build(slug, out)
        n_ok, bad, info = verify(m, bk, path, out)
        mark = "✅" if not bad else "❌"
        print(f"{mark} {m.SLUG:<22} 수식 대조 {n_ok}칸 맞음 · 틀림 {len(bad)} · 시트 {len(bk.wb.worksheets)} · {os.path.getsize(path)/1024:.0f}KB")
        for b in bad[:30]:
            print("    ✗", b)
        for t in info:
            print("    ·", t)
        total_bad += len(bad)
        if do_render:
            render(path, out)
    sys.exit(1 if total_bad else 0)


if __name__ == "__main__":
    main()
