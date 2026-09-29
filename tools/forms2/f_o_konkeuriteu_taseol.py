# -*- coding: utf-8 -*-
"""구조물별 콘크리트 타설현황 — 원본 틀([별지 제35호 서식] 구조물명 · 타설일자 · 타설 부위 · 설계량 · 타설량 · 누계/증감 · 배합 · 납품회사 ·
타설방법 · 타설시간 · 타설시 온도 · 시험결과(슬럼프 · 공기량 · 염분량 · 28일 압축강도) · 시험자(품질관리자 · 감리원))을 새로 만든 것.
■ 원본 파일의 다른 시트(실제 현장의 타설 기록)는 쓰지 않음.
■ 수식: 증감 = 타설량 − 설계량 · 누계 = 같은 구조물의 타설량 누계 · 배합(예: 25-24-150)에서 호칭강도·목표 슬럼프를 읽어
        판정(슬럼프 허용차 · 공기량 4.5±1.5% · 염화물 0.30kg/㎥ 이하 · 28일 강도 1회 ≥ 호칭강도의 85%, 35MPa 넘으면 90%) ·
        합계 · 배합별 합계 · 부적합 건수. 기준값은 노란 칸(현장·발주처 기준에 맞게 고쳐 쓰기)."""
import datetime

from fb import F, I

SLUG = "o-konkeuriteu-taseol"
TITLE = "구조물별 콘크리트 타설현황"
WHERE = "orig"
PREV = [(1, "빈 서식"), (2, "작성 예시")]
PAGES = 1
NSHEETS = 1
SHORT = ("구조물별 콘크리트 타설현황([별지 제35호 서식]). 설계량·타설량을 적으면 증감·구조물별 누계·배합별 합계가 나오고, "
         "슬럼프·공기량·염화물·28일 강도를 적으면 배합(예: 25-24-150)에 맞춰 적합/부적합을 판정합니다.")
NOTE = "노란 칸(공기량 4.5±1.5% · 염화물 0.30kg/㎥ · 강도 85%)은 KS F 4009 기준값입니다 — 발주처 기준이 다르면 고쳐 쓰세요."

D = datetime.date
ROWS = 25
W = [9, 9.5, 20, 6.8, 6.8, 6.2, 6.8, 9.5, 10, 7, 9.5, 6, 6.6, 6.2, 7, 7.8, 11, 7, 6, 7, 6]
N = len(W)
H = {"kind": "head", "size": 9}
EX = [("암거 기초", D(2026, 4, 8), "STA.0+120 ~ 0+160", 42, 44, "25-18-80", "펌프카", "08:30~10:10", 14, 90, 4.2, 0.031, 21.6),
      ("암거 기초", D(2026, 4, 15), "STA.0+160 ~ 0+200", 42, 43, "25-18-80", "펌프카", "08:20~09:50", 16, 85, 4.6, 0.028, 20.3),
      ("집수정", D(2026, 4, 22), "1~4호 바닥", 12, 12, "25-24-150", "백호", "09:00~10:00", 17, 160, 4.8, 0.042, 26.1),
      ("집수정", D(2026, 4, 29), "1~4호 벽체", 18, 19, "25-24-150", "펌프카", "08:40~10:20", 18, 190, 5.1, 0.036, 25.4),
      ("옹벽", D(2026, 5, 6), "기초 L=40m", 36, 38, "25-24-150", "펌프카", "08:30~10:30", 19, 145, 4.4, 0.025, 22.8),
      ("옹벽", D(2026, 5, 13), "벽체 L=40m", 30, 31, "25-24-150", "펌프카", "08:10~10:00", 20, 155, 4.0, 0.033, 25.9),
      ("암거 기초", D(2026, 5, 20), "STA.0+200 ~ 0+240", 42, 45, "25-18-80", "펌프카", "08:30~10:20", 21, 95, 4.3, 0.029, None),
      ("보도 경계석 기초", D(2026, 5, 27), "좌측 L=120m", 14, 13, "25-18-80", "백호", "13:00~14:30", 23, 80, 3.9, 0.038, None)]
QM, SM = "김품질", "박감리"
CR = 'SUBSTITUTE({b}," ","")'


def draw(bk, ex):
    p = bk.page("타설현황", W, ex=ex, fit_height=1, landscape=True, sheetname="타설현황" if not ex else "예시-타설현황", margins=(0.35, 0.35, 0.45, 0.45))
    p.row([(N - 4, "구조물별 콘크리트 타설 현황", {"kind": "title", "size": 20}), (4, "[별지 제35호 서식]", {"kind": "free", "size": 9, "align": "right"})], h=40)
    p.row([(2, "현 장 명 :", {"kind": "free", "bold": True, "align": "right"}), (7, I("현장명", ex="가나지구 배수로 정비공사", size=11), {"border": False}),
           (N - 9, "", {"kind": "free"})], h=26)
    p.gap(4)
    p.row([(1, "구 조\n물 명", dict(H, rs=3)), (1, "타설 일자", dict(H, rs=3)), (1, "타 설  부 위", dict(H, rs=3)), (1, "설계량\n(㎥)", dict(H, rs=3)),
           (1, "타설량\n(㎥)", dict(H, rs=3)), (1, "증감\n(㎥)", dict(H, rs=3)), (1, "구조물\n누계\n(㎥)", dict(H, rs=3)), (1, "콘크리트\n배합종류", dict(H, rs=3)),
           (1, "납품\n회사", dict(H, rs=3)), (1, "타 설\n방 법", dict(H, rs=3)), (1, "타 설\n시 간", dict(H, rs=3)), (1, "타설시\n온도\n(℃)", dict(H, rs=3)),
           (5, "시   험   결   과", H), (4, "시    험    자", H)], h=22)
    p.row([(1, "슬럼프\n(mm)", dict(H, rs=2)), (1, "공기량\n(%)", dict(H, rs=2)), (1, "염분량\n(kg/㎥)", dict(H, rs=2)), (1, "28일\n압축강도\n(MPa)", dict(H, rs=2)),
           (1, "판  정", dict(H, rs=2)), (2, "품질관리자", H), (2, "감 리 원", H)], h=22, start=13)
    p.row([(1, "성명", H), (1, "서명", H), (1, "성명", H), (1, "서명", H)], h=18, start=18)
    first = p.r
    # 기준값 칸 위치는 아래에서 정해지므로 이름으로 부름
    for i in range(1, ROWS + 1):
        e = EX[i - 1] if (ex and i <= len(EX)) else (None,) * 13
        b = f"{{배합#{i}}}"
        cb = CR.format(b=b)
        # 배합 «굵은골재-호칭강도-슬럼프» 에서 호칭강도 · 슬럼프 읽기(못 읽으면 빈칸)
        fck = f'IFERROR(VALUE(MID({cb},FIND("-",{cb})+1,FIND("-",{cb},FIND("-",{cb})+1)-FIND("-",{cb})-1)),"")'
        slp = f'IFERROR(VALUE(MID({cb},FIND("-",{cb},FIND("-",{cb})+1)+1,9)),"")'
        cells = [(1, I(f"구조물#{i}", ex=e[0], align="center", size=9)),
                 (1, I(f"일자#{i}", ex=e[1], fmt="ymd", align="center", size=9)),
                 (1, I(f"부위#{i}", ex=e[2], size=9)),
                 (1, I(f"설계#{i}", ex=e[3], fmt="num", size=9)),
                 (1, I(f"타설#{i}", ex=e[4], fmt="num", size=9)),
                 (1, F(f"증감#{i}", f'=IF(OR({{설계#{i}}}="",{{타설#{i}}}=""),"",ROUND({{타설#{i}}}-{{설계#{i}}},3))', fmt='+#,##0.###;[Red]-#,##0.###;0', size=9)),
                 (1, F(f"누계#{i}", f'=IF(OR({{구조물#{i}}}="",{{타설#{i}}}=""),"",SUMIF({{구조물#1}}:{{구조물#{i}}},{{구조물#{i}}},{{타설#1}}:{{타설#{i}}}))',
                       fmt="num", size=9)),
                 (1, I(f"배합#{i}", ex=e[5], align="center", size=9)),
                 (1, I(f"회사#{i}", ex="예시레미콘(주)" if (ex and i <= len(EX)) else None, align="center", size=8.5)),
                 (1, I(f"방법#{i}", ex=e[6], align="center", size=9)),
                 (1, I(f"시간#{i}", ex=e[7], align="center", size=8.5)),
                 (1, I(f"온도#{i}", ex=e[8], fmt='0.#', align="center", size=9)),
                 (1, I(f"슬럼프#{i}", ex=e[9], fmt='0', align="center", size=9)),
                 (1, I(f"공기#{i}", ex=e[10], fmt='0.0', align="center", size=9)),
                 (1, I(f"염분#{i}", ex=e[11], fmt='0.000', align="center", size=9)),
                 (1, I(f"강도#{i}", ex=e[12], fmt='0.0', align="center", size=9))]
        s_bad = (f'AND(ISNUMBER({{슬럼프#{i}}}),ISNUMBER({{목표#{i}}}),ABS({{슬럼프#{i}}}-N({{목표#{i}}}))>IF(N({{목표#{i}}})>=80,25,IF(N({{목표#{i}}})>=50,15,10)))')
        a_bad = f'AND(ISNUMBER({{공기#{i}}}),ABS({{공기#{i}}}-{{공기기준}})>{{공기허용}})'
        c_bad = f'AND(ISNUMBER({{염분#{i}}}),{{염분#{i}}}>{{염분한도}})'
        f_bad = (f'AND(ISNUMBER({{강도#{i}}}),ISNUMBER({{호칭#{i}}}),{{강도#{i}}}<N({{호칭#{i}}})*IF(N({{호칭#{i}}})>35,{{강도비35}},{{강도비}}))')
        f_low = f'AND(ISNUMBER({{강도#{i}}}),ISNUMBER({{호칭#{i}}}),{{강도#{i}}}<N({{호칭#{i}}}))'
        pan = (f'=IF(COUNT({{슬럼프#{i}}},{{공기#{i}}},{{염분#{i}}},{{강도#{i}}})=0,"",IF(OR({s_bad},{a_bad},{c_bad},{f_bad}),'
               f'"부적합:"&IF({s_bad}," 슬럼프","")&IF({a_bad}," 공기량","")&IF({c_bad}," 염분","")&IF({f_bad}," 강도",""),'
               f'IF({f_low},"적합(강도 3회 평균 확인)",IF(ISNUMBER({{강도#{i}}}),"적합","적합(28일 강도 대기)"))))')
        cells += [(1, F(f"판정#{i}", pan, align="center", size=7.5)),
                  (1, I(f"품질#{i}", ex=QM if (ex and i <= len(EX)) else None, align="center", size=9)), (1, "", {"kind": "text"}),
                  (1, I(f"감리#{i}", ex=SM if (ex and i <= len(EX)) else None, align="center", size=9)), (1, "", {"kind": "text"})]
        p.row(cells, h=24)
    last = p.r - 1
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import Alignment, Font
    from openpyxl.styles.differential import DifferentialStyle
    p.ws.conditional_formatting.add(f"Q{first}:Q{last}", Rule(type="expression", formula=[f'LEFT(Q{first},3)="부적합"'],
                                    dxf=DifferentialStyle(font=Font(color="DC2626", bold=True))))
    for r in range(first, last + 1):          # 납품회사 · 타설시간은 두 줄로 쪼개지 말고 칸에 맞춰 줄임
        for c in (9, 11):
            p.ws.cell(row=r, column=c).alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    # 합계
    p.row([(3, "합        계", {"kind": "sum"}),
           (1, F("설계계", f'=IF(COUNT(D{first}:D{last})=0,"",SUM(D{first}:D{last}))', fmt="num", size=9, bold=True)),
           (1, F("타설계", f'=IF(COUNT(E{first}:E{last})=0,"",SUM(E{first}:E{last}))', fmt="num", size=9, bold=True)),
           (1, F("증감계", f'=IF(OR({{설계계}}="",{{타설계}}=""),"",ROUND({{타설계}}-{{설계계}},3))', fmt='+#,##0.###;[Red]-#,##0.###;0', size=9, bold=True)),
           (6, "", {"kind": "sum"}),
           (4, F("시험수", f'=IF(COUNTIF(Q{first}:Q{last},"적합*")+COUNTIF(Q{first}:Q{last},"부적합*")=0,"","시험 "&(COUNTIF(Q{first}:Q{last},"적합*")+COUNTIF(Q{first}:Q{last},"부적합*"))&"건 · 부적합 "'
                          f'&COUNTIF(Q{first}:Q{last},"부적합*")&"건")', align="center", size=9, bold=True)),
           (1, "", {"kind": "sum"}), (4, "", {"kind": "sum"})], h=24)
    p.gap(6)
    # 배합별 합계 + 기준값
    p.row([(3, "배합별 합계", {"kind": "label"}), (1, "설계량", H), (1, "타설량", H), (1, "증감", H), (1, "", {"kind": "free"}),
           (5, "판정 기준값 (KS F 4009 — 발주처 기준이 다르면 고치세요)", {"kind": "label", "size": 9}), (2, "", {"kind": "free"}), (7, "", {"kind": "free"})], h=20)
    exm = ["25-18-80", "25-24-150", None, None]
    crit = [("공기량 기준(%)", "공기기준", 4.5, "0.0"), ("공기량 허용차(±%)", "공기허용", 1.5, "0.0"), ("염화물 한도(kg/㎥)", "염분한도", 0.30, "0.00"),
            ("28일 강도 1회 ≥ 호칭강도 ×", "강도비", 0.85, "0.00")]
    for j in range(1, 5):
        cells = [(3, I(f"배합종류#{j}", ex=exm[j - 1] if ex else None, align="center", size=9)),
                 (1, F(f"배합설계#{j}", f'=IF({{배합종류#{j}}}="","",SUMIF(H{first}:H{last},{{배합종류#{j}}},D{first}:D{last}))', fmt="num", size=9)),
                 (1, F(f"배합타설#{j}", f'=IF({{배합종류#{j}}}="","",SUMIF(H{first}:H{last},{{배합종류#{j}}},E{first}:E{last}))', fmt="num", size=9)),
                 (1, F(f"배합증감#{j}", f'=IF({{배합종류#{j}}}="","",ROUND({{배합타설#{j}}}-{{배합설계#{j}}},3))', fmt='+#,##0.###;[Red]-#,##0.###;0', size=9)),
                 (1, "", {"kind": "free"})]
        lab, key, v, fm = crit[j - 1]
        cells += [(4, lab, {"kind": "text", "size": 9}), (1, I(key, blank=v, rate=True, fmt=fm, align="center", size=9)), (2, "", {"kind": "free"}), (7, "", {"kind": "free"})]
        p.row(cells, h=20)
    p.row([(7, "", {"kind": "free"}), (4, "호칭강도 35MPa 넘으면 ×", {"kind": "text", "size": 9}),
           (1, I("강도비35", blank=0.90, rate=True, fmt="0.00", align="center", size=9)), (9, "", {"kind": "free"})], h=20)
    p.end_print()
    # 인쇄 밖 도움 칸: 배합에서 읽은 호칭강도 · 목표 슬럼프
    hc = N + 2
    p.ws.column_dimensions[_cl(hc)].width = 8
    p.ws.column_dimensions[_cl(hc + 1)].width = 8
    p.put(first - 1, hc, hc, "호칭강도", kind="head", size=8)
    p.put(first - 1, hc + 1, hc + 1, "목표슬럼프", kind="head", size=8)
    for i in range(1, ROWS + 1):
        r = first + i - 1
        b = f"{{배합#{i}}}"
        cb = CR.format(b=b)
        fck = f'=IF({b}="","",IFERROR(VALUE(MID({cb},FIND("-",{cb})+1,FIND("-",{cb},FIND("-",{cb})+1)-FIND("-",{cb})-1)),""))'
        slp = f'=IF({b}="","",IFERROR(VALUE(MID({cb},FIND("-",{cb},FIND("-",{cb})+1)+1,9)),""))'
        p.put(r, hc, hc, F(f"호칭#{i}", fck, align="center", size=8))
        p.put(r, hc + 1, hc + 1, F(f"목표#{i}", slp, align="center", size=8))
    p.after_note(["노란 칸 = 판정 기준값(KS F 4009 레디믹스트 콘크리트). 하늘색 칸 = 자동(증감 · 구조물 누계 · 판정 · 합계 · 배합별 합계).",
                  "판정: 슬럼프 허용차(목표 80mm 이상 ±25 · 50~65 ±15 · 25 ±10) · 공기량 기준±허용차 · 염화물 한도 이하 · 28일 강도 1회 값 ≥ 호칭강도×0.85(35MPa 넘으면 0.90). "
                  "강도가 호칭강도보다 낮으면 «3회 평균 확인»(3회 평균은 호칭강도 이상이어야 함).",
                  "배합은 «굵은골재-호칭강도-슬럼프» 꼴(예: 25-24-150)로 적으면 오른쪽 인쇄 밖 칸에서 호칭강도·목표 슬럼프를 읽습니다."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 공사 · 회사 · 사람)."])
    return p


def _cl(c):
    from openpyxl.utils import get_column_letter
    return get_column_letter(c)


def expect(inp):
    e = {}
    cum = {}
    tests = bad_n = 0
    qa, qt = inp["공기기준"], inp["공기허용"]
    for i in range(1, ROWS + 1):
        s, d, t = inp.get(f"구조물#{i}"), inp.get(f"설계#{i}"), inp.get(f"타설#{i}")
        e[f"증감#{i}"] = "" if (d is None or t is None) else round(t - d, 3)
        if s and t is not None:
            cum[s] = cum.get(s, 0) + t
            e[f"누계#{i}"] = cum[s]
        else:
            e[f"누계#{i}"] = ""
        b = inp.get(f"배합#{i}")
        if b:
            parts = b.replace(" ", "").split("-")
            fck, slp = float(parts[1]), float(parts[2])
            e[f"호칭#{i}"] = fck
            e[f"목표#{i}"] = slp
        else:
            fck = slp = None
            e[f"호칭#{i}"] = ""
            e[f"목표#{i}"] = ""
        sl, ai, ch, st = (inp.get(f"슬럼프#{i}"), inp.get(f"공기#{i}"), inp.get(f"염분#{i}"), inp.get(f"강도#{i}"))
        if all(v is None for v in (sl, ai, ch, st)):
            e[f"판정#{i}"] = ""
            continue
        tests += 1
        tol = 25 if (slp or 0) >= 80 else (15 if (slp or 0) >= 50 else 10)
        sb = sl is not None and slp is not None and abs(sl - slp) > tol
        ab = ai is not None and abs(ai - qa) > qt
        cb = ch is not None and ch > inp["염분한도"]
        fb_ = st is not None and fck is not None and st < fck * (inp["강도비35"] if fck > 35 else inp["강도비"])
        fl = st is not None and fck is not None and st < fck
        if sb or ab or cb or fb_:
            bad_n += 1
            e[f"판정#{i}"] = "부적합:" + (" 슬럼프" if sb else "") + (" 공기량" if ab else "") + (" 염분" if cb else "") + (" 강도" if fb_ else "")
        elif fl:
            e[f"판정#{i}"] = "적합(강도 3회 평균 확인)"
        elif st is not None:
            e[f"판정#{i}"] = "적합"
        else:
            e[f"판정#{i}"] = "적합(28일 강도 대기)"
    ds = [inp.get(f"설계#{i}") for i in range(1, ROWS + 1) if inp.get(f"설계#{i}") is not None]
    ts = [inp.get(f"타설#{i}") for i in range(1, ROWS + 1) if inp.get(f"타설#{i}") is not None]
    e["설계계"] = sum(ds) if ds else ""
    e["타설계"] = sum(ts) if ts else ""
    e["증감계"] = round(sum(ts) - sum(ds), 3) if (ds and ts) else ""
    e["시험수"] = f"시험 {tests}건 · 부적합 {bad_n}건" if tests else ""
    for j in range(1, 5):
        k = inp.get(f"배합종류#{j}")
        if not k:
            e[f"배합설계#{j}"] = e[f"배합타설#{j}"] = e[f"배합증감#{j}"] = ""
            continue
        sd = sum(inp.get(f"설계#{i}") or 0 for i in range(1, ROWS + 1) if inp.get(f"배합#{i}") == k)
        st_ = sum(inp.get(f"타설#{i}") or 0 for i in range(1, ROWS + 1) if inp.get(f"배합#{i}") == k)
        e[f"배합설계#{j}"] = sd
        e[f"배합타설#{j}"] = st_
        e[f"배합증감#{j}"] = round(st_ - sd, 3)
    return e
