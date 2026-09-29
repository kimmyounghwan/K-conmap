# -*- coding: utf-8 -*-
"""공사예정공정표(총괄) — 현장에서 쓰던 서식의 틀(번호·공종 · 수량 · 단위 · 날짜 눈금 · 비고, 가로 한 장) 그대로 새로 만든 것.

■ 착공일만 적으면 눈금(월 · 상/하순)이 저절로 채워지고, 공종마다 오른쪽 «막대 칸»(착수일·완료일)을 적으면 막대가 그려집니다.
■ 금액(가중치)을 적으면 맨 아래에 월별 계획 공정률(월간·누계, %)이 나옵니다 — 금액 × (그 달까지 지난 날 ÷ 공종 기간)을 더한 값.
■ 오른쪽 착수일·완료일·금액 칸은 화면에서만 보입니다(인쇄 안 됨) — 인쇄되는 칸은 원본과 같습니다.
원본의 공종 목록(다른 현장 내용)은 지우고, 가상의 배수로 공사로 예시를 새로 적었습니다."""
import datetime

from fb import F, I
from openpyxl.formatting.rule import Rule
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as CL

SLUG = "gongjeongpyo"
TITLE = "공사예정공정표"
WHERE = "forms"
PREV = None
SHORT = "공사공정표(총괄). 착공일을 적으면 월·상하순 눈금이, 공종마다 착수·완료일을 적으면 막대가, 금액을 적으면 월별 계획 공정률(%)이 저절로 나옵니다."
NOTE_REPLACE = {
    "🏗 **현장에서 실제로 쓰던 서식입니다** — 내용만 지우고 틀은 원본 그대로 두었습니다.":
        "🏗 **현장에서 실제로 쓰던 서식의 틀**(공종·수량·단위·날짜 눈금·비고)을 그대로 두고 새로 만들었습니다 — 착공일·착수/완료일·금액만 적으면 눈금·막대·월별 계획 공정률이 자동입니다.",
}

D = datetime.date
MONTHS = 18
NT = MONTHS * 2
N = 32
C0 = 5                       # 눈금 첫 칸(E)
CB = C0 + NT                 # 비고
CS, CE, CW = CB + 1, CB + 2, CB + 3          # 착수 · 완료 · 금액 (화면에서만)
W = [4.5, 22, 7.5, 5] + [2.3] * NT + [9] + [10.5, 10.5, 13]
PC = CB
EX = [  # 번호, 공종, 수량, 단위, 착수, 완료, 금액
    ("1.", "가설공사", None, None, None, None, None),
    (None, "공통가설·안전시설", 1, "식", D(2026, 3, 9), D(2026, 12, 31), 18_000_000),
    ("2.", "토공", None, None, None, None, None),
    (None, "터파기", 3_450, "㎥", D(2026, 3, 16), D(2026, 5, 31), 42_000_000),
    (None, "되메우기", 2_980, "㎥", D(2026, 4, 15), D(2026, 7, 15), 21_000_000),
    ("3.", "배수공", None, None, None, None, None),
    (None, "PE 이중벽관 D600 부설", 862, "m", D(2026, 4, 1), D(2026, 6, 30), 185_000_000),
    (None, "집수정 설치", 3, "개소", D(2026, 5, 1), D(2026, 6, 15), 12_000_000),
    ("4.", "구조물공", None, None, None, None, None),
    (None, "암거 기초", 1, "식", D(2026, 6, 1), D(2026, 6, 30), 36_000_000),
    (None, "암거 본체(철근·콘크리트)", 1, "식", D(2026, 7, 1), D(2026, 9, 15), 168_000_000),
    (None, "암거 뒤채움", 1, "식", D(2026, 9, 16), D(2026, 10, 10), 14_000_000),
    ("5.", "포장공", None, None, None, None, None),
    (None, "아스팔트 재포장", 6_020, "㎡", D(2026, 10, 1), D(2026, 11, 15), 96_000_000),
    ("6.", "부대공", None, None, None, None, None),
    (None, "교통처리·안전관리", 1, "식", D(2026, 3, 9), D(2026, 12, 31), 22_000_000),
    (None, "준공 청소·정리", 1, "식", D(2026, 12, 1), D(2026, 12, 31), 5_000_000),
]
BAR = PatternFill("solid", fgColor="2F5597", bgColor="2F5597")


def draw(bk, ex):
    p = bk.page("공사공정표(총괄)" if not ex else "작성 예시", W, ex=ex, landscape=True, fit_height=1, print_cols=PC,
                margins=(0.3, 0.3, 0.4, 0.4))
    L = {"kind": "label"}
    H = {"kind": "head"}
    p.title("공 사 공 정 표  (총 괄)", h=34)
    p.row([(2, "공  사  명", L), (13, I("공사명", ex="가나지구 배수로 정비공사")), (4, "착 공 일", L),
           (8, I("착공일", ex=D(2026, 3, 9), fmt="date", align="center")), (4, "준 공 일", L),
           (8, I("준공일", ex=D(2026, 12, 31), fmt="date", align="center")),
           (2, F("공기", '=IF(OR(N({착공일})=0,N({준공일})=0),"",{준공일}-{착공일}+1)', fmt='#,##0"일";;""', align="center"))],
          h=24)
    p.gap(4)
    # 머리: 공종 · 수량 · 단위 · (월 눈금) · 비고
    hr = p.r
    cells = [(2, "공        종", {"kind": "head", "rs": 2}), (1, "수 량", {"kind": "head", "rs": 2}), (1, "단위", {"kind": "head", "rs": 2})]
    for m in range(MONTHS):
        cells.append((2, F(f"월#{m + 1}", f'=IF(N({{착공일}})=0,"",EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m}))',
                           fmt='yy"."m;;""', align="center", size=8), {"kind": "head"}))
    cells.append((1, "비  고", {"kind": "head", "rs": 2}))
    cells += [(1, "착수일", {"kind": "head", "rs": 2, "fill": None}), (1, "완료일", {"kind": "head", "rs": 2}),
              (1, "금액(가중치)", {"kind": "head", "rs": 2, "size": 9})]
    p.row(cells, h=20)
    p.row([(1, "상" if k % 2 == 0 else "하", {"kind": "head", "size": 8}) for k in range(NT)], h=16, start=C0)
    # 숨은 줄 두 개: 칸마다 시작일·끝일(막대·공정률 계산용)
    rs_, re_ = p.r, p.r + 1
    for rr, kind in ((rs_, "s"), (re_, "e")):
        for k in range(NT):
            m, half = k // 2, k % 2
            base = f"EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m})"
            if kind == "s":
                f = f"=IF(N({{착공일}})=0,\"\",{base}{'+15' if half else ''})"
            else:
                f = (f"=IF(N({{착공일}})=0,\"\",{base}+14)" if half == 0 else
                     f"=IF(N({{착공일}})=0,\"\",EDATE(DATE(YEAR({{착공일}}),MONTH({{착공일}}),1),{m + 1})-1)")
            p.put(rr, C0 + k, C0 + k, F(f"{kind}{k + 1}", f, fmt="yyyy-mm-dd"), kind="free", size=7)
        p.ws.row_dimensions[rr].height = 12
        p.ws.row_dimensions[rr].hidden = True
    p.r = re_ + 1
    r0 = p.r
    for i in range(1, N + 1):
        e = EX[i - 1] if i <= len(EX) else (None,) * 7
        no, name, q, u, s, t, w = e
        cells = [(1, I(f"번호#{i}", ex=no, align="center", size=9)), (1, I(f"공종#{i}", ex=name, size=9.5, bold=bool(no))),
                 (1, I(f"수량#{i}", ex=q, fmt='#,##0.##;;""', size=9)), (1, I(f"단위#{i}", ex=u, align="center", size=9))]
        cells += [(1, "", {"kind": "text"}) for _ in range(NT)]
        cells += [(1, I(f"비고#{i}", size=8.5)), (1, I(f"착수#{i}", ex=s, fmt="ymd", size=9)), (1, I(f"완료#{i}", ex=t, fmt="ymd", size=9)),
                  (1, I(f"금액#{i}", ex=w, fmt='#,##0;;""', size=9))]
        p.row(cells, h=17)
    r1 = p.r - 1
    # 막대: 착수 ≤ 그 칸 끝 · 완료 ≥ 그 칸 시작
    tl = f"{CL(C0)}{r0}"
    sC, eC = CL(CS), CL(CE)
    rule = Rule(type="expression", dxf=DifferentialStyle(fill=BAR),
                formula=[f"AND(N(${sC}{r0})>0,N(${eC}{r0})>=N(${sC}{r0}),${sC}{r0}<={CL(C0)}${re_},${eC}{r0}>={CL(C0)}${rs_})"])
    p.ws.conditional_formatting.add(f"{tl}:{CL(C0 + NT - 1)}{r1}", rule)
    for k, c in (("착수", CS), ("완료", CE), ("금액", CW)):
        p.k[f"{k}[]"] = f"${CL(c)}${r0}:${CL(c)}${r1}"
    # 월별 계획 공정률: 금액 × (그 달 말까지 지난 날 ÷ 기간) 의 합 ÷ 금액 합
    S = {"kind": "sum"}
    acc = []
    for m in range(MONTHS):
        me = f"{{e{2 * m + 2}}}"
        s, t, w = "{착수[]}", "{완료[]}", "{금액[]}"
        dur = f"(({t}-{s}+1)*(({t}-{s}+1)>0)+(({t}-{s}+1)<=0))"
        upto = f"((({t}<={me})*{t}+({t}>{me})*{me})-{s}+1)"
        ms = f"{{s{2 * m + 1}}}"
        f = (f'=IF(OR(N({{착공일}})=0,SUM({w})=0),"",IF(AND(N({{준공일}})>0,{ms}>{{준공일}}),"",SUMPRODUCT({w}*({s}>0)*({t}>={s})*({upto}>0)*{upto}/{dur})/SUM({w})))')
        acc.append(f)
    p.row([(4, "누계 계획 공정률", S)] + [(2, F(f"누계#{m + 1}", acc[m], fmt='0.0%;;""', align="center", size=8)) for m in range(MONTHS)]
          + [(1, "", S)], h=20)
    p.row([(4, "월간 계획 공정률", S)] +
          [(2, F(f"월간#{m + 1}", (f'=IF({{누계#{m + 1}}}="","",{{누계#{m + 1}}}-' + (f'N({{누계#{m}}})' if m else "0") + ")"),
                 fmt='0.0%;;""', align="center", size=8)) for m in range(MONTHS)] + [(1, "", S)], h=20)
    p.end_print()
    p.after_note([
        "① 착공일을 적으면 월·상하순 눈금이 채워집니다(18개월). ② 공종마다 오른쪽 착수일·완료일을 적으면 막대가 그려집니다.",
        "③ 금액(가중치)을 적으면 맨 아래 누계·월간 계획 공정률이 나옵니다. 오른쪽 세 칸(착수·완료·금액)은 인쇄되지 않습니다.",
        "공종 줄이 모자라면 중간에 줄을 끼워 넣으세요(막대·공정률 범위가 따라 늘어납니다).",
    ] if not ex else ["이 시트는 작성 예시입니다(가상의 공사 · 금액 · 날짜). 실제로는 앞 시트에 적으세요."])
    return p


def _edate(d, m):
    y, mo = divmod(d.month - 1 + m, 12)
    return D(d.year + y, mo + 1, 1)


def expect(inp):
    st = inp["착공일"]
    base = D(st.year, st.month, 1)
    e = {"공기": (inp["준공일"] - st).days + 1}
    for m in range(MONTHS):
        mb = _edate(base, m)
        e[f"월#{m + 1}"] = mb
        e[f"s{2 * m + 1}"] = mb
        e[f"s{2 * m + 2}"] = mb + datetime.timedelta(days=15)
        e[f"e{2 * m + 1}"] = mb + datetime.timedelta(days=14)
        e[f"e{2 * m + 2}"] = _edate(base, m + 1) - datetime.timedelta(days=1)
    rows = [(inp.get(f"착수#{i}"), inp.get(f"완료#{i}"), inp.get(f"금액#{i}") or 0) for i in range(1, N + 1)]
    tw = sum(w for _, _, w in rows)
    prev = 0.0
    for m in range(MONTHS):
        me = e[f"e{2 * m + 2}"]
        if e[f"s{2 * m + 1}"] > inp["준공일"]:
            e[f"누계#{m + 1}"] = ""
            e[f"월간#{m + 1}"] = ""
            continue
        v = 0.0
        for s, t, w in rows:
            if not (s and t and w) or t < s:
                continue
            upto = (min(t, me) - s).days + 1
            if upto > 0:
                v += w * upto / ((t - s).days + 1)
        cum = v / tw
        e[f"누계#{m + 1}"] = cum
        e[f"월간#{m + 1}"] = cum - prev
        prev = cum
    return e
