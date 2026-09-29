# -*- coding: utf-8 -*-
"""
fb.py — 서식 «다시 만들기» 틀 (2026-09-29)

소장님 말씀: 「서식은 틀을 유지하라는 거지 그대로 올리라는 것은 아니였어.
             수식넣고, 틀만 유지하고 다시 새롭게 다 만들어 줘. 틀리면 안돼.」

■ 한 서식 = «빈 서식» 시트(들) + «작성 예시» 시트(들)
  같은 그리기 함수를 두 번 부릅니다(ex=False → 빈 서식, ex=True → 작성 예시).
  그래서 두 시트의 칸 자리·수식이 «한 글자도» 다르지 않습니다.
  예시 값은 I(…, ex=값) 로 입력칸 옆에 같이 적습니다 — 빈 서식에는 안 들어갑니다.

■ 칸 종류
  글(라벨·제목) · I(입력칸) · F(수식칸, 옅은 하늘색 칸) · 요율칸(I(…, rate=True), 노란 바탕)
  수식은 '={직접노무비}*{간노율}' 처럼 칸 «이름»으로 적고, 마지막에 주소로 바꿉니다.
  다른 시트 칸은 '{신청서!계약금액}' — 예시 벌에서는 예시 시트로 알아서 이어집니다.

■ 검증(verify.py)
  LibreOffice 로 다시 계산 → 예시 시트의 수식 결과를 «파이썬으로 따로 계산한 값»과 원 단위까지 대조,
  빈 서식은 수식 칸이 모두 빈칸(또는 0 숨김)이고 오류(#REF! 등)가 없어야 통과.
"""

import math
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "맑은 고딕"
INK = "1F2937"
CALC_INK = INK                 # 자동 계산 글자도 검정(인쇄용). 대신 칸을 아주 옅은 하늘색으로
CALC = PatternFill("solid", fgColor="EEF4FC")
LINE = Side(style="thin", color="5B6573")
HEAD = PatternFill("solid", fgColor="E8EEF7")
LABEL = PatternFill("solid", fgColor="F3F5F9")
RATE = PatternFill("solid", fgColor="FFF1B8")    # 노란 칸 — 요율·기준값
SUMF = PatternFill("solid", fgColor="F6F8FB")

NF = {
    "won": '#,##0;[Red]-#,##0;""',          # 0 은 빈칸처럼 보이게(빈 서식 인쇄가 깨끗)
    "won0": "#,##0",
    "num": '#,##0.###;[Red]-#,##0.###;""',
    "qty": '#,##0.00;[Red]-#,##0.00;""',
    "int": '#,##0;[Red]-#,##0;""',
    "pct": '0.00%;[Red]-0.00%;""',
    "pct1": '0.0%;[Red]-0.0%;""',
    "rate": "0.0##%",
    "date": 'yyyy"년" m"월" d"일"',
    "ymd": "yyyy. m. d.",
    "md": 'm"/"d',
    "hangul": '[DBNum4][$-412]"일금 "General"원정";;""',   # 1,234 → 일금 일천이백삼십사원정
    "wonsign": '"(₩"#,##0")";;""',
    "day": '#,##0"일";;""',
    "text": "@",
}


def TR(x):
    """원 미만 절사. 부동소수 찌꺼기(7.9999999…)로 1원 모자라는 일을 막으려 소수 6자리에서 한 번 반올림."""
    return f"INT(ROUND({x},6))"


class I:
    """입력칸. ex=예시 값(작성 예시 시트에만 들어감). blank=빈 서식에도 미리 넣어 두는 값(법정 요율 등)."""

    def __init__(self, key=None, ex=None, fmt=None, rate=False, align=None, blank=None, dv=None,
                 size=None, bold=False, wrap=True):
        self.key, self.ex, self.fmt, self.rate = key, ex, fmt, rate
        self.align, self.blank, self.dv, self.size, self.bold, self.wrap = align, blank, dv, size, bold, wrap


class F:
    """수식칸. f 는 '={이름}*2' 꼴. key 를 주면 다른 수식·검증에서 부를 수 있습니다."""

    def __init__(self, key, f, fmt=None, align=None, size=None, bold=False, ink=True, wrap=True):
        self.key, self.f, self.fmt, self.align = key, f, fmt, align
        self.size, self.bold, self.ink, self.wrap = size, bold, ink, wrap


def _w(ch):
    if ch == "\n":
        return 0
    return 1.9 if unicodedata.east_asian_width(ch) in "WF" else 1.05


def text_lines(text, width_units):
    """글자가 몇 줄로 접히는지 어림합니다(openpyxl 은 자동 행높이가 없음)."""
    if text is None:
        return 1
    n = 0
    for part in str(text).split("\n"):
        w = sum(_w(c) for c in part)
        n += max(1, math.ceil(w / max(1.0, width_units - 1.2)))
    return n


class Page:
    def __init__(self, book, name, widths, *, landscape=False, ex=False, logical=None, fit_height=0,
                 brand=True, margins=(0.45, 0.35, 0.5, 0.5), print_cols=None):
        self.book, self.ex = book, ex
        self.logical = logical or name
        self.ws = book.wb.create_sheet(name)
        self.name = name
        self.widths = widths
        self.ncol = len(widths)
        self.pc = print_cols or self.ncol       # 인쇄되는 열 수(그 오른쪽 열은 화면에서만 보이는 도움 칸)
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[CL(i)].width = w
        ws = self.ws
        ws.page_setup.paperSize = 9
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = fit_height
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        l, r, t, b = margins
        ws.page_margins.left, ws.page_margins.right = l, r
        ws.page_margins.top, ws.page_margins.bottom = t, b
        ws.page_margins.header = ws.page_margins.footer = 0.2
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100
        if ex:
            ws.sheet_properties.tabColor = "F59E0B"
        self.k = {}             # 이름 → 'B5'
        self.inputs = {}        # 이름 → 예시 값
        self.formulas = []      # (셀, 틀)
        self.calc_keys = {}     # 이름 → 틀
        self.occ = set()        # 세로 병합으로 차 있는 (행, 열)
        self.dvs = {}
        self.intfmt = {}        # 소수 자리 «있으면 보이는» 서식(#,##0.##) 칸 — 정수일 때 끝에 점(862.)이 안 찍히게
        self.r = 1
        self.last_print = None
        if brand:
            self._brand()

    # ── 1행: 우리 표시(지우기 쉬운 «한 줄») ───────────────────────────────
    def _brand(self):
        tot = sum(self.widths[:self.pc])
        acc, half = 0, 1
        for i, w in enumerate(self.widths[:self.pc], 1):
            acc += w
            half = i
            if acc >= tot * 0.5:
                break
        half = min(half, self.pc - 1) if self.pc > 1 else 1
        if self.ex:
            self.put(1, 1, half, "  K-건설맵  |  k-conmap.com  무료 건설 서식", kind="brand", wrap=False)
            self.put(1, half + 1, self.pc, "작성 예시 (가상의 현장)  ", kind="brand_ex", wrap=False)
        else:
            self.put(1, 1, half, "  K-건설맵  |  k-conmap.com  무료 건설 서식", kind="brand", wrap=False)
            self.put(1, half + 1, self.pc, "← 이 1행을 지우고 쓰셔도 됩니다  ", kind="brand_r", wrap=False)
        self.ws.row_dimensions[1].height = 20
        self.r = 2

    # ── 칸 하나(병합) ─────────────────────────────────────────────────────
    def put(self, r, c1, c2, content=None, *, kind="text", rs=1, fmt=None, align=None, size=None,
            bold=None, wrap=True, border=None, fill=None, ink=None, valign="center", indent=0):
        ws = self.ws
        r2 = r + rs - 1
        if c2 > c1 or r2 > r:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(row=r, column=c1)
        key = None
        is_calc = False
        is_input = False
        rate = False
        if isinstance(content, I):
            key = content.key
            is_input = True
            rate = content.rate
            fmt = fmt or content.fmt
            align = align or content.align
            size = size or content.size
            bold = content.bold if bold is None else bold
            wrap = content.wrap
            val = content.ex if self.ex else content.blank
            if self.ex and content.ex is None and content.blank is not None:
                val = content.blank
            if val is not None:
                if isinstance(val, str) and val.startswith("="):
                    self.formulas.append((cell, val))
                    is_calc = True
                else:
                    cell.value = val
            if key:
                self.inputs[key] = val
            if content.dv:
                self.dvs.setdefault(content.dv, []).append(f"{CL(c1)}{r}")
        elif isinstance(content, F):
            key = content.key
            is_calc = True
            fmt = fmt or content.fmt
            align = align or content.align
            size = size or content.size
            bold = content.bold if bold is None else bold
            wrap = content.wrap
            self.formulas.append((cell, content.f))
            if key:
                self.calc_keys[key] = content.f
        elif isinstance(content, str) and content.startswith("="):
            is_calc = True
            self.formulas.append((cell, content))
        elif content is not None:
            cell.value = content
        if key:
            if key in self.k:
                raise ValueError(f"[{self.name}] 이름 겹침: {key}")
            self.k[key] = f"{CL(c1)}{r}"

        # 모양
        k = kind
        if is_input and k == "text":
            k = "in"
        if is_calc and k in ("text", "in"):
            k = "calc"
        st = {
            "brand": dict(size=10, bold=True, ink="0F203F", border=False, align="left"),
            "brand_r": dict(size=8.5, ink="6B7280", border=False, align="right"),
            "brand_ex": dict(size=9, bold=True, ink="B45309", border=False, align="right"),
            "title": dict(size=18, bold=True, ink="111827", border=False, align="center"),
            "sub": dict(size=10, ink="374151", border=False, align="center"),
            "h2": dict(size=11, bold=True, ink="111827", border=False, align="left"),
            "label": dict(size=10, bold=True, ink=INK, fill=LABEL, align="center"),
            "head": dict(size=10, bold=True, ink=INK, fill=HEAD, align="center"),
            "text": dict(size=10, ink=INK, align="left"),
            "ctext": dict(size=10, ink=INK, align="center"),
            "in": dict(size=10, ink="000000", align=None),
            "calc": dict(size=10, ink=CALC_INK, align=None),
            "sum": dict(size=10, bold=True, ink=INK, fill=SUMF, align="center"),
            "free": dict(size=10, ink=INK, border=False, align="left"),
            "cfree": dict(size=10, ink=INK, border=False, align="center"),
            "note": dict(size=8.5, ink="6B7280", border=False, align="left"),
            "warn": dict(size=9, bold=True, ink="DC2626", border=False, align="left"),
        }[k]
        size = size or st["size"]
        bold = st.get("bold", False) if bold is None else bold
        calc_look = is_calc and k not in ("title", "sub", "free", "cfree", "h2", "warn", "note") and getattr(content, "ink", True)
        ink = ink or (CALC_INK if calc_look else st["ink"])
        if calc_look and border is not False and not fill and not st.get("fill"):
            fill = CALC
        if rate:
            fill = fill or RATE
        fill = fill or st.get("fill")
        border = st.get("border", True) if border is None else border
        if align is None:
            align = st.get("align")
        if align is None:
            align = "right" if (fmt and fmt not in ("text", "date", "ymd", "md", "hangul")) else "left"
            if fmt in ("date", "ymd", "md"):
                align = "center"
        cell.font = Font(name=FONT, size=size, bold=bold, color=ink)
        # 숫자 칸은 «칸에 맞춰 줄이기» — 좁은 칸에 억 단위 금액이 들어와도 ### 가 되지 않게 (엑셀·한셀·LibreOffice 공통)
        numeric = bool(fmt) and fmt not in ("text", "date", "hangul") and not str(NF.get(fmt, fmt)).startswith("[DBNum")
        shrink = numeric and (is_calc or is_input) and border is not False
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=(wrap and not shrink), indent=indent,
                                   shrink_to_fit=shrink)
        if fmt:
            code = NF.get(fmt, fmt)
            cell.number_format = code
            if re.search(r"0\.#+", code) and not code.startswith("[DBNum"):
                self.intfmt.setdefault(code, []).append((c1, r))
        if fill:
            cell.fill = fill
        if border:
            b = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
            for rr in range(r, r2 + 1):
                for cc in range(c1, c2 + 1):
                    ws.cell(row=rr, column=cc).border = b
        for rr in range(r + 1, r2 + 1):
            for cc in range(c1, c2 + 1):
                self.occ.add((rr, cc))
        return cell

    # ── 한 줄 ─────────────────────────────────────────────────────────────
    def row(self, cells, h=None, minh=19, start=1):
        """cells: [(칸수, 내용, {선택})…]. 세로 병합으로 차 있는 칸은 알아서 건너뜁니다.
        h 를 안 주면 글자 길이로 높이를 어림합니다."""
        r = self.r
        c = start
        need = minh
        for spec in cells:
            span, content = spec[0], spec[1]
            opt = dict(spec[2]) if len(spec) > 2 else {}
            while (r, c) in self.occ:
                c += 1
            c2 = c + span - 1
            if c2 > self.ncol:
                raise ValueError(f"[{self.name}] {r}행: 칸이 넘칩니다 ({c2}>{self.ncol})")
            self.put(r, c, c2, content, **opt)
            # 높이 어림(세로 병합 칸은 빼고)
            if opt.get("rs", 1) == 1:
                txt = content if isinstance(content, str) and not content.startswith("=") else None
                if isinstance(content, I) and self.ex and isinstance(content.ex, str):
                    txt = content.ex
                if isinstance(content, I) and not self.ex and isinstance(content.blank, str) and not content.blank.startswith("="):
                    txt = content.blank
                if txt:
                    wu = sum(self.widths[c - 1:c2])
                    sz = opt.get("size") or 10
                    lines = text_lines(txt, wu * 10.0 / sz)
                    need = max(need, lines * (sz * 1.45) + 5)
            c = c2 + 1
        self.ws.row_dimensions[r].height = h or need
        self.r += 1
        return r

    def vlabel(self, r1, r2, c, text, rotate=True, kind="label"):
        """세로로 긴 구분 칸(예: 순공사원가 · 재료비). 줄을 다 적은 뒤에 부릅니다."""
        from openpyxl.styles import Alignment
        cell = self.put(r1, c, c, text, kind=kind, rs=r2 - r1 + 1)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True,
                                   textRotation=255 if rotate else 0)
        return cell

    def gap(self, h=6):
        self.ws.row_dimensions[self.r].height = h
        self.r += 1

    def title(self, text, h=40, sub=None):
        self.row([(self.pc, text, {"kind": "title"})], h=h)
        if sub:
            self.row([(self.pc, sub, {"kind": "sub"})], h=18)

    def para(self, text, *, size=10, align="left", h=None, kind="free", indent=0):
        """테두리 없는 글 한 문단(자동 높이)."""
        return self.row([(self.pc, text, {"kind": kind, "size": size, "align": align, "valign": "top" if align == "left" else "center", "indent": indent})], h=h)

    def sign_block(self, rows, *, lab_span, val_span, tail="(인)", left=None, h=24):
        """서명줄: [(라벨, 입력키, 예시[, 꼬리]), …] 를 오른쪽에 붙여 적습니다. 꼬리 기본 «(인)», '' 이면 없음."""
        n = self.pc
        left = n - lab_span - val_span - 1 if left is None else left
        for it in rows:
            lab, key, exv = it[0], it[1], it[2]
            tl = it[3] if len(it) > 3 else tail
            cells = []
            if left > 0:
                cells.append((left, "", {"kind": "free"}))
            cells += [(lab_span, lab, {"kind": "free", "align": "distributed", "bold": True}),
                      (val_span, I(key, ex=exv, align="center"), {"border": False}),
                      (n - left - lab_span - val_span, tl, {"kind": "free", "align": "center", "size": 9})]
            self.row(cells, h=h)

    def date_line(self, key="작성일", ex=None, h=26):
        """가운데 「2026년 9월 29일」 — 빈 서식에서는 「년 월 일」 자리글."""
        n = self.pc
        self.row([(n, I(key, ex=ex, blank="          년        월        일", fmt="date", align="center"),
                   {"border": False})], h=h)

    def end_print(self):
        """여기까지가 인쇄되는 서식. 아래는 파일에서만 보이는 안내."""
        self.last_print = self.r - 1
        self.ws.print_area = f"A1:{CL(self.pc)}{self.last_print}"

    def after_note(self, lines):
        self.r += 1
        for t in lines:
            self.row([(self.pc, t, {"kind": "note", "wrap": True})])

    def dv_list(self, name, items):
        self.book.dv_defs[name] = items

    def col_px(self, c1, c2):
        """열 너비(저장값) → 픽셀 어림(맑은 고딕 기본 7px/글자)."""
        return sum(int(self.widths[i - 1] * 7) for i in range(c1, c2 + 1))

    def row_px(self, r):
        h = self.ws.row_dimensions[r].height or 15
        return h * 96 / 72

    def place_image(self, data, c1, c2, r1, r2, pad=5):
        """사진을 (c1..c2, r1..r2) 칸 안에 비율을 지켜 가운데 맞춰 넣습니다. 행 높이는 먼저 정해져 있어야 합니다."""
        import io
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        from PIL import Image as PImage
        iw, ih = PImage.open(io.BytesIO(data)).size
        bw = self.col_px(c1, c2) - 2 * pad
        bh = sum(self.row_px(r) for r in range(r1, r2 + 1)) - 2 * pad
        sc = min(bw / iw, bh / ih)
        w, h = iw * sc, ih * sc
        ox, oy = pad + (bw - w) / 2, pad + (bh - h) / 2
        c, r = c1, r1
        while c < c2 and ox >= self.col_px(c, c):
            ox -= self.col_px(c, c)
            c += 1
        while r < r2 and oy >= self.row_px(r):
            oy -= self.row_px(r)
            r += 1
        img = XLImage(io.BytesIO(data))
        img.width, img.height = w, h
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=c - 1, colOff=pixels_to_EMU(int(ox)), row=r - 1, rowOff=pixels_to_EMU(int(oy))),
                                   ext=XDRPositiveSize2D(pixels_to_EMU(int(w)), pixels_to_EMU(int(h))))
        self.ws.add_image(img)


def _int_cf(p):
    """엑셀은 #,##0.## 서식에 정수가 오면 «862.» 처럼 끝에 점을 찍습니다(리브레오피스는 안 찍음).
    그래서 그런 칸에 «정수면 #,##0» 조건부 서식을 겁니다 — 세로로 이어진 칸끼리 한 규칙."""
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.styles.numbers import NumberFormat
    ids = p.book.__dict__.setdefault("_cf_ids", {})      # 같은 파일 안에서 서식 번호가 겹치지 않게
    for code, cells in p.intfmt.items():
        ic = re.sub(r"\.#+", "", code)
        nid = ids.setdefault(ic, 300 + len(ids))
        runs = []
        for c, r in sorted(set(cells)):
            if runs and runs[-1][0] == c and runs[-1][2] == r - 1:
                runs[-1][2] = r
            else:
                runs.append([c, r, r])
        for c, r1, r2 in runs:
            top = f"{CL(c)}{r1}"
            rng = top if r1 == r2 else f"{top}:{CL(c)}{r2}"
            rule = Rule(type="expression", formula=[f"AND(ISNUMBER({top}),MOD({top},1)=0)"],
                        dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=nid, formatCode=ic)))
            p.ws.conditional_formatting.add(rng, rule)


class Book:
    """한 파일. sets: [('blank', {...}), ('ex', {...})] 두 벌을 그리고 수식을 이어 붙입니다."""

    def __init__(self, title):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.title = title
        self.pages = {}          # (벌, 논리이름) → Page
        self.dv_defs = {}
        self.wb.properties.creator = "K-건설맵 (k-conmap.com)"
        self.wb.properties.title = title

    def page(self, logical, widths, *, ex, sheetname=None, **kw):
        if sheetname is None:
            sheetname = logical
        bad = '[]:*?/\\'
        sheetname = "".join(ch for ch in sheetname if ch not in bad)[:31]
        p = Page(self, sheetname, widths, ex=ex, logical=logical, **kw)
        self.pages[("ex" if ex else "blank", logical)] = p
        return p

    def _resolve(self, page, tmpl):
        st = "ex" if page.ex else "blank"

        def rep(m):
            k = m.group(1)
            if "!" in k:
                lg, kk = k.split("!", 1)
                other = self.pages.get((st, lg))
                if other is None:
                    raise KeyError(f"[{page.name}] 없는 시트: {lg}")
                if kk not in other.k:
                    raise KeyError(f"[{page.name}] {lg} 에 없는 이름: {kk}")
                return f"'{other.name}'!{other.k[kk]}"
            if k not in page.k:
                raise KeyError(f"[{page.name}] 없는 이름: {k}")
            return page.k[k]

        return re.sub(r"\{([^{}]+)\}", rep, tmpl)

    def finish(self):
        for (st, lg), p in self.pages.items():
            for cell, tmpl in p.formulas:
                cell.value = self._resolve(p, tmpl)
            for name, cells in p.dvs.items():
                items = self.dv_defs.get(name) or name.split(",")
                dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True,
                                    showDropDown=False)
                dv.error = "목록에서 고르세요"
                p.ws.add_data_validation(dv)
                for a in cells:
                    dv.add(a)
            _int_cf(p)
            if p.last_print is None:
                p.end_print()
        return self.wb

    def save(self, path):
        self.finish()
        self.wb.save(path)
