# -*- coding: utf-8 -*-
"""예정공정표(변경) · 동원인력 계획표(변경) · 장비투입계획서(변경) 한 벌 — 현장에서 쓰던 서식의 틀
(공사명·공사기간 · 공종마다 두 줄(당초/변경) · 보할(%)·계 · 착공일로부터 30일~360일 눈금 12칸 · 계·누계 · 비고 · 주소/상호/대표이사) 그대로.

■ 첫 장(예정공정표)에 공사명·착공일·준공일·회사를 적으면 뒤 두 장의 머리·눈금·날인란이 저절로 채워집니다.
■ 칸마다 %(공정표) · 인·일(인력) · 대·일(장비)을 적으면 줄마다 보할·계, 눈금마다 계·누계가 나옵니다.
   공정표는 당초·변경 보할 합계가 100%가 아니면 빨간 경고가 나옵니다.
■ 원본은 칸마다 «보할 ÷ 8» 같은 수식을 손으로 넣어 두었던 것 — 공사마다 다르니 값을 바로 적게 바꿨습니다.
■ 원본 눈금에는 날짜가 없어서 «30일» 아래에 착공일 기준 날짜(예: 3.9~4.7)를 붙였습니다."""
import datetime

from fb import F, I
from openpyxl.formatting.rule import Rule
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as CL

SLUG = "gongjeong-inryeok"
TITLE = "예정공정표·동원인력계획표"
WHERE = "forms"
PREV = [(1, "예정공정표 — 빈 서식"), (4, "예정공정표 — 작성 예시"), (5, "동원인력 계획표 — 작성 예시"), (6, "장비투입계획서 — 작성 예시")]
SHORT = "예정공정표(변경)·동원인력 계획표·장비투입계획서 세 장이 한 파일입니다. 착공일로부터 30~360일 눈금에 값만 적으면 보할·계·누계와 날짜가 저절로 채워집니다."
NOTE_REPLACE = {
    "🏗 **현장에서 실제로 쓰던 서식입니다** — 내용만 지우고 틀은 원본 그대로 두었습니다.":
        "🏗 **현장에서 실제로 쓰던 서식의 틀**(세 장 · 당초/변경 두 줄 · 30~360일 눈금 · 날인란)을 그대로 두고 새로 만들었습니다 — 보할·계·누계·눈금 날짜가 자동입니다.",
}
MULTI = True

D = datetime.date
NB = 12
NR = 8
W = [16, 5.5, 7.5] + [7.2] * NB + [10]
C0 = 4
V = '#,##0.#;[Red]-#,##0.#;""'
GRAY = PatternFill("solid", fgColor="E5E7EB", bgColor="E5E7EB")
BLUE = PatternFill("solid", fgColor="BFD4F2", bgColor="BFD4F2")

PLAN = {  # 공종: (당초, 변경) — %
    "토공": ([3, 4, 4, 4], [3, 4, 4, 4]),
    "배수공": ([0, 8, 10, 10, 7], [0, 6, 8, 9, 7, 5]),
    "구조물공": ([0, 0, 0, 6, 8, 8, 8], [0, 0, 0, 0, 6, 8, 8, 8]),
    "포장공": ([0, 0, 0, 0, 0, 0, 0, 6, 6], [0, 0, 0, 0, 0, 0, 0, 0, 6, 6]),
    "부대공": ([0.8] * 10, [0.8] * 10),
}
MAN = {
    "기능공": ([60, 150, 180, 220, 200, 150, 130, 90, 80, 20], [60, 130, 160, 180, 190, 190, 150, 120, 90, 60]),
    "보통인부": ([40, 90, 110, 130, 120, 90, 80, 60, 50, 20], [40, 80, 100, 110, 115, 115, 90, 70, 60, 40]),
}
EQ = {
    "굴착기": ([18, 24, 24, 24, 20, 10, 8, 6, 6, 2], [18, 22, 22, 22, 22, 14, 8, 6, 6, 4]),
    "대형브레이커": ([0, 4, 6, 6, 4], [0, 4, 4, 6, 6, 2]),
    "크레인": ([0, 0, 0, 4, 6, 6, 4], [0, 0, 0, 0, 4, 6, 6, 4]),
    "플레이트 콤팩터": ([4, 8, 8, 8, 6, 4, 4, 6, 6, 2], [4, 8, 8, 8, 8, 6, 4, 4, 6, 6]),
    "크롤러 드릴": ([0, 0, 2, 2], [0, 0, 0, 2, 2]),
}
SHEETS = [("공정표", "예정공정표(변경)", "예 정 공 정 표  (변 경)", "월별\n공종", "보할(%)", PLAN, "%"),
          ("인력", "동원인력계획표(변경)", "동 원 인 력  계 획 표  (변 경)", "일별\n직종", "계", MAN, "인·일"),
          ("장비", "장비투입계획서(변경)", "장 비 투 입  계 획 서  (변 경)", "일별\n장비", "계", EQ, "대·일")]


def _pull(k, fmt=None, align=None):
    return F(k, f'=IF({{공정표!{k}}}="","",{{공정표!{k}}})', fmt=fmt, align=align)


def _sheet(bk, ex, lg, sname, title, corner, tot_label, data, unit):
    first = lg == "공정표"
    p = bk.page(lg, W, ex=ex, landscape=True, fit_height=1, sheetname=sname if not ex else "예시-" + sname,
                margins=(0.35, 0.35, 0.45, 0.45))
    L = {"kind": "label"}
    H = {"kind": "head"}
    p.title(title, h=34)
    p.row([(1, "공  사  명", L), (7, I("공사명", ex="가나지구 배수로 정비공사") if first else _pull("공사명")),
           (2, "착 공 일", L), (3, I("착공일", ex=D(2026, 3, 9), fmt="date", align="center") if first else _pull("착공일", "date", "center")),
           (1, "준공일", L), (2, I("준공일", ex=D(2026, 12, 31), fmt="date", align="center") if first else _pull("준공일", "date", "center"))],
          h=24)
    s, e = ("{착공일}", "{준공일}") if first else ("{공정표!착공일}", "{공정표!준공일}")
    p.row([(1, "공 사 기 간", L),
           (15, F("공사기간", f'=IF(OR(N({s})=0,N({e})=0),"",TEXT({s},"yyyy. m. d.")&" ~ "&TEXT({e},"yyyy. m. d.")&"  (총 "&TEXT({e}-{s}+1,"#,##0")&"일)")',
                  align="left"))], h=22)
    p.gap(4)
    p.row([(2, corner, {"kind": "head", "rs": 3}), (1, tot_label, {"kind": "head", "rs": 3}), (NB, "공   사   기   간", H),
           (1, "비  고", {"kind": "head", "rs": 3})], h=20)
    p.row([(1, f"{30 * (k + 1)}일", {"kind": "head", "size": 9}) for k in range(NB)], h=18, start=C0)
    p.row([(1, F(f"기간#{k + 1}", f'=IF(N({s})=0,"",MONTH({s}+{30 * k})&"."&DAY({s}+{30 * k})&"~"&MONTH({s}+{30 * (k + 1) - 1})&"."&DAY({s}+{30 * (k + 1) - 1}))',
                 align="center", size=7), {"kind": "head"}) for k in range(NB)], h=16, start=C0)
    names = list(data)
    r0 = p.r
    for i in range(1, NR + 1):
        nm = names[i - 1] if i <= len(names) else None
        rows = data.get(nm, ([], []))
        for half, tag in ((0, "당초"), (1, "변경")):
            vals = rows[half]
            cells = []
            if half == 0:
                cells.append((1, I(f"이름#{i}", ex=nm, align="center", size=9.5), {"rs": 2}))
            cells.append((1, tag, {"kind": "ctext", "size": 8.5}))
            key = "ab"[half]
            cells.append((1, F(f"{key}{i}계", f"=IF(SUM({{{key}{i}[]}})=0,\"\",SUM({{{key}{i}[]}}))", fmt=V, align="center")))
            rr = p.r
            for k in range(NB):
                v = vals[k] if k < len(vals) and vals[k] else None
                cells.append((1, I(f"{key}{i}_{k + 1}", ex=v, fmt=V, align="center", size=9)))
            cells.append((1, I(f"{key}{i}비고", size=8), {"rs": 1}))
            p.row(cells, h=17)
            p.k[f"{key}{i}[]"] = f"{CL(C0)}{rr}:{CL(C0 + NB - 1)}{rr}"
    r1 = p.r - 1
    for key, fill in (("당초", GRAY), ("변경", BLUE)):
        p.ws.conditional_formatting.add(
            f"{CL(C0)}{r0}:{CL(C0 + NB - 1)}{r1}",
            Rule(type="expression", dxf=DifferentialStyle(fill=fill), formula=[f'AND(ISNUMBER({CL(C0)}{r0}),{CL(C0)}{r0}>0,$B{r0}="{key}")']))
    S = {"kind": "sum"}
    for lab, pre in (("계", "계"), ("누    계", "누")):
        for half, tag in ((0, "당초"), (1, "변경")):
            key = "ab"[half]
            cells = []
            if half == 0:
                cells.append((1, lab, {"kind": "sum", "rs": 2}))
            cells.append((1, tag, {"kind": "ctext", "size": 8.5}))
            if pre == "계":
                cells.append((1, F(f"계{key}", "=" + "+".join(f"N({{{key}{i}계}})" for i in range(1, NR + 1)), fmt=V, align="center", bold=True)))
            else:
                cells.append((1, "", S))
            for k in range(1, NB + 1):
                if pre == "계":
                    f = "=" + "+".join(f"N({{{key}{i}_{k}}})" for i in range(1, NR + 1))
                    cells.append((1, F(f"계{key}_{k}", f"=IF({f[1:]}=0,\"\",{f[1:]})", fmt=V, align="center", bold=True)))
                else:
                    f = f"=IF(N({{계{key}_{k}}})=0,\"\",SUM({{계{key}_1}}:{{계{key}_{k}}}))"
                    cells.append((1, F(f"누{key}_{k}", f, fmt=V, align="center")))
            cells.append((1, "", S))
            p.row(cells, h=18)
    if first:
        p.row([(3 + NB + 1, F("경고", '=IF(AND(N({계a})>0,ABS({계a}-100)>0.001),"⚠ 당초 보할 합계가 "&TEXT({계a},"0.0")&"% — 100%가 되어야 합니다.  ","")'
                                     '&IF(AND(N({계b})>0,ABS({계b}-100)>0.001),"⚠ 변경 보할 합계가 "&TEXT({계b},"0.0")&"% — 100%가 되어야 합니다.","")',
                               align="left"), {"kind": "warn", "border": False})], h=16)
    else:
        p.row([(3 + NB + 1, f"※ 단위: {unit} (눈금 한 칸 = 30일)", {"kind": "note"})], h=16)
    p.gap(8)
    for lab, k, exv in (("주       소 :", "주소", "가나시 마바로 45, 2층"), ("상       호 :", "상호", "예시건설(주)"),
                        ("대 표 이 사 :", "대표", "홍길동")):
        val = I(k, ex=exv) if first else _pull(k)
        p.row([(9, "", {"kind": "free"}), (2, lab, {"kind": "free", "bold": True, "align": "right"}),
               (4, val, {"border": False}), (1, "(인)" if k == "대표" else "", {"kind": "free", "align": "center"})], h=22)
    p.end_print()
    if first:
        p.after_note(["칸마다 그 30일 동안의 계획 공정률(%)을 적습니다. 줄 끝 보할(%)과 아래 계·누계는 자동입니다.",
                      "공사명·착공일·준공일·회사는 이 장에만 적으면 뒤 두 장(동원인력·장비)에 저절로 들어갑니다."] if not ex else
                     ["이 시트는 작성 예시입니다(가상의 공사 · 수치)."])
    else:
        p.after_note([f"칸마다 그 30일 동안의 {unit} 을 적습니다. 줄 끝 계와 아래 계·누계는 자동입니다."] if not ex else
                     ["이 시트는 작성 예시입니다(가상의 공사 · 수치)."])
    return p


def draw(bk, ex):
    for lg, sname, title, corner, tot, data, unit in SHEETS:
        _sheet(bk, ex, lg, sname, title, corner, tot, data, unit)


def _fmt(d):
    return f"{d.month}.{d.day}"


def expect(pages):
    src = pages["공정표"]
    s, e = src["착공일"], src["준공일"]
    out = {}
    for lg, _sn, _t, _c, _tot, data, _u in SHEETS:
        pg = pages[lg]
        x = {"공사기간": f"{s.year}. {s.month}. {s.day}. ~ {e.year}. {e.month}. {e.day}.  (총 {(e - s).days + 1:,}일)"}
        for k in range(NB):
            a = s + datetime.timedelta(days=30 * k)
            b = s + datetime.timedelta(days=30 * (k + 1) - 1)
            x[f"기간#{k + 1}"] = f"{_fmt(a)}~{_fmt(b)}"
        for key in "ab":
            tot = 0
            cum = 0
            for i in range(1, NR + 1):
                row = [pg.get(f"{key}{i}_{k}") or 0 for k in range(1, NB + 1)]
                x[f"{key}{i}계"] = sum(row) if sum(row) else ""
                tot += sum(row)
            x[f"계{key}"] = tot
            for k in range(1, NB + 1):
                c = sum(pg.get(f"{key}{i}_{k}") or 0 for i in range(1, NR + 1))
                x[f"계{key}_{k}"] = c if c else ""
                cum += c
                x[f"누{key}_{k}"] = cum if c else ""
        if lg == "공정표":
            x["경고"] = ""
        else:
            for k in ("공사명", "착공일", "준공일", "주소", "상호", "대표"):
                x[k] = src[k]
        out[lg] = x
    return out
