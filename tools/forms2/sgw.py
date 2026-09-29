# -*- coding: utf-8 -*-
"""sgw.py — 한글·엑셀로 받은 «공종별 시공계획서»(방수 · 미장 · 도장 · 팽이기초 · 지반개량 …)를 같은 틀로 새로 만드는 공통 부품 (2026-09-29)
  표지 · 목차 · 공사개요(공사명 한 번 → 모든 장, 공사기간 일수, 공사금액 한글) · 위치도 자리 · 주요공사 물량(단위별 합계) ·
  현장 조직도(인원 합계) · 투입인원 표(계) · 검측 흐름도 · 예정공정표(반달 막대 18개월) · 글 덩이(제목 · 글 · 목록 · 표)."""
import calendar
import datetime

from openpyxl.formatting.rule import Rule
from openpyxl.styles import PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as CL

import docw as d
from fb import F, I, text_lines

D = datetime.date
W = [7.6] * 12
PC = 12
L = {"kind": "label"}
H = {"kind": "head"}
FR = {"kind": "free"}
S = {"kind": "sum"}
BAR = PatternFill("solid", fgColor="2F5597", bgColor="2F5597")
NB = 36          # 공정표 눈금 — 반달 × 18달


def pg(bk, ex, lg, fit=0, land=False, widths=None):
    return bk.page(lg, widths or W, ex=ex, fit_height=fit, landscape=land, sheetname=lg if not ex else "예시-" + lg,
                   margins=(0.6, 0.6, 0.6, 0.6) if not land else (0.35, 0.35, 0.45, 0.45))


# ── 글 덩이 ─────────────────────────────────────────────────────────
def ch(p, t, new=True, need=120):
    """장 제목(1.0 공사개요 …) — new 면 새 쪽에서, 아니면 need(pt) 가 안 남으면 넘김"""
    if new and p.r > 3:
        p.__dict__.setdefault("_brk", set()).add(p.r)
    elif not new:
        d.keep(p, need)
    p.gap(4)
    p.row([(PC, t, {"kind": "h2", "size": 14})], h=30)
    p.gap(4)


def sec(p, t, need=90):
    d.keep(p, need)
    p.gap(3)
    p.row([(PC, t, {"kind": "h2", "size": 11.5})], h=22)


def b(p, t, need=100, lead=0):
    d.keep(p, need)
    cells = ([(lead, "", FR)] if lead else []) + [(PC - lead, t, {"kind": "free", "bold": True, "valign": "top"})]
    n = text_lines(t, sum(p.widths[lead:PC]) * 10.0 / 10.5)
    p.row(cells, h=max(21, n * 15.5 + 5))


def para(p, t, indent=2, size=10):
    d.para(p, t, indent=indent, size=size)


def li(p, lst, marks=None, lead=1, size=10):
    if marks == "(n)":
        marks = [f"({i})" for i in range(1, len(lst) + 1)]
    elif marks == "가":
        marks = [c + "." for c in "가나다라마바사아자차카타파하"[:len(lst)]]
    elif marks == "o":
        marks = ["◦"] * len(lst)
    elif marks == "①":
        marks = [chr(0x2460 + i) for i in range(len(lst))]
    first = text_lines(lst[0], sum(p.widths[lead:p.pc]) * 10.0 / size) * size * 1.55 + 5
    d.keep(p, first + 24)
    d.items(p, lst, marks=marks, lead=lead, size=size)


def table(p, heads, rows, h=24, size=9.5, need=None):
    if need:
        d.keep(p, need)
    d.table(p, heads, rows, h=h, size=size)


def blank_box(p, t, rows=12, rh=20):
    d.keep(p, rows * rh + 4)
    r0 = p.r
    for _ in range(rows):
        p.gap(rh)
    p.put(r0, 1, PC, t, kind="ctext", size=9, rs=rows)
    p.ws.cell(row=r0, column=1).font = p.ws.cell(row=r0, column=1).font.copy(color="9CA3AF")


def note(p, t):
    n = text_lines(t, sum(p.widths[:PC]) * 10.0 / 8.5)
    p.row([(PC, t, {"kind": "note", "valign": "top"})], h=n * 8.5 * 1.45 + 5)


# ── 표지 · 목차 ─────────────────────────────────────────────────────
def cover(bk, ex, kind, name_ex, month_ex, co_ex, title="시  공  계  획  서"):
    p = pg(bk, ex, "표지", fit=1)
    p.gap(90)
    p.row([(PC, I("공사명", ex=name_ex, blank="(공 사 명)", align="center", size=16, bold=True), {"border": False})], h=40)
    p.gap(70)
    p.row([(PC, title, {"kind": "title", "size": 30})], h=64)
    p.row([(PC, f"[ {kind} ]", {"kind": "sub", "size": 18, "bold": True})], h=40)
    p.gap(120)
    p.row([(PC, I("작성월", ex=month_ex, blank="20    .    .", fmt='yyyy". "m"."', align="center", size=15), {"border": False})], h=32)
    p.gap(110)
    p.row([(PC, I("회사", ex=co_ex, blank="(회 사 명)", align="center", size=18, bold=True), {"border": False})], h=44)
    p.end_print()
    p.after_note(["공사명 · 작성 연월(2026-03 처럼) · 회사명을 여기 한 번만 적으면 뒤 장에 저절로 들어갑니다."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 공사 · 회사 · 사람)."])
    return p


def toc(p, items):
    p.gap(16)
    p.row([(PC, "목          차", {"kind": "title", "size": 22})], h=50)
    p.gap(14)
    for t in items:
        p.row([(2, "", FR), (8, t, {"kind": "free", "bold": True, "size": 12.5}), (2, "", FR)], h=30)


# ── 공사개요 · 물량 · 조직도 ──────────────────────────────────────────
def name_line(p, key="공사명1"):
    """장 «공사명» — 공사명 한 줄(표지에서 따라옴)"""
    p.row([(1, "", FR), (11, F(key, '=IF(OR({표지!공사명}="",{표지!공사명}="(공 사 명)"),"","◦ "&{표지!공사명})', align="left", size=13, bold=True),
            {"kind": "free"})], h=30)


def name_row(p, label="공 사 명"):
    p.row([(3, label, L), (9, F("공사명표시", '=IF(OR({표지!공사명}="",{표지!공사명}="(공 사 명)"),"",{표지!공사명})', align="left"))], h=28)


def overview(p, ex, loc_ex, owner_ex, st_ex, en_ex, amt_ex, scale_ex, rows=None):
    name_row(p)
    p.row([(3, "공 사 위 치", L), (9, I("위치", ex=loc_ex))], h=26)
    p.row([(3, "발  주  처", L), (9, I("발주처", ex=owner_ex))], h=26)
    p.row([(3, "공 사 기 간", L), (2, I("착공일", ex=st_ex, fmt="ymd", align="center")), (1, "~", {"kind": "ctext"}),
           (2, I("준공일", ex=en_ex, fmt="ymd", align="center")),
           (4, F("공사일수", '=IF(OR(N({착공일})=0,N({준공일})=0),"",IF({준공일}<{착공일},"⚠ 준공일이 착공일보다 빠름",{준공일}-{착공일}+1))',
                 fmt='"(총 "#,##0"일)";;""', align="left"))], h=26)
    p.row([(3, "공 사 금 액", L), (3, I("공사금액", ex=amt_ex, fmt='"₩"#,##0;;""')),
           (6, F("금액한글", '=IF(N({공사금액})=0,"",{공사금액})', fmt="hangul", align="left", size=9.5))], h=26)
    p.row([(3, "공 사 규 모", L), (9, I("규모", ex=scale_ex, size=9.5))], h=26)
    for lab, key, v in rows or []:
        p.row([(3, lab, L), (9, I(key, ex=v, size=9.5))], h=26)
    p.row([(3, "시  공  사", L), (9, F("시공사표시", '=IF(OR({표지!회사}="",{표지!회사}="(회 사 명)"),"",{표지!회사})', align="left"))], h=26)


def qty(p, ex, data, n=12, unit="㎡", unit_first=False, h0="구  분", h1="명  칭"):
    """주요공사 물량 — 구분 · 명칭 · 규격 · 수량 · 단위 · 비고(unit_first 면 단위 · 수량 순), 아래에 «unit» 합계"""
    d.keep(p, 24 + 24 * (n + 1))
    hq, hu = (2, "수  량", H), (1, "단위", H)
    p.row([(2, h0, H), (3, h1, H), (2, "규  격", H)] + ([hu, hq] if unit_first else [hq, hu]) + [(2, "비  고", H)], h=24)
    for i in range(1, n + 1):
        e = data[i - 1] if (ex and i <= len(data)) else (None,) * 6
        cq = (2, I(f"수량#{i}", ex=e[3], fmt='#,##0.##;;""'))
        cu = (1, I(f"단위#{i}", ex=e[4], align="center", size=9.5))
        p.row([(2, I(f"물구분#{i}", ex=e[0], align="center", size=9.5)), (3, I(f"명칭#{i}", ex=e[1], size=9.5)),
               (2, I(f"규격#{i}", ex=e[2], align="center", size=9))] + ([cu, cq] if unit_first else [cq, cu])
              + [(2, I(f"물비고#{i}", ex=e[5] or None, size=9))], h=23)
    p.k["단위[]"] = f"{p.k['단위#1']}:{p.k[f'단위#{n}']}"
    p.k["수량[]"] = f"{p.k['수량#1']}:{p.k[f'수량#{n}']}"
    uu = "m2" if unit == "㎡" else ("m3" if unit == "㎥" else unit)
    extra = f'+SUMIF({{단위[]}},"{uu}",{{수량[]}})' if uu != unit else ""
    f = f'SUMIF({{단위[]}},"{unit}",{{수량[]}}){extra}'
    tot = (2, F("물량합계", f'=IF({f}=0,"",{f})', fmt='#,##0.##', bold=True))
    p.row([(7, f"합   계 (단위가 {unit} 인 줄)", S)] + ([(1, unit, S), tot] if unit_first else [tot, (1, unit, S)]) + [(2, "", S)], h=24)
    return unit


def _box(p, c1, c2, role, key, ex, name, h=24):
    p.row(([(c1 - 1, "", FR)] if c1 > 1 else []) + [(c2 - c1 + 1, role, {"kind": "label", "size": 10})]
          + ([(PC - c2, "", FR)] if c2 < PC else []), h=h)
    p.row(([(c1 - 1, "", FR)] if c1 > 1 else []) + [(c2 - c1 + 1, I(f"조직:{key}", ex=name if ex else None, align="center"))]
          + ([(PC - c2, "", FR)] if c2 < PC else []), h=h)


def _line(p, t="│", h=14):
    p.row([(PC, t, {"kind": "cfree", "size": 9})], h=h)


def org(p, ex, names, crew, co_ex="예시방수(주)", boss="작 업 반 장"):
    """현장 조직도 — 현장대리인 → 품질 · 안전관리자 → 시공업체 → 작업반장 → 직종별 인원(합계)"""
    d.keep(p, 24 * 2 * 4 + 14 * 3 + 24 * 4)
    _box(p, 5, 8, "현 장 대 리 인", "대리인", ex, names[0])
    _line(p)
    p.row([(1, "", FR), (4, "품 질 관 리 자", {"kind": "label"}), (2, "─┼─", {"kind": "cfree"}), (4, "안 전 관 리 자", {"kind": "label"}), (1, "", FR)], h=24)
    p.row([(1, "", FR), (4, I("조직:품질", ex=names[1] if ex else None, align="center")), (2, "│", {"kind": "cfree"}),
           (4, I("조직:안전", ex=names[2] if ex else None, align="center")), (1, "", FR)], h=24)
    _line(p)
    p.row([(4, "", FR), (4, "시 공 업 체", {"kind": "label"}), (4, "", FR)], h=24)
    p.row([(4, "", FR), (4, I("조직:업체", ex=co_ex if ex else None, align="center")), (4, "", FR)], h=24)
    _line(p)
    _box(p, 5, 8, boss, "반장", ex, names[3])
    _line(p, "┌──────────┼──────────┐")
    n = len(crew)
    sp = PC // (n + 1)
    rest = PC - sp * (n + 1)
    p.row([(sp, t, H) for t, _v in crew] + [(sp + rest, "합  계", S)], h=24)
    p.row([(sp, I(f"인원#{i}", ex=v, blank=v, fmt='#,##0"인";;""', align="center")) for i, (_t, v) in enumerate(crew, 1)]
          + [(sp + rest, F("인원합계", "=IF(SUM(" + ",".join(f"{{인원#{i}}}" for i in range(1, n + 1)) + ")=0,\"\",SUM("
                           + ",".join(f"{{인원#{i}}}" for i in range(1, n + 1)) + "))", fmt='#,##0"인"', align="center", bold=True))], h=24)


def crew(p, key, rows, need_head=None):
    """투입인원(일단위) — 구분 · 단위 · 수량 · 비고, 계 (원본 인원을 미리 넣어 둠 · 현장에 맞게 고쳐 씀)"""
    d.keep(p, (need_head or 0) + 24 * (len(rows) + 2))
    p.row([(4, "구    분", H), (2, "단  위", H), (2, "수  량", H), (4, "비    고", H)], h=24)
    for i, (t, u, v) in enumerate(rows, 1):
        p.row([(4, I(f"{key}직종#{i}", ex=t, blank=t, align="center")), (2, u, {"kind": "ctext"}),
               (2, I(f"{key}수#{i}", ex=v, blank=v, fmt='#,##0;;""', align="center")), (4, I(f"{key}비고#{i}", size=9))], h=24)
    p.row([(4, "계", S), (2, "인", S),
           (2, F(f"{key}계", "=IF(SUM(" + ",".join(f"{{{key}수#{i}}}" for i in range(1, len(rows) + 1)) + ")=0,\"\",SUM("
                 + ",".join(f"{{{key}수#{i}}}" for i in range(1, len(rows) + 1)) + "))", fmt="#,##0", align="center", bold=True)), (4, "", S)], h=24)


def flow(p, owner="감 리 원(공사감독자)"):
    """검측 흐름도 — 원본의 흐름 그대로(시공 완료 → 시공자 점검 → 검측요청서 → 감리원 현장검측 → 결과 통보 → 합격 시 다음 공종)"""
    d.keep(p, 30 * 6 + 14 * 5 + 60)
    d.flow(p, [("현장 시공 완료", None), ("시공자 · 담당기술자 점검", "첨부 : 시공자 점검표"), ("검측요청서 제출", "첨부 : 공사참여자 명부 · 관련 도면 · 사진"),
               (f"{owner} 현장 검측", None), ("검측 결과 통보", "불합격 시 → 재시공 · 보완 후 다시 검측 요청"), ("합격 시 다음 단계 공종 착수", None)])


# ── 예정공정표 ────────────────────────────────────────────────────────
def gongjeong(bk, ex, title, data, n=14, src="개요"):
    Wg = [20, 8.5, 8.5, 5.5] + [1.9] * NB + [7]
    nc = len(Wg)
    C0 = 5
    p = bk.page("공정표", Wg, ex=ex, landscape=True, fit_height=1, sheetname="공정표" if not ex else "예시-공정표",
                margins=(0.35, 0.35, 0.45, 0.45))
    p.row([(nc, title, {"kind": "title", "size": 16})], h=34)
    p.row([(1, "착 공 일", L), (2, F("착공일표시", f'=IF(N({{{src}!착공일}})=0,"",{{{src}!착공일}})', fmt="date", align="center")),
           (NB + 2, "※ 착공일은 공사개요에서 따라옵니다. 공종마다 착수 · 완료일을 적으면 반달 칸에 막대가 그려집니다.", {"kind": "note"})], h=22)
    p.row([(1, "공    종", {"kind": "head", "rs": 2}), (1, "착 수", {"kind": "head", "rs": 2}), (1, "완 료", {"kind": "head", "rs": 2}),
           (1, "일 수", {"kind": "head", "rs": 2})]
          + [(2, F(f"월#{m + 1}", f'=IF(N({{착공일표시}})=0,"",EDATE(DATE(YEAR({{착공일표시}}),MONTH({{착공일표시}}),1),{m}))',
                   fmt='yy"."m;;""', align="center", size=7.5), H) for m in range(NB // 2)]
          + [(1, "비 고", {"kind": "head", "rs": 2})], h=20)
    p.row([(1, "상" if k % 2 == 0 else "하", {"kind": "head", "size": 7}) for k in range(NB)], h=15, start=C0)
    rs_, re_ = p.r, p.r + 1
    for rr, kind in ((rs_, "s"), (re_, "e")):
        for k in range(NB):
            m, half = k // 2, k % 2
            base = f"EDATE(DATE(YEAR({{착공일표시}}),MONTH({{착공일표시}}),1),{m})"
            f = (f"=IF(N({{착공일표시}})=0,\"\",{base}{'+15' if half else ''})" if kind == "s" else
                 (f"=IF(N({{착공일표시}})=0,\"\",{base}+14)" if half == 0 else
                  f"=IF(N({{착공일표시}})=0,\"\",EDATE(DATE(YEAR({{착공일표시}}),MONTH({{착공일표시}}),1),{m + 1})-1)"))
            p.put(rr, C0 + k, C0 + k, F(f"{kind}{k + 1}", f, fmt="yyyy-mm-dd"), kind="free", size=7)
        p.ws.row_dimensions[rr].height = 12
        p.ws.row_dimensions[rr].hidden = True
    p.r = re_ + 1
    r0 = p.r
    for i in range(1, n + 1):
        g = data[i - 1] if i <= len(data) else (None, None, None)
        p.row([(1, I(f"공종#{i}", ex=g[0], blank=g[0], size=9)), (1, I(f"착수#{i}", ex=g[1], fmt="ymd", size=8.5, align="center")),
               (1, I(f"완료#{i}", ex=g[2], fmt="ymd", size=8.5, align="center")),
               (1, F(f"일수#{i}", f'=IF(OR(N({{착수#{i}}})=0,N({{완료#{i}}})=0),"",IF({{완료#{i}}}<{{착수#{i}}},"⚠",{{완료#{i}}}-{{착수#{i}}}+1))',
                     fmt='#,##0;;""', align="center", size=8.5))]
              + [(1, "", {"kind": "text"}) for _ in range(NB)] + [(1, I(f"공정비고#{i}", size=8))], h=22)
    r1 = p.r - 1
    p.ws.conditional_formatting.add(
        f"{CL(C0)}{r0}:{CL(C0 + NB - 1)}{r1}",
        Rule(type="expression", dxf=DifferentialStyle(fill=BAR),
             formula=[f"AND(N($B{r0})>0,N($C{r0})>=N($B{r0}),$B{r0}<={CL(C0)}${re_},$C{r0}>={CL(C0)}${rs_})"]))
    p.row([(nc, "※ 일수 = 완료 − 착수 + 1. 막대는 반달(1~15일 · 16일~말일) 단위입니다. 공정은 현장 여건 · 날씨에 따라 바뀔 수 있습니다.", {"kind": "note"})], h=16)
    d.done(p)
    p.after_note(["하늘색 칸 = 자동(착공일 · 달 머리 · 일수). 공종 이름은 원본 차례를 미리 넣어 두었습니다 — 현장에 맞게 고쳐 쓰세요."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 일정)."])
    return p


# ── expect 도움 ────────────────────────────────────────────────────────
def _edate(dd, m):
    y, mo = dd.year + (dd.month - 1 + m) // 12, (dd.month - 1 + m) % 12 + 1
    return D(y, mo, min(dd.day, calendar.monthrange(y, mo)[1]))


def exp_overview(cv, g):
    days = (g["준공일"] - g["착공일"]).days + 1
    return {"공사명1": "◦ " + cv["공사명"], "공사명표시": cv["공사명"], "시공사표시": cv["회사"], "공사일수": days, "금액한글": g["공사금액"]}


def exp_qty(g, n=12, unit="㎡"):
    uu = {"㎡": "m2", "㎥": "m3"}.get(unit, unit)
    s = sum(g.get(f"수량#{i}") or 0 for i in range(1, n + 1) if str(g.get(f"단위#{i}") or "").lower() in (unit, uu))
    return {"물량합계": s or ""}


def exp_org(g, n):
    s = sum(g.get(f"인원#{i}") or 0 for i in range(1, n + 1))
    return {"인원합계": s or ""}


def exp_crew(g, key, n):
    s = sum(g.get(f"{key}수#{i}") or 0 for i in range(1, n + 1))
    return {f"{key}계": s or ""}


def exp_gongjeong(s, st, n=14):
    z = {"착공일표시": st}
    b0 = D(st.year, st.month, 1)
    for m in range(NB // 2):
        z[f"월#{m + 1}"] = _edate(b0, m)
    for k in range(NB):
        m, half = k // 2, k % 2
        base = _edate(b0, m)
        z[f"s{k + 1}"] = base + datetime.timedelta(days=15 if half else 0)
        z[f"e{k + 1}"] = base + datetime.timedelta(days=14) if half == 0 else _edate(b0, m + 1) - datetime.timedelta(days=1)
    for i in range(1, n + 1):
        a, bb = s.get(f"착수#{i}"), s.get(f"완료#{i}")
        z[f"일수#{i}"] = ("⚠" if bb < a else (bb - a).days + 1) if a and bb else ""
    return z


# ── 품질 · 환경 · 안전관리 계획 ─────────────────────────────────────
def quality_std(p, a):
    """품질관리 계획(방수 · 미장 공통) — a = 장 번호"""
    ch(p, f"{a}.0   품 질 관 리  계 획")
    sec(p, f"{a}-1  품질관리 일반 및 적용범위")
    li(p, ["공사 진행 중 필요에 따라 각종 승인도면 · 제작도면 · 제작요령서 등을 작성하여 공사감독자(감리원)의 승인을 받는다.",
             "공사용 재료는 도면 · 공사시방서 및 공사감독자의 지시에 따라, 쓰기 전에 견본 또는 자료를 제출하여 승인을 받은 뒤 쓴다.",
             "품질시험은 「건설기술 진흥법」 및 같은 법 시행령 · 시행규칙과 공사시방서에서 정한 바에 따른다."], marks="(n)")
    sec(p, f"{a}-2  공사용 재료의 품질")
    li(p, ["설계도면 · 공사시방서 및 공사감독자의 따로 지시가 없으면 시방서에서 정한 품질과 규격에 맞는 재료를 쓴다.",
             "시방서에 품질 규격이 없으면 한국산업표준(KS) 표시품 또는 이에 준하는 품질과 규격의 재료를 쓴다.",
             "기성품을 포함한 공사용 재료는 현장 반입 전에 자재승인 서류(견본 · 카탈로그 · 시험성적서 등)를 제출하여 승인을 받고, 공사감독자의 지시에 따라 품질을 확인할 수 있는 증빙서류를 낸다.",
             "재료가 현장에 들어오면 공사감독자의 인수검사를 받고, 합격한 재료는 작업 · 통행에 지장이 없는 곳에 보관하여 수시 점검이 쉽게 한다.",
             "「건설기술 진흥법」에 규정된 품질시험을 하고, 시험실 규모 · 시험장비 · 시험요원 배치기준에 따라 시험실을 운영한다.",
             "검사 또는 시험에 불합격한 재료는 지체 없이 현장 밖으로 반출한다."], marks="(n)")
    sec(p, f"{a}-3  시공 확인 및 검사")
    li(p, ["주요 공사단계가 끝났을 때나 공사감독자가 지시할 때는 시공의 정확성과 품질을 확인받는다.",
             "검사에 필요한 자료 작성 · 측량 등은 검사자의 지시에 따른다."], marks="(n)")
    sec(p, f"{a}-4  자재관리")
    li(p, ["들어온 자재는 자재수불대장으로 입 · 출고를 관리한다.", "자재는 재질별 · 규격별로 나누어 보관한다.",
             "운반 · 야적 때 변형되지 않도록 전용 운반 가설대 또는 전용 야적대를 쓴다.",
             "반입 자재는 받침목 · 지지목 등으로 지면에서 200mm 이상 띄우고 손상 · 변형 · 침수되지 않게 한다.",
             "변형을 막기 위해 받침목 높이를 맞추고 1단 적재를 기본으로 한다."], marks="(n)")


def env_std(p, bb, waste, residue):
    """환경관리 계획(공통) — bb = 장 번호"""
    ch(p, f"{bb}.0   환 경 관 리  계 획")
    sec(p, f"{bb}-1  개  요")
    para(p, "현장의 폐기물 · 유류 · 쓰레기 및 소음 · 진동 · 비산먼지 등에 의한 오염을 방지한다.")
    sec(p, f"{bb}-2  중점관리 방침")
    li(p, ["자율적인 환경관리", "쾌적한 작업환경 관리", "민원 발생 방지"], marks="(n)")
    sec(p, f"{bb}-3  실시계획")
    li(p, ["작업 전후 정리정돈 : 작업 전후 정리정돈을 철저히 하여 쾌적한 작업환경을 유지한다.",
             f"분리수거 : 분리수거를 철저히 하고, {waste}은 시공업체 부담으로 적법하게 처리하여 쓰레기를 줄인다.",
             "민원 발생 억제 : 소음 · 진동 · 비산먼지 · 냄새를 최대한 억제하여 인근 주민의 생활에 피해가 없도록 하고, 적극적으로 환경관리 활동을 한다."], marks="(n)")
    b(p, "(4) 환경관리 대책", lead=1)
    li(p, ["쓰레기 처리 : 현장 안은 금연하고 쓰레기 수거함을 둔다.",
             residue,
             "비산먼지 : 현장 안 주요 도로와 작업장 주변은 하루 2회 이상 살수차로 살수하고, 장비 이동 때 먼지가 나지 않도록 서행한다."], marks="①", lead=2)


def safety_std(p, c, kind, rules_extra=()):
    """안전관리 계획(방수 · 미장 공통) — c = 장 번호"""
    ch(p, f"{c}.0   안 전 관 리  계 획")
    sec(p, f"{c}-1  안전관리 목적")
    para(p, "공종의 특수성을 고려하여 현장의 위험요소를 미리 없애 사고를 막음으로써 귀중한 인명을 보호하고 재산 손실을 최소화하며, 작업 능률을 높여 계획 공정을 원만히 달성하는 데 목적이 있다.")
    sec(p, f"{c}-2  현장의 안전관리 방침")
    li(p, ["선 안전, 후 시공", "안전교육 내실화로 위험요인 사전 제거", "안전지도 · 점검 강화로 위험요인 사전 제거", "취약 지점 점검 · 보강"], marks="①")
    sec(p, f"{c}-3  중점 관리 사항")
    li(p, ["위험예지 활동 실시", "개인 안전보호구 착용 생활화", "신규 채용자는 특별안전교육 후 투입", "고소 및 위험지역 작업 전후 안전대책 강구",
             "각종 도구 · 시설물 · 보호장구 상태 및 사용 실태 수시 점검", "작업장 정리정돈 생활화", "관리감독자 순찰 활동 강화", "장비의 안전상태 점검",
             "장비 운전원은 유자격자 선임"], marks="①")
    sec(p, f"{c}-4  {kind}의 안전수칙")
    li(p, ["모든 작업자는 작업 때 안전보호구(안전모 · 안전화 · 보호장갑 · 작업복 등)를 반드시 착용한다.",
             "작업 전 장비 안전점검을 반드시 하고, 작업 중에도 중요 부분을 수시로 점검하여 이상이 있으면 예방조치를 한다.",
             "가동 중인 장비를 두고 자리를 뜨거나 장난 · 한눈팔기를 하지 않는다.",
             "기사는 작업자와 주위 사람 · 상황을 정확히 살피고 신호수의 신호에 따라 장비를 움직인다.",
             "높은 곳에서 작업할 때는 반드시 안전대를 착용한다."] + list(rules_extra), marks="①")


def gwanri(bk, ex, kind, waste, residue, rules_extra=(), n0=8):
    """품질 · 환경 · 안전관리 계획(방수 · 미장 계획서 공통 틀) — n0 = 품질관리 장 번호"""
    p = pg(bk, ex, "관리계획")
    quality_std(p, n0)
    env_std(p, n0 + 1, waste, residue)
    safety_std(p, n0 + 2, kind, rules_extra)
    d.done(p)
    p.after_note(["법령 이름은 현행(「건설기술 진흥법」)으로 바꿨습니다. 그 밖의 문구는 원본 계획서를 다듬은 것입니다."] if not ex else ["이 시트는 작성 예시입니다."])
    return p


# ── 검측 체크리스트(원본 틀: 검사항목 · 검사기준 · 1차/2차 합격 · 불합격 · 조치사항 · 시공자 점검 / 감리원 검측 서명) ─────
def checklist(p, ex, key, trade, sub, items, res=None, loc_ex=None, date_ex=None, code="", first=False, insp="감 리 원  검 측"):
    """res: 예시 결과 [(1차, 2차, 조치)] — 1차 · 2차는 "o"(합격) · "x"(불합격) · None"""
    if not first:
        p.__dict__.setdefault("_brk", set()).add(p.r)
    p.row([(PC, "검 측 체 크 리 스 트", {"kind": "title", "size": 17})], h=36)
    p.row([(2, "공종 CODE NO.", L), (2, I(f"{key}코드", ex=code or None, blank=code or None, align="center", size=9)), (2, trade, {"kind": "ctext", "bold": True}),
           (2, "위치 및 부위", L), (4, I(f"{key}위치", ex=loc_ex, size=9.5))], h=26)
    p.row([(2, "공종(세부공종)", L), (4, sub, {"kind": "ctext", "size": 9.5}), (2, "검 측 일 자", L),
           (4, I(f"{key}일자", ex=date_ex, fmt="date", blank=None, align="center"))], h=34)
    p.row([(5, "검  사  항  목", dict(H, rs=3)), (2, "검 사 기 준\n(시방서 · 도면 등)", dict(H, rs=3, size=9)), (4, "검  사  결  과", H),
           (1, "조치\n사항", dict(H, rs=3, size=9))], h=20)
    p.row([(2, "1  차", H), (2, "2  차", H)], h=18, start=8)
    p.row([(1, "합격", dict(H, size=9)), (1, "불합격", dict(H, size=8.5)), (1, "합격", dict(H, size=9)), (1, "불합격", dict(H, size=8.5))], h=18, start=8)
    n = len(items)
    for i, (q, std) in enumerate(items, 1):
        r = (res[i - 1] if (ex and res and i <= len(res)) else None) or (None, None, None)
        mk = lambda v, want: ("○" if v == want else None)
        p.row([(5, f"{i}. {q}", {"kind": "text", "size": 9.5}) if q else (5, I(f"{key}항목#{i}", size=9.5)), (2, std, {"kind": "ctext", "size": 9}),
               (1, I(f"{key}1o#{i}", ex=mk(r[0], "o"), align="center")), (1, I(f"{key}1x#{i}", ex=mk(r[0], "x"), align="center")),
               (1, I(f"{key}2o#{i}", ex=mk(r[1], "o"), align="center")), (1, I(f"{key}2x#{i}", ex=mk(r[1], "x"), align="center")),
               (1, I(f"{key}조치#{i}", ex=r[2], size=8))], h=34)
    for c in ("1o", "1x", "2o", "2x"):
        p.k[f"{key}{c}[]"] = f"{p.k[f'{key}{c}#1']}:{p.k[f'{key}{c}#{n}']}"
    a, x, a2, x2 = (f"COUNTA({{{key}{c}[]}})" for c in ("1o", "1x", "2o", "2x"))
    f = (f'=IF({a}+{x}=0,"","1차 합격 "&{a}&" · 불합격 "&{x}&IF({x}>0," / 2차 합격 "&{a2}&" · 불합격 "&{x2},"")'
         f'&IF({a}+{x}<{n}," · 빈칸 "&({n}-{a}-{x}),"")&"  →  "'
         f'&IF({x}=0,IF({a}+{x}<{n},"점검 중","합격"),IF(AND({x2}=0,{a2}>={x}),"합격(재검측)","보완 후 재검측")))')
    p.row([(3, "판  정", S), (9, F(f"{key}판정", f, align="left", bold=True))], h=26)
    p.row([(6, "시 공 자  점 검", H), (6, insp, H)], h=22)
    p.row([(3, "점   검", H), (3, "재 점 검", H), (3, "검   측", H), (3, "재 검 측", H)], h=20)
    p.row([(3, "(인)", {"kind": "text", "align": "right", "valign": "bottom", "size": 9}) for _ in range(4)], h=44)


def exp_checklist(g, key, n):
    cnt = lambda c: sum(1 for i in range(1, n + 1) if g.get(f"{key}{c}#{i}") not in (None, ""))
    a, x, a2, x2 = cnt("1o"), cnt("1x"), cnt("2o"), cnt("2x")
    if a + x == 0:
        return {f"{key}판정": ""}
    t = f"1차 합격 {a} · 불합격 {x}" + (f" / 2차 합격 {a2} · 불합격 {x2}" if x > 0 else "")
    t += (f" · 빈칸 {n - a - x}" if a + x < n else "") + "  →  "
    t += ("점검 중" if a + x < n else "합격") if x == 0 else ("합격(재검측)" if (x2 == 0 and a2 >= x) else "보완 후 재검측")
    return {f"{key}판정": t}


# ── 일 단위 공정표(팽이기초처럼 며칠짜리 공종) ──────────────────────────
def gongjeong_day(bk, ex, title, data, start_ex, nd=30, n=10):
    """착수일 + 공정마다 «시작 일차 · 작업 일수» → 종료 일차 · 날짜 · 일 단위 막대. data: [(공정, 시작 일차, 일수)]"""
    Wg = [22, 5, 5, 5, 11] + [2.15] * nd + [7]
    nc = len(Wg)
    C0 = 6
    p = bk.page("공정표", Wg, ex=ex, landscape=True, fit_height=1, sheetname="공정표" if not ex else "예시-공정표",
                margins=(0.35, 0.35, 0.45, 0.45))
    p.row([(nc, title, {"kind": "title", "size": 16})], h=34)
    p.row([(1, "착 수 일", L), (2, I("착수일", ex=start_ex, fmt="ymd", align="center")),
           (nd + 3, "※ 착수일과 공정마다 시작 일차 · 작업 일수를 적으면 끝나는 날과 일 단위 막대가 그려집니다.", {"kind": "note"})], h=22)
    p.row([(1, "공    정", {"kind": "head", "rs": 2}), (1, "시작\n일차", {"kind": "head", "rs": 2, "size": 8.5}),
           (1, "일수", {"kind": "head", "rs": 2, "size": 8.5}), (1, "종료\n일차", {"kind": "head", "rs": 2, "size": 8.5}),
           (1, "기    간", {"kind": "head", "rs": 2, "size": 9})]
          + [(1, k + 1, {"kind": "head", "size": 7.5}) for k in range(nd)] + [(1, "비 고", {"kind": "head", "rs": 2})], h=18)
    p.row([(1, F(f"날#{k + 1}", f'=IF(N({{착수일}})=0,"",{{착수일}}+{k})', fmt='d;;""', align="center", size=6.5), {"kind": "head"}) for k in range(nd)],
          h=14, start=C0)
    hr = p.r - 2
    r0 = p.r
    for i in range(1, n + 1):
        e = data[i - 1] if i <= len(data) else (None, None, None)
        p.row([(1, I(f"공정#{i}", ex=e[0], blank=e[0], size=9)), (1, I(f"시작#{i}", ex=e[1], fmt='0;;""', align="center", size=9)),
               (1, I(f"일수#{i}", ex=e[2], fmt='0;;""', align="center", size=9)),
               (1, F(f"종료#{i}", f'=IF(OR(N({{시작#{i}}})=0,N({{일수#{i}}})=0),"",{{시작#{i}}}+{{일수#{i}}}-1)', fmt='0;;""', align="center", size=9)),
               (1, F(f"기간#{i}", f'=IF(OR({{종료#{i}}}="",N({{착수일}})=0),"",MONTH({{착수일}}+{{시작#{i}}}-1)&"/"&DAY({{착수일}}+{{시작#{i}}}-1)&" ~ "&MONTH({{착수일}}+{{종료#{i}}}-1)&"/"&DAY({{착수일}}+{{종료#{i}}}-1))',
                     align="center", size=8.5))]
              + [(1, "", {"kind": "text"}) for _ in range(nd)] + [(1, I(f"공정비고#{i}", size=8))], h=24)
    r1 = p.r - 1
    p.k["종료[]"] = f"{p.k['종료#1']}:{p.k[f'종료#{n}']}"
    p.row([(4, "총 작업 일수", S), (1, F("총일수", '=IF(COUNT({종료[]})=0,"",MAX({종료[]}))', fmt='#,##0"일";;""', align="center", bold=True)),
           (nd, F("끝날", '=IF(OR({총일수}="",N({착수일})=0),"",{착수일}+{총일수}-1)', fmt='"마치는 날 : "yyyy"년 "m"월 "d"일";;""', align="left"), {"kind": "free"}),
           (1, "", FR)], h=22)
    p.ws.conditional_formatting.add(
        f"{CL(C0)}{r0}:{CL(C0 + nd - 1)}{r1}",
        Rule(type="expression", dxf=DifferentialStyle(fill=BAR),
             formula=[f"AND(N($B{r0})>0,N($D{r0})>=N($B{r0}),{CL(C0)}${hr}>=$B{r0},{CL(C0)}${hr}<=$D{r0})"]))
    d.done(p)
    p.after_note(["하늘색 칸 = 자동(종료 일차 · 기간 · 날짜 · 총 작업 일수). 공정 이름은 원본 차례를 미리 넣어 두었습니다."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 일정)."])
    return p


def exp_gongjeong_day(s, nd=30, n=10):
    st = s["착수일"]
    z = {f"날#{k + 1}": st + datetime.timedelta(days=k) for k in range(nd)}
    ends = []
    for i in range(1, n + 1):
        a, b_ = s.get(f"시작#{i}"), s.get(f"일수#{i}")
        if a and b_:
            e = a + b_ - 1
            ends.append(e)
            s1, s2 = st + datetime.timedelta(days=a - 1), st + datetime.timedelta(days=e - 1)
            z[f"종료#{i}"] = e
            z[f"기간#{i}"] = f"{s1.month}/{s1.day} ~ {s2.month}/{s2.day}"
        else:
            z[f"종료#{i}"] = ""
            z[f"기간#{i}"] = ""
    z["총일수"] = max(ends) if ends else ""
    z["끝날"] = st + datetime.timedelta(days=max(ends) - 1) if ends else ""
    return z


def img(p, data, rows=8, row_h=20, c1=1, c2=None, caption=None):
    """설명 그림 — 쪽 끝에서 잘리지 않게 먼저 자리를 확인하고 넣음"""
    d.keep(p, rows * row_h + (18 if caption else 0) + 4)
    d.image(p, data, rows=rows, row_h=row_h, c1=c1, c2=c2, caption=caption)


# ── 물량 기반 공정표(총수량 ÷ 1일 작업량 = 작업일수) ─────────────────────
def gongjeong_qty(bk, ex, title, data, start_ex, nd=60, n=10):
    """data: [(공종, 총수량, 단위, 1일 작업량, 시작 일차)] → 작업일수 = 올림(총수량 ÷ 1일 작업량) · 종료 일차 · 일 단위 막대 · 총 공기"""
    Wg = [17, 7, 4.5, 6.5, 5, 5, 5] + [1.75] * nd + [6]
    nc = len(Wg)
    C0 = 8
    p = bk.page("공정표", Wg, ex=ex, landscape=True, fit_height=1, sheetname="공정표" if not ex else "예시-공정표",
                margins=(0.3, 0.3, 0.45, 0.45))
    p.row([(nc, title, {"kind": "title", "size": 16})], h=34)
    p.row([(1, "착 수 일", L), (2, I("착수일", ex=start_ex, fmt="ymd", align="center")),
           (nd + 4, "※ 공종마다 총수량 · 1일 작업량 · 시작 일차를 적으면 작업일수(올림) · 끝나는 일차 · 막대 · 총 공기가 저절로 나옵니다.", {"kind": "note"})], h=22)
    p.row([(1, "공    종", {"kind": "head", "rs": 2}), (1, "총수량", {"kind": "head", "rs": 2, "size": 9}), (1, "단위", {"kind": "head", "rs": 2, "size": 8.5}),
           (1, "1일\n작업량", {"kind": "head", "rs": 2, "size": 8.5}), (1, "작업\n일수", {"kind": "head", "rs": 2, "size": 8.5}),
           (1, "시작\n일차", {"kind": "head", "rs": 2, "size": 8.5}), (1, "종료\n일차", {"kind": "head", "rs": 2, "size": 8.5})]
          + [(1, k + 1, {"kind": "head", "size": 5.5}) for k in range(nd)] + [(1, "비 고", {"kind": "head", "rs": 2})], h=16)
    p.row([(1, F(f"날#{k + 1}", f'=IF(N({{착수일}})=0,"",{{착수일}}+{k})', fmt='d;;""', align="center", size=5.5), {"kind": "head"}) for k in range(nd)],
          h=13, start=C0)
    hr = p.r - 2
    r0 = p.r
    for i in range(1, n + 1):
        e = data[i - 1] if i <= len(data) else (None,) * 5
        p.row([(1, I(f"공종#{i}", ex=e[0], blank=e[0], size=8.5)), (1, I(f"총수량#{i}", ex=e[1], fmt='#,##0.##;;""', size=8.5)),
               (1, I(f"단위#{i}", ex=e[2], blank=e[2], align="center", size=8.5)), (1, I(f"일량#{i}", ex=e[3], fmt='#,##0.##;;""', size=8.5)),
               (1, F(f"일수#{i}", f'=IF(OR(N({{총수량#{i}}})=0,N({{일량#{i}}})=0),"",ROUNDUP({{총수량#{i}}}/{{일량#{i}}},0))', fmt='0;;""', align="center", size=8.5)),
               (1, I(f"시작#{i}", ex=e[4], fmt='0;;""', align="center", size=8.5)),
               (1, F(f"종료#{i}", f'=IF(OR({{일수#{i}}}="",N({{시작#{i}}})=0),"",{{시작#{i}}}+{{일수#{i}}}-1)', fmt='0;;""', align="center", size=8.5))]
              + [(1, "", {"kind": "text"}) for _ in range(nd)] + [(1, I(f"공정비고#{i}", size=7))], h=24)
    r1 = p.r - 1
    p.k["종료[]"] = f"{p.k['종료#1']}:{p.k[f'종료#{n}']}"
    p.row([(6, "총 공기 (가장 늦게 끝나는 일차)", S), (1, F("총일수", '=IF(COUNT({종료[]})=0,"",MAX({종료[]}))', fmt='#,##0"일";;""', align="center", bold=True)),
           (nd, F("끝날", '=IF(OR({총일수}="",N({착수일})=0),"",{착수일}+{총일수}-1)', fmt='"마치는 날 : "yyyy"년 "m"월 "d"일";;""', align="left"), {"kind": "free"}),
           (1, "", FR)], h=22)
    p.ws.conditional_formatting.add(
        f"{CL(C0)}{r0}:{CL(C0 + nd - 1)}{r1}",
        Rule(type="expression", dxf=DifferentialStyle(fill=BAR),
             formula=[f"AND(N($F{r0})>0,N($G{r0})>=N($F{r0}),{CL(C0)}${hr}>=$F{r0},{CL(C0)}${hr}<=$G{r0})"]))
    d.done(p)
    p.after_note(["하늘색 칸 = 자동(작업일수 = 총수량 ÷ 1일 작업량 올림 · 종료 일차 · 날짜 · 총 공기). 공종 이름 · 단위는 원본 차례를 넣어 두었습니다."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 물량 · 일정)."])
    return p


def exp_gongjeong_qty(s, nd=60, n=10):
    import math
    st = s["착수일"]
    z = {f"날#{k + 1}": st + datetime.timedelta(days=k) for k in range(nd)}
    ends = []
    for i in range(1, n + 1):
        q, w, a = s.get(f"총수량#{i}"), s.get(f"일량#{i}"), s.get(f"시작#{i}")
        dd = math.ceil(q / w - 1e-12) if (q and w) else ""
        z[f"일수#{i}"] = dd
        if dd != "" and a:
            z[f"종료#{i}"] = a + dd - 1
            ends.append(a + dd - 1)
        else:
            z[f"종료#{i}"] = ""
    z["총일수"] = max(ends) if ends else ""
    z["끝날"] = st + datetime.timedelta(days=max(ends) - 1) if ends else ""
    return z
