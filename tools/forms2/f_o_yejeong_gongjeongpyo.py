# -*- coding: utf-8 -*-
"""예정공정표 — 원본 틀(구분 · 공종 · 금액(천원) · 보할(%) · 수량 · 소요기간 · 참조사항 · 주간 칸 · 보할 합계/누계 · 공정곡선, 총괄분 + 월간)을 새로 만든 것.
■ 원본은 주마다 보할(%)을 손으로 나눠 적고 막대·곡선을 손으로 그렸습니다 →
  공종마다 금액과 시작/끝 날짜만 적으면: 보할 = 금액 ÷ 총계 × 100 · 주간 칸 = 보할 × (그 주에 걸친 날 ÷ 공종 기간) · 막대(조건부 서식) ·
  주별 보할 합계 · 누계 · 공정곡선(S커브) · 월간(일별) 시트까지 저절로.
■ 기간 밖으로 나간 공종은 «차이» 칸에 남은 보할이 보이고 빨간 안내가 뜹니다."""
import datetime

from fb import F, I

SLUG = "o-yejeong-gongjeongpyo"
TITLE = "예정공정표"
WHERE = "orig"
PREV = [(1, "총괄분(주간) — 빈 서식"), (4, "총괄분(주간) — 작성 예시"), (5, "월간(일별) — 작성 예시"), (6, "공정곡선 — 작성 예시")]
PAGES = 3
NSHEETS = 3
MULTI = True
SHORT = ("금액·시작·끝만 적으면 보할(%)·주간 배분·막대·보할 합계·누계·공정곡선(S커브)·월간(일별) 공정표까지 저절로 채워지는 예정공정표입니다.")
NOTE = "공종마다 금액(천원)과 시작·끝 날짜만 적으세요. 보할과 주·일 배분, 막대, 누계, 곡선은 저절로 셈합니다."

D = datetime.date
NI = 20            # 공종 줄 수
NW = 30            # 주 수
ND = 31            # 월간 날 수
EX_START = D(2026, 3, 2)
EX = [("토공사", "토공", "벌개제근", 5100, 1, "식", D(2026, 3, 2), D(2026, 3, 13), ""),
      ("", "토공", "터파기", 55400, 3200, "㎥", D(2026, 3, 9), D(2026, 4, 24), ""),
      ("", "토공", "되메우기", 18300, 2100, "㎥", D(2026, 4, 6), D(2026, 5, 29), ""),
      ("", "토공", "사토", 42700, 1100, "㎥", D(2026, 3, 16), D(2026, 5, 8), "운반거리 12km"),
      ("배수공사", "암거공", "PC 암거 거치", 186000, 120, "개", D(2026, 3, 23), D(2026, 5, 22), "하루 3개"),
      ("", "관로공", "PE 이중벽관 D300", 64500, 460, "m", D(2026, 5, 11), D(2026, 6, 26), ""),
      ("", "구조물공", "집수정", 23800, 12, "개소", D(2026, 6, 1), D(2026, 7, 10), ""),
      ("포장공사", "포장공", "보도블록 포장", 38200, 850, "㎡", D(2026, 7, 6), D(2026, 8, 14), ""),
      ("", "포장공", "아스콘 포장", 29600, 1430, "㎡", D(2026, 8, 17), D(2026, 8, 28), "170ton"),
      ("부대공사", "가설공", "가설사무소", 4200, 1, "식", D(2026, 3, 2), D(2026, 9, 25), ""),
      ("", "가설공", "세륜·축중기", 9800, 1, "식", D(2026, 3, 2), D(2026, 8, 28), ""),
      ("", "운반", "각종 운반", 31000, 1, "식", D(2026, 3, 9), D(2026, 8, 28), ""),
      ("품질시험비", "", "", 1500, None, "", D(2026, 3, 2), D(2026, 9, 25), ""),
      ("안전관리비", "", "", 9300, None, "", D(2026, 3, 2), D(2026, 9, 25), "")]
EX_MONTH = D(2026, 4, 1)
BLUE = "93C5FD"
L = {"kind": "label"}
H = {"kind": "head", "size": 8.5}
WD = 'CHOOSE(WEEKDAY({d}),"일","월","화","수","목","금","토")'


def CL(c):
    from openpyxl.utils import get_column_letter
    return get_column_letter(c)


def _cf_fill(p, rng, formula, color=BLUE):
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    p.ws.conditional_formatting.add(rng, Rule(type="expression", formula=[formula], dxf=DifferentialStyle(fill=PatternFill("solid", bgColor=color))))


# ─────────────────────────── 총괄분(주간) ───────────────────────────
WW = [7.5, 8, 12, 9, 6, 7, 4.6, 6.4, 6.4, 5, 12] + [4.4] * NW + [5.6, 5.6]
C_W0 = 12                      # 첫 주 칸(L)
C_SUM, C_DIFF = C_W0 + NW, C_W0 + NW + 1


def weekly(bk, ex):
    N = len(WW)
    p = bk.page("총괄분", WW, ex=ex, fit_height=1, landscape=True, sheetname="총괄분" if not ex else "예시-총괄분",
                margins=(0.3, 0.3, 0.4, 0.4))
    p.row([(N, "예      정      공      정      표   (총괄분)", {"kind": "title", "size": 18})], h=34)
    p.row([(2, "공 사 명 :", {"kind": "free", "bold": True, "align": "right"}), (9, I("공사명", ex="가나지구 배수로 정비공사", size=11, bold=True)),
           (5, "첫 주 시작일", {"kind": "free", "bold": True, "align": "right", "size": 9}),
           (4, I("시작주", ex=EX_START, fmt="ymd", align="center")),
           (4, "(그 주 월요일)", {"kind": "free", "size": 8.5}),
           (N - 24, F("기간표시", '=IF(N({시작주})=0,"","공정표 기간 : "&YEAR({시작주})&". "&MONTH({시작주})&". "&DAY({시작주})&". ~ "'
                                   f'&YEAR({{시작주}}+{NW * 7 - 1})&". "&MONTH({{시작주}}+{NW * 7 - 1})&". "&DAY({{시작주}}+{NW * 7 - 1})&". ({NW}주)")',
                     align="left", size=9.5), {"border": False})], h=24)
    p.gap(4)
    r1 = p.r
    heads = [("구 분", 1), ("공 종", 1), ("세부 공종", 1), ("금액\n(천원)", 1), ("보할\n(%)", 1), ("수 량", 1), ("단위", 1), ("시 작", 1), ("끝", 1),
             ("기간\n(일)", 1), ("참 조 사 항", 1)]
    cells = [(s, t, dict(H, rs=3)) for t, s in heads]
    for k in range(1, NW + 1):
        cells.append((1, F(f"월#{k}", (f'=IF(N({{시작주}})=0,"",MONTH({{시작주}}+{7 * (k - 1)})&"월")' if k == 1 else
                                      f'=IF(N({{시작주}})=0,"",IF(MONTH({{시작주}}+{7 * (k - 1)})<>MONTH({{시작주}}+{7 * (k - 2)}),MONTH({{시작주}}+{7 * (k - 1)})&"월",""))'),
                         align="center", size=8.5, bold=True), H))
    cells += [(1, "반영\n(%)", dict(H, rs=3)), (1, "차이", dict(H, rs=3))]
    p.row(cells, h=16)
    cells = [(1, F(f"주#{k}", f'=IF(N({{시작주}})=0,"",INT((DAY({{시작주}}+{7 * (k - 1)})-1)/7)+1&"주")', align="center", size=8), H) for k in range(1, NW + 1)]
    p.row(cells, h=14)
    cells = [(1, F(f"주시작#{k}", f'=IF(N({{시작주}})=0,"",{{시작주}}+{7 * (k - 1)})', fmt='m"/"d', align="center", size=7.5), H) for k in range(1, NW + 1)]
    p.row(cells, h=14)
    rdate = p.r - 1
    # 공종 줄(두 줄씩: 위 = 막대, 아래 = 주간 보할)
    first = p.r
    tot = f"$D${first + 2 * NI}"                       # 총계 칸(아래에서 만듦)
    for i in range(1, NI + 1):
        e = EX[i - 1] if (ex and i <= len(EX)) else (None,) * 9
        ra, rb = p.r, p.r + 1
        cells = [(1, I(f"구분#{i}", ex=e[0] or None, align="center", size=9), {"rs": 2}),
                 (1, I(f"공종#{i}", ex=e[1] or None, align="center", size=9), {"rs": 2}),
                 (1, I(f"세부#{i}", ex=e[2] or None, size=9), {"rs": 2}),
                 (1, I(f"금액#{i}", ex=e[3], fmt="int", size=9), {"rs": 2}),
                 (1, F(f"보할#{i}", f'=IF(OR(N({{금액#{i}}})=0,N({tot})=0),"",{{금액#{i}}}/{tot}*100)', fmt='0.00;;""', align="center", size=9), {"rs": 2}),
                 (1, I(f"수량#{i}", ex=e[4], fmt="num", size=9), {"rs": 2}),
                 (1, I(f"단위#{i}", ex=e[5] or None, align="center", size=9), {"rs": 2}),
                 (1, I(f"시작#{i}", ex=e[6], fmt="md", align="center", size=9), {"rs": 2}),
                 (1, I(f"끝#{i}", ex=e[7], fmt="md", align="center", size=9), {"rs": 2}),
                 (1, F(f"기간#{i}", f'=IF(OR(N({{시작#{i}}})=0,N({{끝#{i}}})=0),"",{{끝#{i}}}-{{시작#{i}}}+1)', fmt='0;;""', align="center", size=9), {"rs": 2}),
                 (1, I(f"참조#{i}", ex=e[8] or None, size=8), {"rs": 2})]
        cells += [(1, "", {"kind": "text"})] * NW
        cells += [(1, F(f"반영#{i}", f'=IF({{보할#{i}}}="","",SUM({CL(C_W0)}{rb}:{CL(C_W0 + NW - 1)}{rb}))', fmt='0.00;;""', align="center", size=8.5),
                   {"rs": 2}),
                  (1, F(f"차이#{i}", f'=IF({{보할#{i}}}="","",ROUND({{보할#{i}}}-{{반영#{i}}},2))', fmt='0.00;[Red]-0.00;""', align="center", size=8.5),
                   {"rs": 2})]
        p.row(cells, h=11)
        cells = []
        for k in range(1, NW + 1):
            wc = f"{CL(C_W0 + k - 1)}${rdate}"
            cells.append((1, F(f"w#{i}#{k}", f'=IF(OR({{보할#{i}}}="",{{기간#{i}}}="",N({wc})=0),"",MAX(0,MIN({{끝#{i}}},{wc}+6)-MAX({{시작#{i}}},{wc})+1)'
                                               f'/{{기간#{i}}}*{{보할#{i}}})', fmt='[<0.005]"";0.00', align="center", size=7, ink=False), {"kind": "text"}))
        p.row(cells, h=13, start=C_W0)
    last = p.r - 1
    p.row([(3, "총        계", {"kind": "sum"}), (1, F("금액계", f'=SUM(D{first}:D{last})', fmt="int", size=9, bold=True)),
           (1, F("보할계", f'=IF(N({{금액계}})=0,"",SUM(E{first}:E{last}))', fmt='0.00;;""', align="center", size=9, bold=True)),
           (6, "", {"kind": "sum"})] + [(1, "", {"kind": "sum"})] * NW + [(2, "", {"kind": "sum"})], h=18)
    assert p.r - 1 == first + 2 * NI, (p.r - 1, first + 2 * NI)
    cells = [(10, "보    할  (%)", dict(L, rs=2)), (1, "합  계", {"kind": "ctext", "size": 9})]
    for k in range(1, NW + 1):
        c = CL(C_W0 + k - 1)
        cells.append((1, F(f"합#{k}", f'=IF(N({{금액계}})=0,"",SUM({c}{first}:{c}{last}))', fmt='0.00;;""', align="center", size=7.5)))
    cells.append((2, "", {"kind": "text"}))
    p.row(cells, h=15)
    cells = [(1, "누  계", {"kind": "ctext", "size": 9})]
    for k in range(1, NW + 1):
        prev = f"+{{누#{k - 1}}}" if k > 1 else ""
        cells.append((1, F(f"누#{k}", f'=IF(N({{금액계}})=0,"",{{합#{k}}}{prev})', fmt='0.00;;""', align="center", size=7.5, bold=True)))
    cells.append((2, "", {"kind": "text"}))
    p.row(cells, h=15, start=11)
    rcum = p.r - 1
    dr = f"{CL(C_DIFF)}{first}:{CL(C_DIFF)}{last}"
    p.row([(N, F("경고", f'=IF(COUNTIF({dr},">0.004")+COUNTIF({dr},"<-0.004")>0,'
                        f'"⚠ 공정표 기간 밖으로 나간 공종이 있습니다(«차이» 칸) — 첫 주 시작일이나 공종 날짜를 확인하세요.",'
                        f'IF(SUMPRODUCT((H{first}:H{last}>0)*(I{first}:I{last}>0)*(I{first}:I{last}<H{first}:H{last}))>0,"⚠ 끝 날짜가 시작보다 빠른 공종이 있습니다.",""))',
                 align="left"), {"kind": "warn"})], h=16)
    # 막대 — 아래 줄 값이 있으면 위 줄을 칠함
    w1, w2 = CL(C_W0), CL(C_W0 + NW - 1)
    _cf_fill(p, f"{w1}{first}:{w2}{last}", f'AND(MOD(ROW()-{first},2)=0,N({w1}{first + 1})>0)')
    p.end_print()
    p.after_note(["흰 칸(구분 · 공종 · 금액 · 수량 · 시작 · 끝 · 참조사항)만 적으면 하늘색 칸과 주간 칸 · 막대 · 합계 · 누계가 저절로 채워집니다.",
                  "주간 칸 = 보할 × (그 주(월~일)에 걸친 날 수 ÷ 공종 기간). 휴무 주를 빼려면 공종을 둘로 나눠 적으세요. 공정곡선 시트는 누계를 그립니다."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 공사)."])
    p.k["_first"], p.k["_last"], p.k["_rcum"], p.k["_rdate"] = first, last, rcum, rdate
    return p


# ─────────────────────────── 월간(일별) ───────────────────────────
MW = [7.5, 8, 12, 6, 6] + [3.1] * ND + [6, 6]
C_D0 = 6


def monthly(bk, ex):
    N = len(MW)
    p = bk.page("월간", MW, ex=ex, fit_height=1, landscape=True, sheetname="월간" if not ex else "예시-월간", margins=(0.3, 0.3, 0.4, 0.4))
    p.row([(N, "예      정      공      정      표   (월간)", {"kind": "title", "size": 18})], h=34)
    p.row([(2, "공 사 명 :", {"kind": "free", "bold": True, "align": "right"}),
           (12, F("공사명표시", '=IF({총괄분!공사명}="","",{총괄분!공사명})', align="left", size=11, bold=True)),
           (6, "이 달 (1일)", {"kind": "free", "bold": True, "align": "right", "size": 9}),
           (6, I("월초", ex=EX_MONTH, fmt='yyyy"년" m"월"', align="center")),
           (N - 26, "← 달의 1일을 적으세요(총괄분 공종·보할을 날마다 나눠 보여 줌)", {"kind": "free", "size": 8.5})], h=24)
    p.gap(4)
    cells = [(1, "구 분", dict(H, rs=2)), (1, "공 종", dict(H, rs=2)), (1, "세부 공종", dict(H, rs=2)), (1, "보할\n(%)", dict(H, rs=2)),
             (1, "지난달\n까지", dict(H, rs=2))]
    for d in range(1, ND + 1):
        cells.append((1, F(f"일#{d}", f'=IF(N({{월초}})=0,"",IF(MONTH({{월초}}+{d - 1})<>MONTH({{월초}}),"",{{월초}}+{d - 1}))', fmt="d", align="center", size=8), H))
    cells += [(1, "이 달", dict(H, rs=2)), (1, "누 계", dict(H, rs=2))]
    p.row(cells, h=16)
    rdate = p.r - 1
    cells = []
    for d in range(1, ND + 1):
        dc = f"{{일#{d}}}"
        cells.append((1, F(f"요일#{d}", f'=IF(N({dc})=0,"",' + WD.format(d=dc) + ')', align="center", size=8), H))
    p.row(cells, h=14, start=C_D0)
    first = p.r
    for i in range(1, NI + 1):
        g = f"{{총괄분!구분#{i}}}"
        cells = [(1, F(f"구분#{i}", f'=IF({g}="","",{g})', align="center", size=8.5)),
                 (1, F(f"공종#{i}", f'=IF({{총괄분!공종#{i}}}="","",{{총괄분!공종#{i}}})', align="center", size=8.5)),
                 (1, F(f"세부#{i}", f'=IF({{총괄분!세부#{i}}}="","",{{총괄분!세부#{i}}})', align="left", size=8.5)),
                 (1, F(f"보할#{i}", f'=IF({{총괄분!보할#{i}}}="","",{{총괄분!보할#{i}}})', fmt='0.00;;""', align="center", size=8.5)),
                 (1, F(f"지난#{i}", f'=IF(OR({{보할#{i}}}="",{{총괄분!기간#{i}}}="",N({{월초}})=0),"",MAX(0,MIN({{총괄분!끝#{i}}},{{월초}}-1)-{{총괄분!시작#{i}}}+1)'
                                    f'/{{총괄분!기간#{i}}}*{{보할#{i}}})', fmt='0.00;;""', align="center", size=8))]
        for d in range(1, ND + 1):
            dc = f"{CL(C_D0 + d - 1)}${rdate}"
            cells.append((1, F(f"d#{i}#{d}", f'=IF(OR({{보할#{i}}}="",{{총괄분!기간#{i}}}="",N({dc})=0),"",IF(AND({dc}>={{총괄분!시작#{i}}},{dc}<={{총괄분!끝#{i}}}),'
                                                f'{{보할#{i}}}/{{총괄분!기간#{i}}},""))', fmt='[<0.005]"";0.00', align="center", size=6.5, ink=False), {"kind": "text"}))
        cells += [(1, F(f"이달#{i}", f'=IF({{보할#{i}}}="","",SUM({CL(C_D0)}{p.r}:{CL(C_D0 + ND - 1)}{p.r}))', fmt='0.00;;""', align="center", size=8)),
                  (1, F(f"누적#{i}", f'=IF({{보할#{i}}}="","",{{지난#{i}}}+{{이달#{i}}})', fmt='0.00;;""', align="center", size=8))]
        p.row(cells, h=15)
    last = p.r - 1
    cells = [(3, "일 별  합 계", {"kind": "sum"}), (1, F("보할계", f'=IF(COUNT(D{first}:D{last})=0,"",SUM(D{first}:D{last}))', fmt='0.00;;""', align="center", size=8.5, bold=True)),
             (1, F("지난계", f'=IF(COUNT(E{first}:E{last})=0,"",SUM(E{first}:E{last}))', fmt='0.00;;""', align="center", size=8.5, bold=True))]
    for d in range(1, ND + 1):
        c = CL(C_D0 + d - 1)
        cells.append((1, F(f"일합#{d}", f'=IF(COUNT({c}{first}:{c}{last})=0,"",SUM({c}{first}:{c}{last}))', fmt='0.00;;""', align="center", size=6.5)))
    cells += [(1, F("이달계", f'=IF(COUNT({CL(C_D0 + ND)}{first}:{CL(C_D0 + ND)}{last})=0,"",SUM({CL(C_D0 + ND)}{first}:{CL(C_D0 + ND)}{last}))', fmt='0.00;;""',
                    align="center", size=8.5, bold=True)),
              (1, F("누적계", f'=IF(COUNT({CL(C_D0 + ND + 1)}{first}:{CL(C_D0 + ND + 1)}{last})=0,"",SUM({CL(C_D0 + ND + 1)}{first}:{CL(C_D0 + ND + 1)}{last}))',
                    fmt='0.00;;""', align="center", size=8.5, bold=True))]
    p.row(cells, h=16)
    cells = [(5, "누   계 (%)", {"kind": "sum"})]
    for d in range(1, ND + 1):
        cells.append((1, F(f"일누#{d}", f'=IF({{일합#{d}}}="","",N({{지난계}})+SUM({{일합#1}}:{{일합#{d}}}))', fmt='0.0;;""', align="center", size=6.5, bold=True)))
    cells += [(2, "", {"kind": "sum"})]
    p.row(cells, h=16)
    d1, d2 = CL(C_D0), CL(C_D0 + ND - 1)
    _cf_fill(p, f"{d1}{first}:{d2}{last}", f'N({d1}{first})>0')
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import Font
    from openpyxl.styles.differential import DifferentialStyle
    p.ws.conditional_formatting.add(f"{d1}{rdate}:{d2}{rdate + 1}", Rule(type="expression", formula=[f'AND(ISNUMBER({d1}${rdate}),WEEKDAY({d1}${rdate})=1)'],
                                    dxf=DifferentialStyle(font=Font(color="DC2626", bold=True))))
    p.end_print()
    p.after_note(["이 시트는 적는 칸이 «이 달 (1일)» 하나뿐입니다 — 공종·보할·날짜는 총괄분에서 가져와 날마다 나눕니다(보할 ÷ 공종 기간)."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 공사)."])
    return p


# ─────────────────────────── 공정곡선 ───────────────────────────
def curve(bk, ex, wk):
    W = [4] * 30
    p = bk.page("공정곡선", W, ex=ex, fit_height=1, landscape=True, sheetname="공정곡선" if not ex else "예시-공정곡선", margins=(0.4, 0.4, 0.5, 0.5))
    p.row([(30, "공 정 곡 선 (S-커브)", {"kind": "title", "size": 18})], h=34)
    p.row([(4, "공 사 명 :", {"kind": "free", "bold": True, "align": "right"}),
           (26, F("공사명표시", '=IF({총괄분!공사명}="","",{총괄분!공사명})', align="left", size=11, bold=True), {"border": False})], h=24)
    for _ in range(26):
        p.gap(16)
    p.end_print()
    from openpyxl.chart import LineChart, Reference
    ch = LineChart()
    ch.title = None
    ch.y_axis.title = "누계 공정률(%)"
    ch.x_axis.title = "주 (시작일)"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 100
    ch.y_axis.majorUnit = 10
    ch.height, ch.width = 11.5, 22
    ch.x_axis.tickLblSkip = 2
    ch.legend = None
    ws = wk.ws
    data = Reference(ws, min_col=C_W0, max_col=C_W0 + NW - 1, min_row=wk.k["_rcum"], max_row=wk.k["_rcum"])
    cats = Reference(ws, min_col=C_W0, max_col=C_W0 + NW - 1, min_row=wk.k["_rdate"], max_row=wk.k["_rdate"])
    ch.add_data(data, from_rows=True, titles_from_data=False)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties.line.solidFill = "DC2626"
    s.graphicalProperties.line.width = 28000
    s.graphicalProperties.line.dashStyle = "dash"
    s.smooth = False
    ch.x_axis.number_format = 'm"/"d'
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    p.ws.add_chart(ch, "B4")
    p.after_note(["총괄분 «누계» 줄을 그린 곡선입니다. 총괄분만 적으면 저절로 바뀝니다."] if not ex else ["이 시트는 작성 예시입니다(가상의 공사)."])
    return p


def draw(bk, ex):
    wk = weekly(bk, ex)
    monthly(bk, ex)
    curve(bk, ex, wk)


def expect(pages):
    wk, mo = pages["총괄분"], pages["월간"]
    dd = datetime.timedelta
    s0 = wk["시작주"]
    items = []
    for i in range(1, NI + 1):
        items.append((wk.get(f"금액#{i}"), wk.get(f"시작#{i}"), wk.get(f"끝#{i}")))
    tot = sum(a or 0 for a, _, _ in items)
    e = {"금액계": tot, "보할계": 100.0,
         "기간표시": f"공정표 기간 : {s0.year}. {s0.month}. {s0.day}. ~ {(s0 + dd(days=NW * 7 - 1)).year}. {(s0 + dd(days=NW * 7 - 1)).month}. "
                  f"{(s0 + dd(days=NW * 7 - 1)).day}. ({NW}주)"}
    wsum = [0.0] * NW
    for k in range(1, NW + 1):
        ws_ = s0 + dd(days=7 * (k - 1))
        prev = s0 + dd(days=7 * (k - 2))
        e[f"월#{k}"] = f"{ws_.month}월" if (k == 1 or ws_.month != prev.month) else ""
        e[f"주#{k}"] = f"{(ws_.day - 1) // 7 + 1}주"
        e[f"주시작#{k}"] = ws_
    bh = {}
    for i, (a, st, en) in enumerate(items, 1):
        if not a:
            e[f"보할#{i}"] = ""
            e[f"기간#{i}"] = "" if not (st and en) else (en - st).days + 1
            e[f"반영#{i}"] = ""
            e[f"차이#{i}"] = ""
            for k in range(1, NW + 1):
                e[f"w#{i}#{k}"] = ""
            continue
        b = a / tot * 100
        bh[i] = b
        dur = (en - st).days + 1
        e[f"보할#{i}"] = b
        e[f"기간#{i}"] = dur
        s = 0.0
        for k in range(1, NW + 1):
            ws_ = s0 + dd(days=7 * (k - 1))
            ov = max(0, (min(en, ws_ + dd(days=6)) - max(st, ws_)).days + 1)
            v = ov / dur * b
            e[f"w#{i}#{k}"] = v
            s += v
            wsum[k - 1] += v
        e[f"반영#{i}"] = s
        e[f"차이#{i}"] = round(b - s, 2)
    cum = 0.0
    for k in range(1, NW + 1):
        e[f"합#{k}"] = wsum[k - 1]
        cum += wsum[k - 1]
        e[f"누#{k}"] = cum
    e["경고"] = "⚠ 공정표 기간 밖으로 나간 공종이 있습니다(«차이» 칸) — 첫 주 시작일이나 공종 날짜를 확인하세요." if any(
        e[f"차이#{i}"] not in ("", 0) for i in bh) else ""
    # 월간
    m0 = mo["월초"]
    m = {"공사명표시": wk["공사명"]}
    days = []
    for d in range(1, ND + 1):
        x = m0 + dd(days=d - 1)
        days.append(x if x.month == m0.month else None)
        m[f"일#{d}"] = x if x.month == m0.month else ""
        m[f"요일#{d}"] = "" if x.month != m0.month else "월화수목금토일"[x.weekday()]
    dsum = [None] * ND
    tot_prev = 0.0
    tot_mon = 0.0
    anyb = False
    for i in range(1, NI + 1):
        a, st, en = items[i - 1]
        g = wk.get(f"구분#{i}")
        m[f"구분#{i}"] = g or ""
        m[f"공종#{i}"] = wk.get(f"공종#{i}") or ""
        m[f"세부#{i}"] = wk.get(f"세부#{i}") or ""
        if i not in bh:
            m[f"보할#{i}"] = ""
            m[f"지난#{i}"] = ""
            m[f"이달#{i}"] = ""
            m[f"누적#{i}"] = ""
            for d in range(1, ND + 1):
                m[f"d#{i}#{d}"] = ""
            continue
        anyb = True
        b = bh[i]
        dur = (en - st).days + 1
        prev = max(0, (min(en, m0 - dd(days=1)) - st).days + 1) / dur * b
        m[f"보할#{i}"] = b
        m[f"지난#{i}"] = prev
        mon = 0.0
        for d in range(1, ND + 1):
            x = days[d - 1]
            v = (b / dur) if (x and st <= x <= en) else ""
            m[f"d#{i}#{d}"] = v
            if v != "":
                mon += v
                dsum[d - 1] = (dsum[d - 1] or 0) + v
        m[f"이달#{i}"] = mon
        m[f"누적#{i}"] = prev + mon
        tot_prev += prev
        tot_mon += mon
    m["보할계"] = sum(bh.values()) if anyb else ""
    m["지난계"] = tot_prev if anyb else ""
    m["이달계"] = tot_mon if anyb else ""
    m["누적계"] = tot_prev + tot_mon if anyb else ""
    run = tot_prev
    for d in range(1, ND + 1):
        v = dsum[d - 1]
        m[f"일합#{d}"] = "" if v is None else v
        if v is None:
            m[f"일누#{d}"] = ""
        else:
            run += v
            m[f"일누#{d}"] = run
    return {"총괄분": e, "월간": m, "공정곡선": {"공사명표시": wk["공사명"]}}
