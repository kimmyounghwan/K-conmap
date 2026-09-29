# -*- coding: utf-8 -*-
"""자재 공급원 승인 요청 · 결과통보 — 원본 틀([별지 제37호 서식], 요청서 + 검토결과 통보서 한 장)을 새로 만든 것.
■ □ 칸은 목록에서 ■ 를 고르면 됩니다(글자 칸이라 엑셀·한셀 어디서나 똑같이 보임).
■ 수식: 공종·판정을 목록에서 고르면 «건축□ 기계□ 토목■ 기타□» 모양으로 저절로 · 첨부 서류 수 = ■ 개수 ·
        판정이 «부적합» 이면 빨간 안내 · 통보서의 공사명·품명은 요청서에서 따라옴."""
import datetime

from fb import F, I

SLUG = "o-jajae-seungin"
TITLE = "자재 공급원 승인요청서(시공자용)"
WHERE = "orig"
PREV = [(1, "빈 서식"), (2, "작성 예시")]
PAGES = 1
NSHEETS = 1
SHORT = ("시공자가 감리(건설사업관리기술인)에게 자재 공급원 승인을 요청하고 검토결과를 통보받는 양식([별지 제37호 서식]). "
         "공종·판정을 목록에서 고르면 □/■ 표시가 저절로 되고, 첨부 서류 수를 세어 줍니다.")
NOTE = "□ 칸은 목록에서 ■ 를 고르세요. 공종과 판정은 목록에서 고르면 «건축□ 기계□ 토목■ 기타□» 모양으로 저절로 적힙니다."

D = datetime.date
W = [7.6] * 12
N = 12
L = {"kind": "label"}
FR = {"kind": "free"}
ATT = ["인증서", "사업자등록증", "CATALOG", "공장등록증", "시험성적서", "납품실적", "견본", "기타"]
ATT_EX = {"인증서", "사업자등록증", "CATALOG", "시험성적서", "납품실적"}
KS = ["KS", "비KS", "환경표지", "GR"]


def _chk(key, on):
    return I(key, ex="■" if on else "□", blank="□", dv="체크", align="center", size=11)


def _show(key, opts):
    return F(key + "표시", "=" + '&"   "&'.join([f'"{o}"&IF({{{key}}}="{o}","■","□")' for o in opts]), align="left")


def _sign(p, lab, key, exv):
    p.row([(5, "", FR), (3, lab, {"kind": "free", "align": "right", "bold": True}), (3, I(key, ex=exv, align="center"), {"border": False}),
           (1, "(인)", {"kind": "cfree", "size": 9})], h=22)


def draw(bk, ex):
    p = bk.page("승인요청", W, ex=ex, fit_height=1, sheetname="승인요청" if not ex else "예시-승인요청", margins=(0.5, 0.5, 0.45, 0.45))
    p.dv_list("체크", ["□", "■"])
    p.dv_list("공종", ["건축", "기계", "토목", "기타"])
    p.dv_list("판정", ["적합", "조건부적합", "부적합"])
    p.row([(N, "[별지 제37호 서식]", {"kind": "free", "size": 9})], h=16)
    p.title("자재 공급원 승인 요청 · 결과통보 내용", h=34)
    p.gap(4)
    p.row([(2, "문서번호", L), (4, I("문서번호", ex="예시건설 제2026-045호", size=9.5)), (2, "수    신", L),
           (4, I("수신", ex="책임건설사업관리기술인", blank="책임건설사업관리기술인", align="center", size=9.5))], h=26)
    p.row([(2, "공 사 명", L), (10, I("공사명", ex="가나지구 배수로 정비공사"))], h=26)
    p.row([(2, "공    종", L), (2, I("공종", ex="토목", dv="공종", align="center")), (8, _show("공종", ["건축", "기계", "토목", "기타"]))], h=26)
    p.row([(2, "품    명", L), (4, I("품명", ex="PC 암거")), (2, "규    격", L), (4, I("규격", ex="1.5m × 1.5m × 1.0m", size=9.5))], h=26)
    p.row([(2, "제조회사명", L), (10, I("제조사", ex="예시콘크리트(주)"))], h=26)
    cells = [(2, "KS · 녹색제품\n유무", {"kind": "label", "size": 9})]
    for j, k in enumerate(KS, 1):
        cells += [(1, _chk(f"KS#{j}", k in ("KS",))), (1, k, {"kind": "text", "size": 9, "align": "center"})]
    cells.append((2, "", {"kind": "text"}))
    p.row(cells, h=28)
    p.row([(2, "시공자 의견", L), (10, I("시공자의견", ex="설계도서의 규격(KS F 4413)과 같고 납품 실적이 충분하여 공급원으로 승인을 요청합니다.", size=9.5))], h=48)
    r0 = p.r
    for half in (ATT[:4], ATT[4:]):
        cells = [] if p.r > r0 else [(2, "첨    부", {"kind": "label", "rs": 2})]
        for k in half:
            cells += [(1, _chk(f"첨부:{k}", k in ATT_EX)), (1, k, {"kind": "text", "size": 8.5, "align": "center"})]
        if p.r == r0:
            cells.append((2, F("첨부수", '=IF(COUNTIF({첨부[]},"■")=0,"","첨부 "&COUNTIF({첨부[]},"■")&"종")', align="center", bold=True),
                           {"rs": 2}))
        p.row(cells, h=24)
    c1, c2 = p.k[f"첨부:{ATT[0]}"], p.k[f"첨부:{ATT[-1]}"]
    p.k["첨부[]"] = f"{c1}:{c2}"
    # 좁은 칸의 긴 이름(사업자등록증 · CATALOG · 공장등록증 …)은 두 줄로 쪼개지 말고 칸에 맞춰 줄임
    from openpyxl.styles import Alignment
    for rr in range(r0 - 2, p.r):
        for cc in range(1, N + 1):
            v = p.ws.cell(row=rr, column=cc).value
            if isinstance(v, str) and v in ATT + KS:
                p.ws.cell(row=rr, column=cc).alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
    p.row([(2, "특기사항", L), (10, I("특기사항1", ex="시험성적서는 공인시험기관 발행분(최근 6개월 이내)", size=9.5))], h=34)
    p.gap(6)
    _sign(p, "담  당  자", "담당자", "이공무")
    _sign(p, "현장대리인", "현장대리인", "홍길동")
    p.gap(4)
    p.row([(N, "상기 자재에 대한 검토를 요청하오니 결과를 통보하여 주시기 바랍니다.", {"kind": "cfree"})], h=22)
    p.row([(N, I("요청일", ex=D(2026, 4, 2), fmt="date", blank="20    년      월      일", align="center"), {"border": False})], h=22)
    p.gap(8)
    from openpyxl.styles import Border, Side
    for c in range(1, N + 1):
        p.ws.cell(row=p.r - 1, column=c).border = Border(bottom=Side(style="medium", color="111827"))
    p.gap(6)
    p.row([(N, "자재승인 검토결과 통보서", {"kind": "title", "size": 15})], h=30)
    p.row([(2, "문서번호", L), (4, I("통보번호", ex="가나감리 제2026-112호", size=9.5)), (2, "수    신", L),
           (4, I("통보수신", ex="현장대리인", blank="현장대리인", align="center"))], h=26)
    p.row([(2, "공사명 · 품명", L), (10, F("통보대상", '=IF({공사명}="","",{공사명}&IF({품명}="",""," — "&{품명}&IF({규격}="",""," ("&{규격}&")")))',
                                              align="left", size=9.5))], h=26)
    p.row([(2, "검토 의견", L), (10, I("검토의견", ex="제출 서류 확인 결과 설계도서 규격과 같고 품질관리 체계가 적정함.", size=9.5))], h=48)
    p.row([(2, "판    정", L), (2, I("판정", ex="적합", dv="판정", align="center")), (8, _show("판정", ["적합", "조건부적합", "부적합"]))], h=26)
    p.row([(N, F("판정안내", '=IF({판정}="부적합","⚠ 부적합 — 다른 공급원으로 다시 요청해야 합니다.",IF({판정}="조건부적합","※ 조건부적합 — 조건(특기사항)을 지켜 쓰고 결과를 보고합니다.",""))',
                  align="left"), {"kind": "warn"})], h=16)
    p.row([(2, "특기사항", L), (10, "1. 공장방문 후 검토결과를 통보할 경우 해당 공장방문 검사 Check List 첨부\n"
                                  "2. 자체시험 및 외부 의뢰시험을 실시하는 경우 시험 결과치 기록 또는 시험성적서 첨부", {"kind": "text", "size": 9})], h=36)
    p.gap(4)
    p.row([(N, "상기 검토요청에 대한 검토결과를 통보합니다.", {"kind": "cfree"})], h=22)
    p.row([(N, I("통보일", ex=D(2026, 4, 6), fmt="date", blank="20    년      월      일", align="center"), {"border": False})], h=22)
    _sign(p, "담당 건설사업관리기술인", "담당감리", "김철수")
    _sign(p, "책임건설사업관리기술인", "책임감리", "박책임")
    p.end_print()
    p.after_note(["□ 칸은 목록에서 ■ 를 고르세요. 하늘색 칸 = 자동(공종·판정 표시 · 첨부 수 · 통보서 공사명)."] if not ex else
                 ["이 시트는 작성 예시입니다(가상의 공사 · 회사 · 사람)."])
    return p


def expect(inp):
    def show(v, opts):
        return "   ".join(f"{o}{'■' if v == o else '□'}" for o in opts)
    n = sum(1 for k in ATT if inp.get(f"첨부:{k}") == "■")
    return {"공종표시": show(inp["공종"], ["건축", "기계", "토목", "기타"]),
            "판정표시": show(inp["판정"], ["적합", "조건부적합", "부적합"]),
            "첨부수": f"첨부 {n}종" if n else "",
            "통보대상": f"{inp['공사명']} — {inp['품명']} ({inp['규격']})",
            "판정안내": ""}
