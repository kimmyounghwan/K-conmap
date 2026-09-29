# -*- coding: utf-8 -*-
"""발파 작업계획서 — 원본 틀(결재 · 1 발파 일시 · 2 발파 작업 계획 · 3 화약류 운반 · 4 발파관리 · 5 운반/발파 신고 · 6 책임자 ·
7 발파 위치도 · 8 작업 방법 및 안전대책)을 새로 만든 것.
■ 원본 위치도는 실제 현장 도면이라 쓰지 않고, 작성 예시에는 가상의 현장을 그린 그림(img/o-balpa-jageop/make_map.py)을 넣음.
  그 그림 속 «출입금지 · 위험 발파중» 간판 그림은 원본 그대로.
■ 수식: 발파 요일 · 폭약 사용량 = 천공수 × 공당장약량 · 뇌관 = 천공수(바꿔 적어도 됨) · 천공 연장 = 천공수 × 천공심도 ·
        신고일이 발파일보다 늦으면 안내 · 책임자 칸."""
import datetime
import os

from fb import F, I

SLUG = "o-balpa-jageop"
TITLE = "발파 작업계획서"
WHERE = "orig"
PREV = [(1, "빈 서식 첫 쪽"), (3, "작성 예시 첫 쪽"), (4, "작성 예시 — 위치도 · 작업 방법")]
PAGES = 2
NSHEETS = 1
SHORT = ("발파 작업계획서(발파 일시 · 천공/장약 계획 · 화약류 운반 · 발파관리 · 신고 · 위치도 · 작업 방법과 안전대책). "
         "천공수와 공당장약량을 적으면 폭약·뇌관 사용량과 천공 연장을 셈해 줍니다.")
NOTE = "발파 위치도 칸에는 현장 도면을 붙여 넣으세요(작성 예시 그림은 가상의 현장). 신고 기한은 관할 경찰서 안내를 따르세요."

D = datetime.date
T = datetime.time
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "img", SLUG)
W = [15, 12, 14, 13, 14, 12, 12]
N = 7
from openpyxl.styles import PatternFill
LO = {"kind": "label", "fill": PatternFill("solid", fgColor="FED7AA")}   # 원본의 주황 머리칸(노란 칸 = 기준값과 헷갈리지 않게 옅은 주황)
LG = {"kind": "label", "fill": PatternFill("solid", fgColor="D9F99D")}   # 원본의 연두 머리칸
FR = {"kind": "free"}
H1 = {"kind": "free", "bold": True, "size": 12}
WKD = 'CHOOSE(WEEKDAY({d}),"일","월","화","수","목","금","토")'
KO = "월화수목금토일"
METHOD = [("천공", "암반 노면을 최대한 평탄하게 하여 드릴의 전도사고가 발생하지 않도록 하며, 설계에 맞는 공간격 · 저항선 · 천공장으로 천공한다."),
          ("화약운반", "경찰서에서 화약류 운반필증을 수령한 후 화약고(저장소)로 가서 당일 필요한 화약류를 운반한다."),
          ("장약", "화약과 뇌관으로 전폭약을 만들고, 천공 구멍에 전폭약과 천공장에 알맞은 화약을 장약한다."),
          ("전색", "전색물(메지)로 장약 후 남은 빈 공간을 밀실하게 채우고 다진다."),
          ("발파", "뇌관선을 서로 연결하고 뇌관 누락 · 연결 누락을 확인하기 위해 도통시험을 한 후 발파모선과 뇌관선을 연결한다. "
                  "그 후 다시 도통시험을 하고 모선을 발파기에 연결한 후 사이렌 등으로 발파 경고를 한 뒤 발파한다."),
          ("발파결과 확인", "발파 후 15분이 지난 뒤 발파 지역에 들어가 불발 화약과 지반 상태를 확인하고 이상 유무를 판단한 후 후속 공정을 진행한다.")]
SAFETY = [("폭발위험", "화약류 운반 · 장약 · 발파 때 절대 흡연 금지, 악천후에는 작업 중지 / 화약과 뇌관은 따로 보관 / 불발 화약 확인 철저"),
          ("비산위험", "발파암 비산 방호 매트를 견고하게 충분히 덮어 비산을 막음 / 발파 때 근로자는 충분한 안전거리 확보"),
          ("낙하위험", "발파매트 이동 · 설치 · 해체 때 굴착기 후크(볼트)를 체결하여 매트 낙하 방지")]


def draw(bk, ex):
    p = bk.page("발파 작업계획서", W, ex=ex, fit_height=0, sheetname="발파 작업계획서" if not ex else "예시-발파 작업계획서", margins=(0.5, 0.5, 0.45, 0.45))
    # 제목 + 결재
    p.row([(4, "발 파  작 업  계 획 서", {"kind": "title", "size": 20, "rs": 3}),
           (1, "입  안", {"kind": "head"}), (1, "심  사", {"kind": "head"}), (1, "결  정", {"kind": "head"})], h=20)
    p.row([(1, "", {"kind": "text"})] * 3, h=34, start=5)
    p.row([(1, "     /", {"kind": "ctext", "size": 9})] * 3, h=16, start=5)
    p.gap(6)
    p.row([(2, "1. 발파 일시 :", H1), (2, I("발파일", ex=D(2026, 6, 10), fmt="date", align="left"), {"border": False}),
           (1, F("요일", '=IF(N({발파일})=0,"","("&' + WKD.format(d="{발파일}") + '&")")', align="left"), {"border": False}),
           (1, I("발파시각", ex=T(11, 30), fmt='hh:mm', align="center", blank=None)), (1, "(시 : 분)", {"kind": "free", "size": 8.5})], h=24)
    p.gap(4)
    p.row([(N, "2. 발파 작업 계획", H1)], h=22)
    p.row([(1, "발    파", dict(LO, rs=3)), (1, "천공수", LG), (2, I("천공수", ex=480, fmt='#,##0" 공"', align="center")),
           (1, "천공심도", LG), (2, I("심도", ex=3.0, fmt='0.0#" m"', align="center"))], h=24)
    p.row([(1, "공 간격", LG), (2, I("공간격", ex=1.2, fmt='0.0#" m"', align="center")), (1, "공당장약량", LG),
           (2, I("공당", ex=0.37, fmt='0.0##" kg"', align="center"))], h=24)
    p.row([(1, "발파횟수", LG), (2, I("횟수", ex=2, fmt='0" 회"', align="center")), (1, "천공 연장", LG),
           (2, F("연장", '=IF(OR(N({천공수})=0,N({심도})=0),"",{천공수}*{심도})', fmt='#,##0.0#" m"', align="center"))], h=24)
    p.row([(1, "폭약사용량", LO), (1, "폭    약", LG),
           (2, F("폭약", '=IF(OR(N({천공수})=0,N({공당})=0),"",ROUND({천공수}*{공당},1))', fmt='#,##0.0" kg"', align="center", bold=True)),
           (1, "뇌    관", LG), (2, F("뇌관", '=IF(N({천공수})=0,"",{천공수})', fmt='#,##0" 개"', align="center", bold=True))], h=24)
    p.row([(N, "※ 폭약 = 천공수 × 공당장약량 · 뇌관 = 천공수(한 공에 뇌관 1개 기준 — 다르면 그 칸에 직접 적으세요) · 천공 연장 = 천공수 × 천공심도",
            {"kind": "note"})], h=16)
    p.gap(4)
    p.row([(N, "3. 화약류 운반", H1)], h=22)
    for lab, k, exv in (("시    간", "운반시간", "08:00 ~ 09:00"), ("책 임 자", "운반책임자", "김화약"), ("차량번호", "차량번호", "가나 12가 3456"),
                        ("경 유 지", "경유지", "화약고 → 가나로 → 현장 정문")):
        p.row([(1, lab, LO), (6, I(k, ex=exv if ex else None, size=9.5))], h=22)
    p.gap(4)
    p.row([(N, "4. 발파관리", H1)], h=22)
    for lab, k, exv in (("업 체 명", "업체명", "예시발파(주)"), ("관련기술사", "기술사", "이기술 (화약류관리기술사)"), ("연 락 처", "연락처", "02-0000-0000")):
        p.row([(1, lab, LO), (6, I(k, ex=exv if ex else None, size=9.5))], h=22)
    p.gap(4)
    p.row([(N, "5. 화약류운반(발파)신고", H1)], h=22)
    p.row([(1, "운반신고", dict(LO, rs=2)), (1, "신 고 일", LG), (5, I("운반신고일", ex=D(2026, 6, 9) if ex else None, fmt="date", align="left"))], h=22)
    p.row([(1, "신고기한", LG), (5, I("운반기한", ex="운반 개시 전(관할 경찰서 안내에 따름)" if ex else None, size=9.5))], h=22)
    p.row([(1, "발파신고", dict(LO, rs=2)), (1, "신 고 일", LG), (5, I("발파신고일", ex=D(2026, 6, 8) if ex else None, fmt="date", align="left"))], h=22)
    p.row([(1, "신고기한", LG), (5, I("발파기한", ex="발파 전(관할 경찰서 안내에 따름)" if ex else None, size=9.5))], h=22)
    p.row([(2, "화약운반(도착)신고기관", LG), (5, I("신고기관", ex="가나경찰서" if ex else None))], h=24)
    p.row([(N, F("신고안내", '=IF(N({발파일})=0,"",IF(OR(N({운반신고일})>{발파일},N({발파신고일})>{발파일}),"⚠ 신고일이 발파일보다 늦습니다 — 날짜를 확인하세요.",""))',
                 align="left"), {"kind": "warn"})], h=16)
    p.row([(2, "6. 현장 발파 관련 책임자 :", H1), (2, I("책임자", ex="박발파" if ex else None, align="center"), {"border": False}),
           (3, I("자격", ex="화약류관리보안책임자 1급", blank="화약류관리보안책임자 1급", align="left", size=9.5), {"border": False})], h=26)
    from openpyxl.worksheet.pagebreak import Break
    p.ws.row_breaks.append(Break(id=p.r - 1))
    p.row([(N, "7. 발파 위치도 (발파 지역, 신호수 배치, 접근금지 간판 설치 장소 등)", H1)], h=26)
    r0 = p.r
    for _ in range(15):
        p.gap(18)
    if ex:
        with open(os.path.join(IMG, "example_map.png"), "rb") as f:
            p.place_image(f.read(), 1, N, r0, p.r - 1, pad=4)
    else:
        p.put(r0, 1, N, "(발파 위치도를 붙이는 곳 — 발파 지역 · 경계 반경 · 신호수 배치 · 접근금지 간판 · 대피 장소를 표시)", kind="ctext", size=9, rs=15)
    p.gap(6)
    p.row([(N, "8. 작업 방법 및 안전대책", H1)], h=22)
    Y = LO
    p.row([(1, "작 업 명", Y), (6, "세 부 작 업 내 용", Y)], h=22)
    from fb import text_lines
    for j, (nm, t) in enumerate(METHOD, 1):
        n = text_lines(t, sum(W[1:]) * 10 / 9.5)
        p.row([(1, nm, {"kind": "ctext", "size": 9.5}), (6, I(f"방법#{j}", ex=t, blank=t, size=9.5))], h=max(24, n * 14 + 8))
    p.row([(1, "안 전 대 책", Y), (6, "예 방 대 책", Y)], h=22)
    for j, (nm, t) in enumerate(SAFETY, 1):
        n = text_lines(t, sum(W[1:]) * 10 / 9.5)
        p.row([(1, nm, {"kind": "ctext", "size": 9.5}), (6, I(f"대책#{j}", ex=t, blank=t, size=9.5))], h=max(24, n * 14 + 8))
    p.end_print()
    p.after_note(["하늘색 칸 = 자동(요일 · 천공 연장 · 폭약 · 뇌관). 뇌관 수가 천공수와 다르면(한 공에 2개 등) 그 칸에 직접 적으세요.",
                  "작업 방법 · 안전대책 글은 원본 뜻 그대로 두고 오타만 고쳤습니다(«비례위험» → «비산위험», «한 수» → «한 후»). 현장에 맞게 고쳐 쓰세요."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 공사 · 회사 · 사람 · 위치도). 간판 그림만 원본 계획서의 것입니다."])
    return p


def expect(inp):
    d = inp["발파일"]
    n, dep, q = inp["천공수"], inp["심도"], inp["공당"]
    late = any(inp.get(k) and inp[k] > d for k in ("운반신고일", "발파신고일"))
    return {"요일": f"({KO[d.weekday()]})", "연장": n * dep, "폭약": round(n * q + 1e-9, 1), "뇌관": n,
            "신고안내": "⚠ 신고일이 발파일보다 늦습니다 — 날짜를 확인하세요." if late else ""}
