# -*- coding: utf-8 -*-
"""회의 공정표(예정공정표) — 원본 틀(공종 · 내역 · 수량 · 단위 · 날짜 칸 막대 · 수량 전주/금주/누계 · 잔여량 · 비고, 가로 한 장)을 새로 만든 것.
■ 원본은 막대를 «선 그림» 으로 그려 손으로 옮겨야 했습니다 → 계획·실적 시작/끝 날짜를 적으면 날짜 칸이 저절로 칠해집니다
  (계획 = 파랑, 실적 = 노랑 · 실적 끝을 비우면 기준일까지).
■ 수식: 기준일 → 전주·금주·내주 21일 날짜 · 누계 = 전주까지 + 금주 · 잔여량 = 수량 − 누계 · 진척률 · 상태(완료·지연·착수 지연·진행·예정)."""
import datetime

from fb import F, I

SLUG = "o-hoeui-gongjeongpyo"
TITLE = "회의 공정표"
WHERE = "orig"
PREV = [(1, "빈 서식"), (2, "작성 예시")]
PAGES = 1
NSHEETS = 1
SHORT = ("주간 공정회의 때 쓰는 예정공정표(전주·금주·내주 3주)입니다. 계획·실적 날짜를 적으면 막대가 저절로 칠해지고, "
         "누계·잔여량·진척률과 상태(완료·지연·착수 지연)를 셈해 줍니다.")
NOTE = "기준일(회의일)을 적으면 3주 날짜가 붙습니다. 공종마다 계획·실적 시작/끝을 적으면 날짜 칸이 칠해집니다(실적 끝을 비우면 기준일까지)."

D = datetime.date
NDAY = 21
ROWS = 16
W = [8, 14, 7, 5, 4.6, 6.4, 6.4] + [3.0] * NDAY + [7, 6.4, 7, 7, 7, 6.6, 10]
N = len(W)
C0 = 8                     # 첫 날짜 칸(H)
H = {"kind": "head", "size": 9}
L = {"kind": "label"}
BLUE, YEL = "60A5FA", "FACC15"
WD = "월화수목금토일"
BASE = D(2026, 5, 13)
EX = [("가시설", "흙막이판 설치", 180, "m", (D(2026, 4, 27), D(2026, 5, 8)), (D(2026, 4, 28), D(2026, 5, 9)), 180, None, ""),
      ("토공", "터파기", 3200, "㎥", (D(2026, 5, 4), D(2026, 5, 12)), (D(2026, 5, 4), None), 1450, 520, "우천 1일"),
      ("배수공", "PC 암거 거치", 120, "개", (D(2026, 5, 6), D(2026, 5, 22)), (D(2026, 5, 7), None), 38, 16, ""),
      ("토공", "되메우기", 2100, "㎥", (D(2026, 5, 11), D(2026, 5, 27)), (D(2026, 5, 12), None), None, 240, ""),
      ("배수공", "PE 이중벽관 D300", 460, "m", (D(2026, 5, 18), D(2026, 5, 29)), (None, None), None, None, "자재 반입 5/15"),
      ("포장공", "보도블록 포장", 850, "㎡", (D(2026, 5, 25), D(2026, 6, 5)), (None, None), None, None, "")]


def _abs(ref):
    import re
    m = re.match(r"([A-Z]+)(\d+)", ref)
    return f"${m.group(1)}${m.group(2)}"


def _per(tag, a, b):
    s = "{금주시작}"
    return f'="{tag}"&IF(N({s})=0,"(   /   ~   /   )","("&MONTH({s}+{a})&"/"&DAY({s}+{a})&"~"&MONTH({s}+{b})&"/"&DAY({s}+{b})&")")'


def draw(bk, ex):
    p = bk.page("공정표", W, ex=ex, fit_height=1, landscape=True, sheetname="공정표" if not ex else "예시-공정표",
                margins=(0.4, 0.4, 0.45, 0.45))
    p.row([(N, "예  정  공  정  표", {"kind": "title", "size": 17})], h=32)
    p.row([(2, "▣ 공 사 명 :", {"kind": "free", "bold": True}), (11, I("공사명", ex="가나지구 배수로 정비공사")),
           (4, "", {"kind": "free"}),
           (4, "기준일(회의일)", {"kind": "free", "bold": True, "align": "right", "size": 9}),
           (4, I("기준일", ex=BASE, fmt="ymd", align="center")),
           (3, "금주 시작", {"kind": "free", "bold": True, "align": "right", "size": 9}),
           (3, F("금주시작", '=IF(N({기준일})=0,"",{기준일}-WEEKDAY({기준일},2)+1)', fmt="ymd", align="center")),
           (1, "계획", {"kind": "cfree", "size": 9, "bold": True}), (1, "", {"border": False}), (1, "실적", {"kind": "cfree", "size": 9, "bold": True}),
           (1, "", {"border": False})], h=22)
    from openpyxl.styles import PatternFill
    lr = p.r - 1
    p.ws.cell(row=lr, column=N - 2).fill = PatternFill("solid", fgColor=BLUE)
    p.ws.cell(row=lr, column=N).fill = PatternFill("solid", fgColor=YEL)
    p.gap(4)
    # 머리 3줄
    p.row([(1, "공  종", dict(H, rs=3)), (1, "내    역", dict(H, rs=3)), (1, "수 량", dict(H, rs=3)), (1, "단위", dict(H, rs=3)),
           (1, "구분", dict(H, rs=3)), (2, "기    간", H),
           (7, F("전주표시", _per("전 주", -7, -1)), H), (7, F("금주표시", _per("금 주", 0, 6)), H), (7, F("내주표시", _per("내 주", 7, 13)), H),
           (3, "수     량", H), (1, "잔여량", dict(H, rs=3)), (1, "진척률", dict(H, rs=3)), (1, "상  태", dict(H, rs=3)),
           (1, "비   고", dict(H, rs=3))], h=20)
    cells = [(1, "시작", dict(H, rs=2)), (1, "끝", dict(H, rs=2))]
    for k in range(1, NDAY + 1):
        cells.append((1, F(f"일#{k}", f'=IF(N({{금주시작}})=0,"",{{금주시작}}+{k - 8})', fmt='d', align="center", size=8.5), H))
    cells += [(1, "전주까지", dict(H, rs=2)), (1, "금 주", dict(H, rs=2)), (1, "누 계", dict(H, rs=2))]
    p.row(cells, h=17)
    rd = p.r - 1
    cells = []
    for k in range(NDAY):
        w = WD[k % 7]
        cells.append((1, w, dict(H, size=8, ink="DC2626" if w == "일" else ("2563EB" if w == "토" else None))))
    p.row(cells, h=15)
    first = p.r
    for i in range(1, ROWS + 1):
        e = EX[i - 1] if (ex and i <= len(EX)) else (None, None, None, None, (None, None), (None, None), None, None, None)
        cells = [(1, I(f"공종#{i}", ex=e[0], align="center", size=9), {"rs": 2}), (1, I(f"내역#{i}", ex=e[1], size=9), {"rs": 2}),
                 (1, I(f"수량#{i}", ex=e[2], fmt="num", size=9), {"rs": 2}), (1, I(f"단위#{i}", ex=e[3], align="center", size=9), {"rs": 2}),
                 (1, "계획", {"kind": "ctext", "size": 8.5}),
                 (1, I(f"계획시작#{i}", ex=e[4][0], fmt="md", align="center", size=9)), (1, I(f"계획끝#{i}", ex=e[4][1], fmt="md", align="center", size=9))]
        cells += [(1, "", {"kind": "text"})] * NDAY
        cells += [(1, I(f"전주까지#{i}", ex=e[6], fmt="num", size=9), {"rs": 2}), (1, I(f"금주#{i}", ex=e[7], fmt="num", size=9), {"rs": 2}),
                  (1, F(f"누계#{i}", f'=IF(AND({{전주까지#{i}}}="",{{금주#{i}}}=""),"",ROUND(N({{전주까지#{i}}})+N({{금주#{i}}}),3))',
                        fmt="num", size=9), {"rs": 2}),
                  (1, F(f"잔여#{i}", f'=IF(OR(N({{수량#{i}}})=0,{{누계#{i}}}=""),"",ROUND({{수량#{i}}}-{{누계#{i}}},3))', fmt="num", size=9), {"rs": 2}),
                  (1, F(f"진척#{i}", f'=IF(OR(N({{수량#{i}}})=0,{{누계#{i}}}=""),"",{{누계#{i}}}/{{수량#{i}}})', fmt="pct1", align="center", size=9),
                   {"rs": 2}),
                  (1, F(f"상태#{i}", f'=IF(AND({{공종#{i}}}="",{{내역#{i}}}=""),"",'
                                     f'IF(AND(N({{수량#{i}}})>0,N({{누계#{i}}})>=N({{수량#{i}}})),"완료",'
                                     f'IF(N({{기준일}})=0,"",'
                                     f'IF(AND(N({{계획끝#{i}}})>0,{{기준일}}>{{계획끝#{i}}}),"지연",'
                                     f'IF(N({{실적시작#{i}}})>0,"진행",'
                                     f'IF(AND(N({{계획시작#{i}}})>0,{{기준일}}>{{계획시작#{i}}}),"착수 지연",IF(N({{계획시작#{i}}})>0,"예정","")))))))',
                        align="center", size=9, bold=True), {"rs": 2}),
                  (1, I(f"비고#{i}", ex=e[8] or None, size=8.5), {"rs": 2})]
        p.row(cells, h=16)
        cells = [(1, "실적", {"kind": "ctext", "size": 8.5}),
                 (1, I(f"실적시작#{i}", ex=e[5][0], fmt="md", align="center", size=9)), (1, I(f"실적끝#{i}", ex=e[5][1], fmt="md", align="center", size=9))]
        cells += [(1, "", {"kind": "text"})] * NDAY
        p.row(cells, h=16)
    last = p.r - 1
    p.end_print()
    # 막대(조건부 서식) — 계획 파랑 · 실적 노랑(끝을 비우면 기준일까지)
    from openpyxl.formatting.rule import Rule
    from openpyxl.styles import Font
    from openpyxl.styles.differential import DifferentialStyle
    from openpyxl.utils import get_column_letter as CL
    c1, c2 = CL(C0), CL(C0 + NDAY - 1)
    base = _abs(p.k["기준일"])
    dref = f"{c1}${rd}"
    rng = f"{c1}{first}:{c2}{last}"
    f_plan = (f'AND($E{first}="계획",ISNUMBER({dref}),N($F{first})>0,{dref}>=$F{first},'
              f'{dref}<=IF(N($G{first})=0,$F{first},$G{first}))')
    f_act = (f'AND($E{first}="실적",ISNUMBER({dref}),N($F{first})>0,{dref}>=$F{first},'
             f'{dref}<=IF(N($G{first})=0,IF(N({base})=0,$F{first},MAX($F{first},{base})),$G{first}))')
    p.ws.conditional_formatting.add(rng, Rule(type="expression", formula=[f_plan], dxf=DifferentialStyle(fill=PatternFill("solid", bgColor=BLUE))))
    p.ws.conditional_formatting.add(rng, Rule(type="expression", formula=[f_act], dxf=DifferentialStyle(fill=PatternFill("solid", bgColor=YEL))))
    p.ws.conditional_formatting.add(f"{c1}{rd}:{c2}{rd}", Rule(type="expression", formula=[f'AND(ISNUMBER({c1}{rd}),{c1}{rd}={base})'],
                                    dxf=DifferentialStyle(fill=PatternFill("solid", bgColor="FDE68A"), font=Font(bold=True, color="B45309"))))
    st1, st2 = p.k["상태#1"], p.k[f"상태#{ROWS}"]
    p.ws.conditional_formatting.add(f"{st1}:{st2}", Rule(type="expression", formula=[f'OR({st1}="지연",{st1}="착수 지연")'],
                                    dxf=DifferentialStyle(font=Font(color="DC2626", bold=True))))
    p.ws.conditional_formatting.add(f"{st1}:{st2}", Rule(type="expression", formula=[f'{st1}="완료"'],
                                    dxf=DifferentialStyle(font=Font(color="15803D", bold=True))))
    p.after_note(["기준일을 적으면 전주 월요일부터 3주(21일) 날짜가 붙고, 기준일 칸이 노랗게 표시됩니다. 주가 월요일에 시작하지 않으면 «금주 시작» 칸에 날짜를 직접 적으세요.",
                  "공종마다 계획·실적의 시작/끝(월/일)을 적으면 날짜 칸이 칠해집니다 — 계획 파랑 · 실적 노랑(실적 끝을 비우면 기준일까지 칠함).",
                  "수량: 누계 = 전주까지 + 금주 · 잔여량 = 수량 − 누계 · 진척률 = 누계 ÷ 수량. 상태: 누계가 수량에 닿으면 완료 · 계획 끝이 지났는데 남았으면 지연 · "
                  "계획 시작이 지났는데 실적 시작이 없으면 착수 지연."]
                 if not ex else ["이 시트는 작성 예시입니다(가상의 공사). 기준일 2026. 5. 13.(수)"])
    return p


def expect(inp):
    b = inp["기준일"]
    s = b - datetime.timedelta(days=b.weekday())
    dd = datetime.timedelta

    def per(tag, a, c):
        x, y = s + dd(days=a), s + dd(days=c)
        return f"{tag}({x.month}/{x.day}~{y.month}/{y.day})"
    e = {"금주시작": s, "전주표시": per("전 주", -7, -1), "금주표시": per("금 주", 0, 6), "내주표시": per("내 주", 7, 13)}
    for k in range(1, NDAY + 1):
        e[f"일#{k}"] = s + dd(days=k - 8)
    for i in range(1, ROWS + 1):
        q, a, w = inp.get(f"수량#{i}"), inp.get(f"전주까지#{i}"), inp.get(f"금주#{i}")
        cum = "" if a is None and w is None else round((a or 0) + (w or 0), 3)
        e[f"누계#{i}"] = cum
        e[f"잔여#{i}"] = "" if (not q or cum == "") else round(q - cum, 3)
        e[f"진척#{i}"] = "" if (not q or cum == "") else cum / q
        if not inp.get(f"공종#{i}") and not inp.get(f"내역#{i}"):
            stt = ""
        elif q and cum != "" and cum >= q:
            stt = "완료"
        else:
            pe, ps, xs = inp.get(f"계획끝#{i}"), inp.get(f"계획시작#{i}"), inp.get(f"실적시작#{i}")
            if pe and b > pe:
                stt = "지연"
            elif xs:
                stt = "진행"
            elif ps and b > ps:
                stt = "착수 지연"
            elif ps:
                stt = "예정"
            else:
                stt = ""
        e[f"상태#{i}"] = stt
    return e
