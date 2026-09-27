# -*- coding: utf-8 -*-
"""
📑 내역 대조(/jeoksan/auto) 예시 내역서 — K-건설맵이 만든 «가상» 내역서 (2026-09-27)
    python tools/내역_예시.py  →  web/public/jeoksan/내역_예시.xlsx

실제 공사가 아닙니다. 예시 도면 둘(토목_예시.dxf · 마감_예시.dxf)과 대조해 보라고
일부러 몇 줄은 도면과 같게, 몇 줄은 다르게, 한 줄은 도면에 없게 적었습니다.
시험: node tools/시험_내역대조.mjs
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', 'web', 'public', 'jeoksan', '내역_예시.xlsx')

줄 = [
    # 품명, 규격, 단위, 수량, 단가
    ('이형철근', 'SD400 D13', 'ton', 0.049, 950000),
    ('이형철근', 'SD400 D16', 'ton', 0.075, 940000),
    ('이형철근', 'SD400 D22', 'ton', 0.077, 930000),
    ('레미콘', '25-24-150', 'm3', 18.4, 98000),
    ('합판거푸집', '3회', 'm2', 60.0, 21000),
    ('기초잡석', '', 'm3', 6.2, 35000),
    ('PE 이중벽관', 'D300', 'm', 64.5, 42000),
    ('토사 깎기', '', 'm3', 589, 3500),
    ('리핑암 깎기', '', 'm3', 110, 12000),
    ('발파암 깎기', '', 'm3', 16, 25000),
    ('노체 쌓기', '', 'm3', 572, 2800),
    ('노상 쌓기', '', 'm3', 150, 3100),
    ('표토제거', '', 'm3', 239, 1900),
    ('떼붙임', '', 'm2', 636, 4200),
    ('L형 측구', '', 'm', 150, 38000),
    ('빗물받이', '', '개소', 4, 450000),
    ('가로등 설치', '', '개', 4, 1850000),
    ('교통안전표지판', '', '개', 2, 320000),
    ('바닥 마감(F1) 비닐타일', '', 'm2', 78, 18000),
    ('벽 마감(W1) 수성페인트', '', 'm2', 130, 9000),
    ('천장 마감(C1) 텍스', '', 'm2', 78, 21000),
    ('걸레받이(B1)', '', 'm', 48.2, 6500),
    ('창호 AW1 알루미늄창', '1800×1500', '개소', 3, 620000),
    ('창호 WD1 목재문', '900×2100', '개소', 2, 380000),
]

wb = Workbook()
ws = wb.active
ws.title = '내역서'
thin = Side(style='thin', color='999999')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ws['A1'] = '가상 공사 내역서 (예시 — 실제 공사가 아닙니다 · K-건설맵)'
ws['A1'].font = Font(bold=True, size=13)
머리 = ['번호', '품명', '규격', '단위', '수량', '단가', '금액', '비고']
for j, h in enumerate(머리, 1):
    c = ws.cell(row=3, column=j, value=h)
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='EDF2F7')
    c.alignment = Alignment(horizontal='center')
    c.border = box
for i, (품명, 규격, 단위, 수량, 단가) in enumerate(줄, 1):
    r = 3 + i
    vals = [i, 품명, 규격, 단위, 수량, 단가, None, '']
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.border = box
    ws.cell(row=r, column=7, value='=ROUND(E%d*F%d,0)' % (r, r)).border = box
    ws.cell(row=r, column=5).number_format = '#,##0.000'
    ws.cell(row=r, column=6).number_format = '#,##0'
    ws.cell(row=r, column=7).number_format = '#,##0'
last = 3 + len(줄)
ws.cell(row=last + 1, column=2, value='합계').font = Font(bold=True)
ws.cell(row=last + 1, column=7, value='=SUM(G4:G%d)' % last).number_format = '#,##0'
for col, w in zip('ABCDEFGH', [6, 28, 14, 7, 12, 12, 14, 16]):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet('일위대가')
ws2['A1'] = '일위대가 (예시 — 대조에서는 처음에 꺼 둡니다)'
for j, h in enumerate(['호표', '품명', '규격', '단위', '수량', '단가'], 1):
    ws2.cell(row=3, column=j, value=h)
ws2.append([1, '보통인부', '', '인', 0.12, 169000])
ws2.append([2, '특별인부', '', '인', 0.05, 220000])

wb.save(OUT)
print('저장:', os.path.abspath(OUT), os.path.getsize(OUT), 'bytes')
