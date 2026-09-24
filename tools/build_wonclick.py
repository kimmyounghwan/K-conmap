# -*- coding: utf-8 -*-
"""K-건설맵 공사서류 원클릭 — 한 번 입력하면 착공부터 하자까지 서류가 채워지는 엑셀 (매크로 없음)

■ 우리 것으로 새로 만든 것입니다 (다른 기관 프로그램을 고친 것이 아님).
■ 매크로를 쓰지 않습니다 — 인터넷에서 받은 매크로 파일은 윈도우가 막기 때문입니다. 수식만 씁니다.
■ 법령 요율(보증금률·지체상금률·하자기간)은 넣어 두지 않습니다. 계약서에 적힌 값을 이용자가 넣습니다.
■ 출력: wonclick.xlsx (빈 프로그램) + wonclick_meta.json (사이트 도구가 입력칸 주소를 읽음)

쓰는 법: python build_wonclick.py <출력폴더> [시험값.json]
"""
import json, math, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties

FONT = '맑은 고딕'
NC = 12                 # 서식 칸 수 (B..M)
C0 = 2                  # 첫 칸 = B
COLW = 6.4
INPUT_FILL = PatternFill('solid', fgColor='FFF7D6')
CALC_FILL = PatternFill('solid', fgColor='EEF2F7')
HEAD_FILL = PatternFill('solid', fgColor='F2F2F2')
SEC_FILL = PatternFill('solid', fgColor='1F3A5F')
thin = Side(style='thin', color='808080')
med = Side(style='medium', color='404040')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MARK = '&"맑은 고딕,보통"&8&K808080K-건설맵 · k-conmap.com'


def f(size=11, bold=False, color='000000'):
    return Font(name=FONT, size=size, bold=bold, color=color)


# ───────────────────────── 입력 칸 ─────────────────────────
# (키, 이름, 종류, 예시, 도움말, [고를 것])
SECTIONS = [
    ('1. 공사', [
        ('공사명', '공 사 명', 'text', '○○동 ○○시설 개보수공사', '계약서에 적힌 이름 그대로', None),
        ('현장위치', '현장 위치', 'text', '○○시 ○○구 ○○로 00', '공사 현장 주소', None),
        ('공사개요', '공사 개요', 'text', '내부 마감 개보수, 설비 교체', '주요 공사 내용 한 줄 (공사대장에 들어갑니다)', None),
        ('계약번호', '계약 번호', 'text', '제2026-000호', '없으면 비워 두세요', None),
        ('계약방법', '계약 방법', 'list', '제한경쟁', '고르세요', ['일반경쟁', '제한경쟁', '지명경쟁', '수의계약', '민간계약', '기타']),
        ('계약일', '계 약 일', 'date', '2026-03-02', '예: 2026-03-02', None),
        ('착공일', '착 공 일', 'date', '2026-03-05', '', None),
        ('준공기한', '준공 기한 (계약상)', 'date', '2026-06-30', '계약서의 준공 예정일', None),
        ('계약금액', '계약 금액 (원, 부가세 포함)', 'money', 123456000, '숫자만 — 예: 123456000', None),
    ]),
    ('2. 발주기관 (발주자)', [
        ('발주기관', '발주기관 이름', 'text', '○○시', '민간 공사면 건축주·회사 이름', None),
        ('받는사람', '서류 받는 사람 (귀하 앞)', 'text', '○○시장', '서류 끝 「○○○ 귀하」 에 들어갈 이름·직함', None),
        ('발주주소', '발주기관 주소', 'text', '', '', None),
        ('감독자', '공사 감독자 (직·성명)', 'text', '', '예: 시설팀장 ○○○', None),
        ('검사자', '준공 검사자 (직·성명)', 'text', '', '', None),
        ('입회자', '검사 입회자 (직·성명)', 'text', '', '', None),
        ('설계자', '설계자', 'text', '', '설계사무소 (공사대장용, 없으면 비움)', None),
        ('감리자', '감리자', 'text', '', '감리단 (없으면 비움)', None),
    ]),
    ('3. 계약상대자 (우리 회사)', [
        ('상호', '상    호', 'text', '○○건설㈜', '', None),
        ('대표자', '대 표 자', 'text', '', '', None),
        ('사업자번호', '사업자등록번호', 'text', '', '000-00-00000', None),
        ('업체주소', '주    소', 'text', '', '', None),
        ('업체전화', '전화번호', 'text', '', '', None),
        ('업종', '업종 (면허)', 'text', '', '예: 실내건축공사업', None),
        ('은행', '은 행 명', 'text', '', '대금 받을 계좌 (계약 때 등록한 계좌)', None),
        ('예금주', '예 금 주', 'text', '', '', None),
        ('계좌', '계좌번호', 'text', '', '', None),
    ]),
    ('4. 현장대리인', [
        ('대리인', '성    명', 'text', '', '', None),
        ('대리인생일', '생년월일', 'text', '', '예: 1980. 1. 1.', None),
        ('대리인자격', '자격 종목·등급', 'text', '', '예: 건축기사 / 건설기술인 중급', None),
        ('대리인번호', '자격(등록) 번호', 'text', '', '', None),
        ('대리인전화', '연 락 처', 'text', '', '', None),
    ]),
    ('5. 보증·요율 — 계약서에 적힌 값을 넣으세요', [
        ('계약보증률', '계약보증금률 (%)', 'num', '', '예: 10 → 10% (공고서·계약서 확인)', None),
        ('계약보증방법', '계약보증 방법', 'list', '', '', ['보증서(증권)', '현금', '지급각서', '면제']),
        ('하자보증률', '하자보수보증금률 (%)', 'num', '', '계약서 확인 (공종마다 다름)', None),
        ('하자년', '하자담보책임기간 (년)', 'num', '', '계약서 확인 (공종마다 다름)', None),
        ('하자보증방법', '하자보수보증 방법', 'list', '', '', ['보증서(증권)', '현금', '지급각서', '면제']),
        ('지체율', '지체상금률 (1,000분의)', 'num', '', '예: 0.5 → 1,000분의 0.5 (계약서 확인)', None),
    ]),
    ('6. 기성 (이번 회차)', [
        ('기성회차', '기성 회차', 'num', '', '예: 1', None),
        ('기성시작', '기성 대상 기간 — 시작', 'date', '', '', None),
        ('기성끝', '기성 대상 기간 — 끝', 'date', '', '', None),
        ('기성요청일', '기성검사 요청일', 'date', '', '', None),
        ('기성청구일', '기성금 청구일', 'date', '', '비우면 기성검사 요청일', None),
        ('전회누계', '전회까지 기성 누계 (원)', 'money', '', '처음이면 0', None),
        ('금회기성', '이번 회 기성금액 (원)', 'money', '', '', None),
        ('금회선금정산', '이번 회 선금 정산액 (원)', 'money', '', '선금이 없으면 0', None),
    ]),
    ('7. 공기 연장 (필요할 때만)', [
        ('연장신청일', '연장 신청일', 'date', '', '', None),
        ('연장기한', '연장 요청 준공기한', 'date', '', '', None),
        ('연장사유', '연장 사유', 'text', '', '짧게 — 자세한 것은 붙임으로', None),
    ]),
    ('8. 준공', [
        ('준공일', '실제 준공일', 'date', '', '', None),
        ('검사요청일', '준공검사 요청일', 'date', '', '비우면 실제 준공일', None),
        ('검사일', '준공검사일', 'date', '', '', None),
        ('검사결과', '준공검사 결과', 'list', '', '', ['합격', '보완 요구']),
        ('변경기한', '승인된 변경 준공기한', 'date', '', '공기 연장이 승인됐을 때만 (지체일수 계산용)', None),
        ('정산금액', '최종 계약금액 (정산, 원)', 'money', '', '비우면 계약 금액', None),
        ('기성지급누계', '기성 지급 누계 (원)', 'money', '', '준공 전까지 받은 기성금 합계', None),
        ('준공선금정산', '선금 잔액 정산액 (원)', 'money', '', '선금이 없으면 0', None),
        ('지체상금입력', '지체상금 (원)', 'money', '', '비우면 자동 계산(최종 계약금액×요율×지체일수). 면제·감면·기성 인수분 공제가 있으면 그 금액(0 포함)을 넣으세요', None),
        ('준공청구일', '준공금 청구일', 'date', '', '비우면 준공검사일', None),
        ('하자시작', '하자담보 시작일', 'date', '', '비우면 준공검사일', None),
    ]),
    ('9. 하자 검사', [
        ('하자검사일', '하자 검사일', 'date', '', '', None),
        ('하자결과', '하자 검사 결과', 'list', '', '', ['하자 없음', '하자 있음 (보수 요구)']),
    ]),
]

# 자동 계산 (입력 시트 아래) — (키, 이름, 수식 틀, 형식)
# 수식 틀 안의 {키} 는 입력 칸 주소로 바뀝니다.
CALCS = [
    ('공기일수', '공사 기간 (일)', 'IF(OR({착공일}="",{준공기한}=""),"",{준공기한}-{착공일}+1)', '0"일"'),
    ('계약보증금', '계약보증금 (원)', 'IF(OR({계약금액}="",{계약보증률}=""),"",ROUNDUP({계약금액}*{계약보증률}/100,0))', '#,##0'),
    ('정산액', '최종 계약금액 (원)', 'IF({정산금액}="",{계약금액},{정산금액})', '#,##0'),
    ('하자보증금', '하자보수보증금 (원)', 'IF(OR({정산액}="",{하자보증률}=""),"",ROUNDUP({정산액}*{하자보증률}/100,0))', '#,##0'),
    ('검사요청eff', '준공검사 요청일', 'IF({검사요청일}="",{준공일},{검사요청일})', 'yyyy-mm-dd'),
    ('하자시작eff', '하자담보 시작일', 'IF({하자시작}="",{검사일},{하자시작})', 'yyyy-mm-dd'),
    ('하자끝', '하자담보 끝나는 날', 'IF(OR({하자시작eff}="",{하자년}=""),"",EDATE({하자시작eff},{하자년}*12)-1)', 'yyyy-mm-dd'),
    ('기한eff', '지체 계산 기준 준공기한', 'IF({변경기한}="",{준공기한},{변경기한})', 'yyyy-mm-dd'),
    ('지체일수', '지체 일수', 'IF(OR({준공일}="",{기한eff}=""),"",MAX(0,{준공일}-{기한eff}))', '0"일"'),
    ('지체자동', '지체상금 (자동 계산)', 'IF(OR({지체일수}="",{지체율}=""),"",ROUNDDOWN({정산액}*{지체율}/1000*{지체일수},0))', '#,##0'),
    ('지체상금', '지체상금 (서류에 쓰는 값)', 'IF({지체상금입력}="",IF({지체자동}="",0,{지체자동}),{지체상금입력})', '#,##0'),
    ('기성누계', '이번 회까지 기성 누계 (원)', 'IF({금회기성}="","",N({전회누계})+{금회기성})', '#,##0'),
    ('기성률', '기성률 (누계 ÷ 계약 금액)', 'IF(OR({기성누계}="",N({계약금액})=0),"",{기성누계}/{계약금액})', '0.00%'),
    ('기성청구액', '기성금 청구액 (원)', 'IF({금회기성}="","",{금회기성}-N({금회선금정산}))', '#,##0'),
    ('기성청구eff', '기성금 청구일', 'IF({기성청구일}="",{기성요청일},{기성청구일})', 'yyyy-mm-dd'),
    ('준공청구액', '준공금 청구액 (원)', 'IF({정산액}="","",{정산액}-N({기성지급누계})-N({준공선금정산})-N({지체상금}))', '#,##0'),
    ('준공청구eff', '준공금 청구일', 'IF({준공청구일}="",{검사일},{준공청구일})', 'yyyy-mm-dd'),
    ('연장일수', '연장 요청 일수', 'IF(OR({연장기한}="",{준공기한}=""),"",{연장기한}-{준공기한})', '0"일"'),
    ('증감액', '정산 증감액 (원)', 'IF(OR({정산액}="",{계약금액}=""),"",{정산액}-{계약금액})', '#,##0;[Red]-#,##0'),
]
# 금액을 한글로 (계산 시트)
HANGUL = ['계약금액', '정산액', '계약보증금', '하자보증금', '기성청구액', '준공청구액', '금회기성', '기성누계', '증감액abs']

REQUIRED = ['공사명', '현장위치', '계약일', '착공일', '준공기한', '계약금액', '발주기관', '받는사람', '상호', '대표자', '업체주소', '대리인']

REF = {}      # 키 → "'입력'!$C$n"
CELL = {}     # 키 → "C n"


def R(k):
    return REF[k]


def fx(tpl):
    """수식 틀 {키} → 칸 주소"""
    out = tpl
    for k in sorted(REF, key=len, reverse=True):
        out = out.replace('{' + k + '}', REF[k])
    assert '{' not in out, out
    return out


# ─── 글자로 바꾸는 작은 수식들 ───
def DS(r):       # 2026. 3. 2.
    return f'IF({r}="","",YEAR({r})&". "&MONTH({r})&". "&DAY({r})&".")'


def DL(r):       # 2026년 3월 2일 (비면 손으로 쓰는 빈칸)
    return f'IF({r}="","          년      월      일",YEAR({r})&"년  "&MONTH({r})&"월  "&DAY({r})&"일")'


def WON(r):      # 123,456,000원
    return f'IF({r}="","",FIXED({r},0)&"원")'


def KRW(k):      # 일금 일억…원정 (₩123,456,000)
    r = REF[k] if k in REF else k
    return f'IF({r}="","",HANGUL_{k}&"원정 (₩"&FIXED({r},0)&")")'


def PERIOD(a, b):
    ra, rb = R(a), R(b)
    return f'IF(OR({ra}="",{rb}=""),"",{DS(ra)}&" ~ "&{DS(rb)}&" ("&({rb}-{ra}+1)&"일)")'


def T(k):        # 입력 글자 그대로
    return f'{R(k)}&""'


# ─────────────────────── 서류 정의 ───────────────────────
# 블록: ('title', 글) ('kv', [(이름, 수식|글, 높이?)]) ('text', 글) ('ftext', 수식) ('table', 머리, 너비, 줄들, 높이)
#       ('date', 수식) ('sign', [(이름, 수식)]) ('to',) ('note', 글) ('gap', pt)
class F(str):
    """수식 표시 (앞에 = 붙여 넣음)"""


def kvF(label, expr, h=None):
    return (label, F(expr), h)


def CONTRACTOR(extra=()):
    rows = list(extra) + [
        ('상    호', F(T('상호'))),
        ('주    소', F(T('업체주소'))),
        ('대 표 자', F(f'{R("대표자")}&"          (인)"')),
    ]
    return ('sign', '계약상대자', rows)


def TO():
    return ('to', F(f'IF({R("받는사람")}="","                  귀하",{R("받는사람")}&"  귀하")'))


ATT = lambda rows: ('table', ['붙임 서류', '부수', '비고'], [7, 2, 3], [[r, '1부' if r else '', ''] for r in rows], 22)

DOCS = []


STAGE = {'계약': 1, '착공': 2, '공사 중': 3, '준공': 4, '관리': 5, '하자': 6}


def doc(no, sheet, title, who, when, pub, blocks, date_key=None, land=False, desc=''):
    DOCS.append(dict(no=no, name=sheet[3:], title=title, who=who, when=when, pub=pub, blocks=blocks,
                     date_key=date_key, land=land, desc=desc))


def number_docs():
    """업체 서류 먼저, 그 안에서 계약 → 착공 → 공사 중 → 준공 순서로 번호를 다시 매김"""
    DOCS.sort(key=lambda d: (d['who'] == '발주기관', STAGE[d['when']], d['no']))
    for i, d in enumerate(DOCS, 1):
        d['no'] = i
        d['sheet'] = f'{i:02d} {d["name"]}'


def build_docs():
    DOCS.clear()
    # ── 착공 ──
    doc(1, '01 착공신고서', '착 공 신 고 서', '업체', '착공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계 약 일', DS(R('계약일'))),
                kvF('계약 금액', KRW('계약금액'), 34), kvF('공사 기간', PERIOD('착공일', '준공기한')),
                kvF('착 공 일', DS(R('착공일'))), kvF('준공 기한', DS(R('준공기한'))), kvF('현장 위치', T('현장위치'), 34),
                kvF('현장대리인', f'{R("대리인")}&IF({R("대리인자격")}="",""," ("&{R("대리인자격")}&")")')]),
        ('text', '위 공사를 계약 조건에 따라 착공하였기에 관계 서류를 붙여 신고합니다.'),
        ATT(['현장대리인 지정 신고서', '공사 예정공정표', '직접시공계획서 (해당 시)', '안전·보건 관리 계획서 (해당 시)',
             '착공 전 현장 사진', '']),
        ('date', '착공일'), CONTRACTOR(), TO()], date_key='착공일',
        desc='공사를 시작했다고 알리는 서류. 공사명·계약금액·공사기간이 입력에서 채워집니다.')

    doc(2, '02 현장대리인계', '현 장 대 리 인 지 정 신 고 서', '업체', '착공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('공사 기간', PERIOD('착공일', '준공기한')),
                kvF('현장 위치', T('현장위치'), 34)]),
        ('table', ['성  명', '생년월일', '자격 종목·등급', '자격(등록) 번호', '연 락 처'], [2, 2, 3, 3, 2],
         [[F(T('대리인')), F(T('대리인생일')), F(T('대리인자격')), F(T('대리인번호')), F(T('대리인전화'))]], 34),
        ('kv', [kvF('배치 기간', f'IF(OR({R("착공일")}="",{R("준공기한")}=""),"",{DS(R("착공일"))}&" ~ "&{DS(R("준공기한"))})')]),
        ('text', '위 사람을 이 공사의 현장대리인으로 지정하여 배치하였기에 신고합니다. 현장대리인은 계약 조건과 관계 법령에 따라 '
                 '이 공사의 현장 관리 업무를 맡습니다.'),
        ATT(['자격증(경력수첩) 사본', '경력증명서', '재직 증명 (4대 보험 가입 증명 등)']),
        ('date', '착공일'), ('sign', '현장대리인', [('성    명', F(f'{R("대리인")}&"          (인)"'))]), CONTRACTOR(), TO()], date_key='착공일',
        desc='현장대리인의 자격·연락처를 알리는 서류.')

    months = [f'IF({R("착공일")}="","",MONTH(EDATE(DATE(YEAR({R("착공일")}),MONTH({R("착공일")}),1),{k}))&"월")'
              for k in range(8)]
    doc(3, '03 예정공정표', '공 사 예 정 공 정 표', '업체', '착공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('공사 기간', PERIOD('착공일', '준공기한'))]),
        ('table', ['공  종', '비중(%)'] + [F(m) for m in months], [3, 1] + [1] * 8,
         [['', ''] + [''] * 8 for _ in range(12)] +
         [['월별 예정공정률(%)', ''] + [''] * 8, ['누계 예정공정률(%)', ''] + [''] * 8], 24),
        ('note', '※ 공종마다 시작하는 달부터 끝나는 달까지 칸을 칠하거나 선(━)을 긋습니다. 달 이름은 착공일에 맞춰 저절로 바뀝니다.'),
        ('date', '착공일'), CONTRACTOR(), TO()], date_key='착공일', desc='공종별 공사 일정을 달 단위로 적는 표.')

    doc(4, '04 직접시공계획서', '직 접 시 공 계 획 서', '업체', '착공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 금액', KRW('계약금액'), 34),
                kvF('공사 기간', PERIOD('착공일', '준공기한')), kvF('현장 위치', T('현장위치'), 34)]),
        ('table', ['공  종', '직접 시공할 내용', '금액 (원)', '비율(%)', '비고'], [2, 4, 3, 1, 2],
         [['', '', '', '', ''] for _ in range(8)] + [['합  계', '', F('SUMCOL'), F('RATIOCOL'), '']], 24),
        ('text', '위와 같이 이 공사를 직접 시공할 계획을 알려 드립니다.'),
        ('note', '※ 직접 시공 대상과 비율은 건설산업기본법(직접 시공 조항)과 계약 조건을 확인하세요.'),
        ('date', '착공일'), CONTRACTOR(), TO()], date_key='착공일', desc='우리 회사가 직접 시공할 공종·금액을 알리는 계획서.')

    doc(5, '05 사용인감계', '사 용 인 감 계', '업체', '계약', False, [
        ('kv', [kvF('상    호', T('상호')), kvF('대 표 자', T('대표자')), kvF('사업자등록번호', T('사업자번호')),
                kvF('주    소', T('업체주소'), 34)]),
        ('kv', [kvF('대상 계약', T('공사명'), 34), ('사용 기간', '이 계약의 체결부터 하자담보책임기간이 끝날 때까지', None)]),
        ('table', ['구  분', '인  감', '비  고'], [3, 5, 4], [['법인(개인) 인감', '', '인감증명서 붙임'], ['사용 인감', '', '']], 70),
        ('text', '위 사용 인감을 이 계약과 관련된 모든 서류에 법인(개인) 인감을 대신하여 사용하고자 신고하며, '
                 '이 인감을 사용하여 생기는 모든 책임은 우리 회사가 집니다.'),
        ('date', '계약일'), CONTRACTOR(), TO()], date_key='계약일', desc='법인 인감 대신 쓸 도장을 신고하는 서류.')

    doc(6, '06 계약보증금 납부서', '계 약 보 증 금 납 부 서', '업체', '계약', True, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('계약금액'), 34),
                kvF('보증금률', f'IF({R("계약보증률")}="","",{R("계약보증률")}&"%")'),
                kvF('계약보증금', KRW('계약보증금'), 34), kvF('납부 방법', T('계약보증방법')),
                kvF('보증 기간', f'IF(OR({R("계약일")}="",{R("준공기한")}=""),"",{DS(R("계약일"))}&" ~ "&{DS(R("준공기한"))})')]),
        ('text', '위 계약의 계약보증금을 위와 같이 납부합니다.'),
        ATT(['계약보증서(보증증권) — 현금이면 납부 영수증']),
        ('note', '※ 보증 기간은 보증서에 적힌 기간을 확인하세요.'),
        ('date', '계약일'), CONTRACTOR(), TO()], date_key='계약일', desc='계약보증금(보증서·현금)을 내는 서류. 금액은 요율로 자동 계산.')

    doc(7, '07 계약보증금 지급각서', '계 약 보 증 금 지 급 각 서', '업체', '계약', True, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('계약금액'), 34),
                kvF('계약보증금', KRW('계약보증금'), 34)]),
        ('text', '우리 회사는 위 계약의 계약보증금 납부를 면제받는 대신, 계약상의 의무를 이행하지 아니하여 계약보증금이 '
                 '발주기관에 귀속되는 사유가 생긴 때에는 위 계약보증금에 해당하는 금액을 발주기관의 청구에 따라 '
                 '즉시 현금으로 납부할 것을 각서합니다.'),
        ('date', '계약일'), CONTRACTOR(), TO()], date_key='계약일', desc='계약보증금 대신 내는 지급각서 (면제 대상일 때).')

    doc(8, '08 청렴계약 이행서약서', '청 렴 계 약 이 행 서 약 서', '업체', '계약', True, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('상    호', T('상호')),
                kvF('대 표 자', T('대표자'))]),
        ('text', '우리 회사는 위 공사의 입찰·계약 체결과 이행에 관련하여 다음 사항을 지킬 것을 서약합니다.'),
        ('text', '1. 관계 공무원(임직원)에게 직접 또는 간접으로 금품·향응 등을 주거나 받지 않겠습니다.'),
        ('text', '2. 담합 등 공정한 경쟁을 방해하는 행위를 하지 않겠습니다.'),
        ('text', '3. 하도급·자재·장비 업체와의 거래에서도 부당한 요구를 하거나 금품을 주고받지 않겠습니다.'),
        ('text', '4. 위 사항을 어긴 때에는 계약 해제·해지, 입찰참가자격 제한 등 관계 법령과 계약 조건에 따른 조치를 받아들이겠습니다.'),
        ('date', '계약일'), CONTRACTOR(), TO()], date_key='계약일', desc='금품·담합을 하지 않겠다는 서약서.')

    doc(9, '09 수의계약 각서', '수 의 계 약 각 서', '업체', '계약', True, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 금액', KRW('계약금액'), 34), kvF('계약 방법', T('계약방법'))]),
        ('text', '우리 회사는 위 공사를 수의계약으로 체결하면서 다음 사항을 지킬 것을 각서합니다.'),
        ('text', '1. 견적서 제출과 계약 체결 과정에서 다른 업체와 짜거나 부정한 방법을 쓰지 않았습니다.'),
        ('text', '2. 수의계약 상대자가 되는 데 결격 사유가 없으며, 제출한 서류는 사실과 다르지 않습니다.'),
        ('text', '3. 계약을 이행하면서 관계 법령과 계약 조건을 성실히 지키겠습니다.'),
        ('text', '4. 위 사항이 사실과 다르거나 이를 어긴 때에는 계약 해제·해지 등 관계 법령에 따른 조치를 받아들이겠습니다.'),
        ('date', '계약일'), CONTRACTOR(), TO()], date_key='계약일', desc='수의계약일 때 내는 각서.')

    doc(10, '10 수도전기료 납부각서', '수 도 · 전 기 료 납 부 각 서', '업체', '착공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('공사 기간', PERIOD('착공일', '준공기한')), kvF('현장 위치', T('현장위치'), 34)]),
        ('text', '우리 회사는 위 공사 기간 동안 발주기관(시설)의 수도·전기 등을 공사에 쓰는 경우, 사용량을 계량하거나 '
                 '발주기관과 협의한 방법으로 산정한 요금을 발주기관이 정한 기한까지 납부할 것을 각서합니다.'),
        ('table', ['구  분', '계량 방법', '시작 지침', '끝 지침', '비  고'], [2, 3, 2, 2, 3],
         [['수  도', '', '', '', ''], ['전  기', '', '', '', ''], ['기  타', '', '', '', '']], 26),
        ('date', '착공일'), CONTRACTOR(), TO()], date_key='착공일', desc='기존 시설 안에서 공사할 때 — 수도·전기 요금 정산 각서.')

    # ── 공사 중 ──
    doc(11, '11 기성검사 요청서', '기 성 부 분 검 사 요 청 서', '업체', '공사 중', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('계약금액'), 34),
                kvF('공사 기간', PERIOD('착공일', '준공기한')),
                kvF('기성 회차', f'IF({R("기성회차")}="","","제 "&{R("기성회차")}&" 회")'),
                kvF('기성 대상기간', f'IF(OR({R("기성시작")}="",{R("기성끝")}=""),"",{DS(R("기성시작"))}&" ~ "&{DS(R("기성끝"))})')]),
        ('table', ['구  분', '금  액 (원)', '계약 금액 대비'], [4, 5, 3],
         [['계약 금액', F(WON(R('계약금액'))), ''],
          ['전회까지 기성 누계', F(WON(R('전회누계'))), F(f'IF(OR({R("전회누계")}="",N({R("계약금액")})=0),"",TEXT({R("전회누계")}/{R("계약금액")},"0.00%"))')],
          ['이번 회 기성', F(WON(R('금회기성'))), F(f'IF(OR({R("금회기성")}="",N({R("계약금액")})=0),"",TEXT({R("금회기성")}/{R("계약금액")},"0.00%"))')],
          ['이번 회까지 누계', F(WON(R('기성누계'))), F(f'IF({R("기성률")}="","",TEXT({R("기성률")},"0.00%"))')]], 26),
        ('text', '위 공사의 기성 부분에 대하여 검사를 요청합니다.'),
        ATT(['기성 내역서', '기성 부분 사진', '품질시험 성과 (해당 시)']),
        ('date', '기성요청일'), CONTRACTOR(), TO()], date_key='기성요청일', desc='기성 부분 검사를 요청하는 서류. 누계·기성률 자동.')

    doc(12, '12 기성대가 청구서', '기 성 대 가 청 구 서', '업체', '공사 중', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')),
                kvF('기성 회차', f'IF({R("기성회차")}="","","제 "&{R("기성회차")}&" 회")')]),
        ('kv', [kvF('이번 회 기성금액', WON(R('금회기성'))), kvF('선금 정산액 (공제)', WON(R('금회선금정산'))),
                kvF('청구 금액', KRW('기성청구액'), 34),
                kvF('공급가액 / 부가세', f'IF({R("기성청구액")}="","",FIXED(ROUND({R("기성청구액")}/1.1,0),0)&"원  /  "&FIXED({R("기성청구액")}-ROUND({R("기성청구액")}/1.1,0),0)&"원")')]),
        ('kv', [kvF('은 행 명', T('은행')), kvF('예 금 주', T('예금주')), kvF('계좌 번호', T('계좌'))]),
        ('text', '위 공사의 기성 대가를 위와 같이 청구합니다.'),
        ('note', '※ 공급가액·부가세는 청구 금액을 1.1로 나눠 계산한 값입니다. 세금계산서 금액과 맞는지 확인하세요.'),
        ('date', '기성청구eff'), CONTRACTOR(), TO()], date_key='기성청구eff', desc='기성금을 청구하는 서류. 선금 정산을 빼고 계산.')

    doc(13, '13 공기연장 신청서', '공 기 연 장 신 청 서', '업체', '공사 중', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('계약금액'), 34),
                kvF('당초 공사기간', PERIOD('착공일', '준공기한')), kvF('연장 요청 기한', DS(R('연장기한'))),
                kvF('연장 일수', f'IF({R("연장일수")}="","",{R("연장일수")}&"일")'),
                kvF('연장 사유', T('연장사유'), 60)]),
        ('text', '위와 같은 사유로 계약 기간 안에 공사를 마칠 수 없어 계약 기간의 연장을 신청합니다.'),
        ATT(['연장 사유 증빙 자료', '변경 예정공정표']),
        ('date', '연장신청일'), CONTRACTOR(), TO()], date_key='연장신청일', desc='준공기한을 늘려 달라는 신청서. 연장 일수 자동.')

    # ── 준공 ──
    doc(14, '14 준공신고서', '준 공 신 고 서', '업체', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('정산액'), 34),
                kvF('공사 기간', PERIOD('착공일', '준공기한')), kvF('착 공 일', DS(R('착공일'))),
                kvF('준 공 일', DS(R('준공일'))), kvF('현장 위치', T('현장위치'), 34), kvF('현장대리인', T('대리인'))]),
        ('text', '위 공사를 계약 조건에 따라 준공하였기에 관계 서류를 붙여 신고합니다.'),
        ATT(['준공 사진', '준공 내역서', '준공 도면 (해당 시)', '품질시험 성과 (해당 시)', '하도급대금 지급 확인서 (해당 시)']),
        ('date', '준공일'), CONTRACTOR(), TO()], date_key='준공일', desc='공사를 마쳤다고 알리는 서류(준공계).')

    doc(15, '15 준공검사 요청서', '준 공 검 사 요 청 서', '업체', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('계약 금액', KRW('정산액'), 34),
                kvF('착 공 일', DS(R('착공일'))), kvF('준 공 일', DS(R('준공일'))), kvF('현장 위치', T('현장위치'), 34)]),
        ('text', '위 공사를 준공하였으므로 준공검사를 하여 주시기 바랍니다.'),
        ('date', '검사요청eff'), CONTRACTOR(), TO()], date_key='검사요청eff', desc='준공검사를 해 달라는 요청서(준공검사원).')

    doc(16, '16 준공대가 청구서', '준 공 대 가 청 구 서', '업체', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('준 공 일', DS(R('준공일'))),
                kvF('준공검사일', DS(R('검사일')))]),
        ('kv', [kvF('최종 계약금액', WON(R('정산액'))), kvF('기성 지급 누계 (공제)', WON(R('기성지급누계'))),
                kvF('선금 잔액 정산 (공제)', WON(R('준공선금정산'))),
                kvF('지체상금 (공제)', f'IF(N({R("지체상금")})=0,"없음",FIXED({R("지체상금")},0)&"원 (지체 "&{R("지체일수")}&"일)")'),
                kvF('청구 금액', KRW('준공청구액'), 34)]),
        ('kv', [kvF('은 행 명', T('은행')), kvF('예 금 주', T('예금주')), kvF('계좌 번호', T('계좌'))]),
        ('text', '위 공사의 준공 대가를 위와 같이 청구합니다.'),
        ('date', '준공청구eff'), CONTRACTOR(), TO()], date_key='준공청구eff', desc='준공금을 청구하는 서류. 기성·선금·지체상금 공제 자동.')

    doc(17, '17 준공정산 동의서', '준 공 정 산 동 의 서', '업체', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('당초 계약금액', KRW('계약금액'), 34),
                kvF('정산 계약금액', KRW('정산액'), 34),
                kvF('증 감 액', f'IF({R("증감액")}="","",IF({R("증감액")}<0,"감액 ","증액 ")&FIXED(ABS({R("증감액")}),0)&"원")')]),
        ('text', '우리 회사는 위 공사의 준공 정산 결과 최종 계약금액을 위와 같이 확정하는 데 동의하며, '
                 '이 금액에 대하여 앞으로 이의를 제기하지 않겠습니다.'),
        ('date', '검사일'), CONTRACTOR(), TO()], date_key='검사일', desc='정산된 최종 금액에 동의하는 서류.')

    doc(18, '18 하자보수보증금 납부서', '하 자 보 수 보 증 금 납 부 서', '업체', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('최종 계약금액', KRW('정산액'), 34),
                kvF('보증금률', f'IF({R("하자보증률")}="","",{R("하자보증률")}&"%")'),
                kvF('하자보수보증금', KRW('하자보증금'), 34), kvF('납부 방법', T('하자보증방법')),
                kvF('하자담보책임기간', f'IF({R("하자끝")}="","",{R("하자년")}&"년 ("&{DS(R("하자시작eff"))}&" ~ "&{DS(R("하자끝"))}&")")', 34)]),
        ('text', '위 공사의 하자보수보증금을 위와 같이 납부합니다.'),
        ATT(['하자보수보증서(보증증권) — 현금이면 납부 영수증']),
        ('date', '검사일'), CONTRACTOR(), TO()], date_key='검사일', desc='하자보수보증금을 내는 서류. 금액·기간 자동.')

    doc(19, '19 하자보수보증금 지급각서', '하 자 보 수 보 증 금 지 급 각 서', '업체', '준공', True, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')), kvF('하자보수보증금', KRW('하자보증금'), 34),
                kvF('하자담보책임기간', f'IF({R("하자끝")}="","",{DS(R("하자시작eff"))}&" ~ "&{DS(R("하자끝"))})')]),
        ('text', '우리 회사는 위 공사의 하자보수보증금 납부를 면제받는 대신, 하자담보책임기간 안에 하자보수 의무를 '
                 '이행하지 아니하여 하자보수보증금이 발주기관에 귀속되는 사유가 생긴 때에는 위 금액을 발주기관의 '
                 '청구에 따라 즉시 현금으로 납부할 것을 각서합니다.'),
        ('date', '검사일'), CONTRACTOR(), TO()], date_key='검사일', desc='하자보수보증금 대신 내는 지급각서.')

    # ── 발주기관이 쓰는 것 ──
    doc(20, '20 공사감독조서', '공 사 감 독 조 서', '발주기관', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')),
                kvF('계약상대자', f'{R("상호")}&IF({R("대표자")}="",""," (대표 "&{R("대표자")}&")")'),
                kvF('계약 금액', KRW('정산액'), 34), kvF('계 약 일', DS(R('계약일'))), kvF('착 공 일', DS(R('착공일'))),
                kvF('준공 기한', DS(R('기한eff'))), kvF('준 공 일', DS(R('준공일'))),
                kvF('지체 일수', f'IF({R("지체일수")}="","",IF({R("지체일수")}=0,"없음",{R("지체일수")}&"일"))')]),
        ('text', '위 공사를 설계서와 계약 조건에 따라 감독한 결과 위와 같이 시공되었음을 확인합니다.'),
        ('date', '검사요청eff'), ('sign', '공사 감독자', [('직·성명', F(f'{R("감독자")}&"          (인)"'))]), TO()],
        date_key='검사요청eff', desc='감독자가 쓰는 감독 결과 조서.')

    doc(21, '21 준공검사조서', '준 공 검 사 조 서', '발주기관', '준공', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('계약 번호', T('계약번호')),
                kvF('계약상대자', f'{R("상호")}&IF({R("대표자")}="",""," (대표 "&{R("대표자")}&")")'),
                kvF('계약 금액', KRW('정산액'), 34), kvF('착 공 일', DS(R('착공일'))), kvF('준 공 일', DS(R('준공일'))),
                kvF('검사 요청일', DS(R('검사요청eff'))), kvF('준공검사일', DS(R('검사일'))),
                kvF('검사 결과', f'IF({R("검사결과")}="","합격  /  보완 요구",{R("검사결과")})')]),
        ('text', '위 공사에 대하여 설계서·계약서와 관계 서류에 따라 준공검사를 한 결과 위와 같음을 확인합니다.'),
        ('date', '검사일'),
        ('sign', '검 사 자', [('직·성명', F(f'{R("검사자")}&"          (인)"'))]),
        ('sign', '입 회 자', [('직·성명', F(f'{R("입회자")}&"          (인)"'))]), TO()],
        date_key='검사일', desc='검사자가 쓰는 준공검사 조서.')

    doc(22, '22 공사대장', '공 사 대 장', '발주기관', '관리', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34), kvF('현장 위치', T('현장위치'), 34), kvF('공사 개요', T('공사개요'), 34),
                kvF('계약 방법', T('계약방법')), kvF('계약 번호', T('계약번호')), kvF('계 약 일', DS(R('계약일'))),
                kvF('당초 계약금액', WON(R('계약금액'))), kvF('최종 계약금액', WON(R('정산액'))),
                kvF('공사 기간', PERIOD('착공일', '준공기한')), kvF('준 공 일', DS(R('준공일'))),
                kvF('준공검사일', DS(R('검사일'))),
                kvF('계약상대자', f'{R("상호")}&IF({R("대표자")}="",""," / 대표 "&{R("대표자")})&IF({R("업체전화")}="",""," / "&{R("업체전화")})', 34),
                kvF('현장대리인', f'{R("대리인")}&IF({R("대리인전화")}="",""," / "&{R("대리인전화")})'),
                kvF('설 계 자', T('설계자')), kvF('감 리 자', T('감리자')), kvF('공사 감독자', T('감독자')),
                kvF('계약보증', f'IF({R("계약보증금")}="",{R("계약보증방법")}&"",FIXED({R("계약보증금")},0)&"원 / "&{R("계약보증방법")})'),
                kvF('하자보수보증', f'IF({R("하자보증금")}="",{R("하자보증방법")}&"",FIXED({R("하자보증금")},0)&"원 / "&{R("하자보증방법")})'),
                kvF('하자담보책임기간', f'IF({R("하자끝")}="","",{DS(R("하자시작eff"))}&" ~ "&{DS(R("하자끝"))})'),
                kvF('지체 일수·상금', f'IF({R("지체일수")}="","",IF({R("지체일수")}=0,"없음",{R("지체일수")}&"일 / "&FIXED(N({R("지체상금")}),0)&"원"))')]),
        ('note', '※ 공사 하나에 한 장 — 계약부터 하자까지 한눈에 보는 관리 카드입니다.')], date_key=None,
        desc='공사 하나를 한 장으로 정리한 관리 카드.')

    doc(23, '23 하자검사조서', '하 자 검 사 조 서', '발주기관', '하자', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34),
                kvF('계약상대자', f'{R("상호")}&IF({R("대표자")}="",""," (대표 "&{R("대표자")}&")")'),
                kvF('준 공 일', DS(R('준공일'))),
                kvF('하자담보책임기간', f'IF({R("하자끝")}="","",{DS(R("하자시작eff"))}&" ~ "&{DS(R("하자끝"))})'),
                kvF('검 사 일', DS(R('하자검사일'))),
                kvF('검사 결과', f'IF({R("하자결과")}="","하자 없음  /  하자 있음 (보수 요구)",{R("하자결과")})')]),
        ('table', ['위  치', '하자 내용', '보수 요구 사항', '비  고'], [2, 4, 4, 2], [['', '', '', ''] for _ in range(6)], 28),
        ('text', '위 공사의 하자담보책임기간 중 하자 검사를 한 결과 위와 같음을 확인합니다.'),
        ('date', '하자검사일'),
        ('sign', '검 사 자', [('직·성명', F(f'{R("검사자")}&"          (인)"'))]),
        ('sign', '입 회 자', [('계약상대자', F(f'{R("상호")}&"          (인)"'))]), TO()],
        date_key='하자검사일', desc='하자 검사 결과를 적는 조서.')

    doc(24, '24 하자대장', '하 자 대 장', '발주기관', '하자', False, [
        ('kv', [kvF('공 사 명', T('공사명'), 34),
                kvF('계약상대자', f'{R("상호")}&IF({R("업체전화")}="",""," / "&{R("업체전화")})'),
                kvF('하자담보책임기간', f'IF({R("하자끝")}="","",{DS(R("하자시작eff"))}&" ~ "&{DS(R("하자끝"))})'),
                kvF('하자보수보증', f'IF({R("하자보증금")}="",{R("하자보증방법")}&"",FIXED({R("하자보증금")},0)&"원 / "&{R("하자보증방법")})')]),
        ('table', ['번호', '접 수 일', '위치 · 하자 내용', '보수 요구일', '보수 완료일', '확 인'], [1, 2, 4, 2, 2, 1],
         [[str(i + 1), '', '', '', '', ''] for i in range(14)], 26)], date_key=None,
        desc='하자 접수부터 보수 완료까지 적는 대장.')


# ───────────────────────── 그리기 ─────────────────────────
def col(i):          # 서식 칸 번호(0..11) → 엑셀 열 번호
    return C0 + i


def box(ws, r1, c1, r2, c2, fill=None, border=True):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if border:
                cell.border = Border(left=thin if c == c1 else None, right=thin if c == c2 else None,
                                     top=thin if r == r1 else None, bottom=thin if r == r2 else None)
            if fill:
                cell.fill = fill


def put(ws, r, c1, c2, val, font=None, align=None, fill=None, border=True, fmt=None):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(r, c1)
    if isinstance(val, F):
        cell.value = '=' + val
    else:
        cell.value = val if val != '' else None
    cell.font = font or f()
    cell.alignment = align or Alignment(vertical='center', wrap_text=True)
    if fmt:
        cell.number_format = fmt
    box(ws, r, c1, r, c2, fill, border)
    return cell


def text_h(s, width_chars=44, size=11):
    n = 0
    for part in str(s).split('\n'):
        w = sum(2 if ord(ch) > 0x2E80 else 1 for ch in part) / 2
        n += max(1, math.ceil(w / width_chars))
    return n * (size * 1.55) + 8


def spread(ws_widths, n=NC):
    """비율 → 칸 수 (합 = 12)"""
    tot = sum(ws_widths)
    acc, out, prev = 0, [], 0
    for w in ws_widths:
        acc += w
        cur = round(acc * n / tot)
        out.append(max(1, cur - prev)); prev = cur
    d = n - sum(out)
    out[-1] += d
    return out


def make_sheet_for_doc(wb, d, hangul_ref):
    """제목 → 블록 순서로 그림"""
    ws = wb.create_sheet(d['sheet'])
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = {'업체': '2F6FB5', '발주기관': 'B7802F'}[d['who']]
    ws.column_dimensions['A'].width = 1.5
    for i in range(NC):
        ws.column_dimensions[ws.cell(1, col(i)).column_letter].width = COLW
    ws.column_dimensions['N'].width = 2
    ws.column_dimensions['O'].width = 18
    first = col(0); last = col(NC - 1)
    ws.row_dimensions[1].height = 8
    title = d['title']
    put(ws, 2, first, last, title, f(20, True), Alignment(horizontal='center', vertical='center'), border=False)
    ws.row_dimensions[2].height = 44
    sub = {'업체': '', '발주기관': ''}[d['who']]
    ws.row_dimensions[3].height = 10
    r = 4
    stretch = []          # 늘려도 되는 줄 (표·항목)
    date_gap = None       # 날짜 앞 빈 줄

    def expr(v):
        if isinstance(v, F):
            s = str(v)
            for k in sorted(hangul_ref, key=len, reverse=True):
                s = s.replace('HANGUL_' + k, hangul_ref[k])
            assert 'HANGUL_' not in s, s
            return F(s)
        return v

    for b in d['blocks']:
        kind = b[0]
        if kind == 'kv':
            r1 = r
            for label, val, h in b[1]:
                put(ws, r, first, first + 2, label, f(10.5, True), Alignment(horizontal='center', vertical='center', wrap_text=True), HEAD_FILL)
                put(ws, r, first + 3, last, expr(val), f(11), Alignment(horizontal='left', vertical='center', wrap_text=True, indent=1))
                ws.row_dimensions[r].height = h or 26
                stretch.append(r)
                r += 1
            ws.row_dimensions[r].height = 8; r += 1
        elif kind == 'text':
            put(ws, r, first, last, b[1], f(11), Alignment(horizontal='left', vertical='center', wrap_text=True), border=False)
            ws.row_dimensions[r].height = text_h(b[1])
            r += 1
        elif kind == 'note':
            put(ws, r, first, last, b[1], f(9.5, color='555555'), Alignment(horizontal='left', vertical='center', wrap_text=True), border=False)
            ws.row_dimensions[r].height = text_h(b[1], 50, 9.5)
            r += 1
        elif kind == 'table':
            heads, widths, rows, h = b[1], b[2], b[3], b[4]
            spans = spread(widths)
            c = first
            for hd, sp in zip(heads, spans):
                put(ws, r, c, c + sp - 1, expr(hd), f(10.5, True),
                    Alignment(horizontal='center', vertical='center', wrap_text=True), HEAD_FILL)
                c += sp
            ws.row_dimensions[r].height = 24
            r0 = r + 1
            r += 1
            for row in rows:
                c = first
                for j, (v, sp) in enumerate(zip(row, spans)):
                    hd = heads[j] if isinstance(heads[j], str) and not isinstance(heads[j], F) else ''
                    if isinstance(v, F) and str(v) == 'SUMCOL':
                        cl = ws.cell(r0, c).column_letter
                        v = F(f'IF(SUM({cl}{r0}:{cl}{r - 1})=0,"",SUM({cl}{r0}:{cl}{r - 1}))')
                    elif isinstance(v, F) and str(v) == 'RATIOCOL':
                        pc = ws.cell(r0, c - spans[j - 1]).column_letter
                        v = F(f'IF(OR({pc}{r}="",N({R("계약금액")})=0),"",{pc}{r}/{R("계약금액")})')
                    bold = isinstance(row[0], str) and row[0].startswith(('합', '월별', '누계'))
                    al = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    if j > 0 and hd.startswith('금'):
                        al = Alignment(horizontal='right', vertical='center', indent=1)
                    if j == 0 and len(heads) > 2 and heads[1] == '금  액 (원)':
                        al = Alignment(horizontal='left', vertical='center', indent=1)
                    cell = put(ws, r, c, c + sp - 1, expr(v), f(10.5, bold), al)
                    if hd.startswith('금액') or hd.startswith('금  액'):
                        cell.number_format = '#,##0'
                    if hd.startswith('비율'):
                        cell.number_format = '0.0%'
                    c += sp
                ws.row_dimensions[r].height = h
                stretch.append(r)
                r += 1
            ws.row_dimensions[r].height = 8; r += 1
        elif kind == 'date':
            ws.row_dimensions[r].height = 8
            date_gap = r
            r += 1
            put(ws, r, first, last, F(DL(R(b[1]))), f(12), Alignment(horizontal='center', vertical='center'), border=False)
            ws.row_dimensions[r].height = 30
            r += 1
            ws.row_dimensions[r].height = 6; r += 1
        elif kind == 'sign':
            role, rows = b[1], b[2]
            for k, (lab, val) in enumerate(rows):
                if k == 0:
                    put(ws, r, first, first + 3, role, f(11, True), Alignment(horizontal='right', vertical='center'), border=False)
                dist = len(lab.replace(' ', '')) <= 5 and '·' not in lab
                put(ws, r, first + 5, first + 6, lab, f(11), Alignment(horizontal='distributed' if dist else 'left', vertical='center'), border=False)
                put(ws, r, first + 7, first + 7, ':', f(11), Alignment(horizontal='center', vertical='center'), border=False)
                put(ws, r, first + 8, last, expr(val), f(11), Alignment(horizontal='left', vertical='center', shrink_to_fit=True), border=False)
                ws.row_dimensions[r].height = 24
                r += 1
            ws.row_dimensions[r].height = 8; r += 1
        elif kind == 'to':
            ws.row_dimensions[r].height = 8; r += 1
            put(ws, r, first, last, expr(b[1]), f(15, True), Alignment(horizontal='left', vertical='center'), border=False)
            ws.row_dimensions[r].height = 30
            r += 1
    last_row = r - 1
    # A4 한 장을 보기 좋게 채움 — 항목 줄을 조금 키우고, 남으면 날짜 앞을 띄움
    TARGET = 700
    tot = sum((ws.row_dimensions[i].height or 15) for i in range(1, last_row + 1))
    if tot < TARGET and stretch:
        sh = sum(ws.row_dimensions[i].height for i in stretch)
        k = min(1.3, 1 + (TARGET - tot) / sh)
        for i in stretch:
            ws.row_dimensions[i].height = round(ws.row_dimensions[i].height * k, 1)
        tot = sum((ws.row_dimensions[i].height or 15) for i in range(1, last_row + 1))
    if tot < TARGET and date_gap:
        ws.row_dimensions[date_gap].height = 8 + min(TARGET - tot, 120)
    # ← 처음으로 (인쇄 범위 밖)
    c = ws.cell(2, 15, '← 처음(목록)으로')
    c.hyperlink = Hyperlink(ref='O2', location="'처음'!A1", display='← 처음(목록)으로')
    c.font = Font(name=FONT, size=10, color='1F5FBF', underline='single')
    c2 = ws.cell(3, 15, '✎ 입력 고치기')
    c2.hyperlink = Hyperlink(ref='O3', location="'입력'!A1", display='✎ 입력 고치기')
    c2.font = Font(name=FONT, size=10, color='1F5FBF', underline='single')
    ws.cell(5, 15, '칸은 「입력」 시트에서 저절로 채워집니다. 여기서 직접 고쳐 써도 됩니다 (그 칸만 수식이 바뀜).').font = f(9, color='777777')
    ws.cell(5, 15).alignment = Alignment(wrap_text=True, vertical='top')
    ws.print_area = f'A1:{ws.cell(1, last).column_letter}{last_row}'
    page(ws, d.get('land'))
    return ws


def page(ws, land=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if land else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.6
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3
    ws.print_options.horizontalCentered = True
    ws.oddHeader.right.text = 'K-건설맵 · k-conmap.com'
    ws.oddHeader.right.font = '맑은 고딕,보통'
    ws.oddHeader.right.size = 8
    ws.oddHeader.right.color = '808080'


# ───────────────────────── 입력 시트 ─────────────────────────
def make_input(wb):
    ws = wb.create_sheet('입력')
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = 'E0A800'
    for c, w in zip('ABCD', (1.5, 30, 42, 60)):
        ws.column_dimensions[c].width = w
    put(ws, 1, 2, 4, '공사서류 원클릭 — 입력', f(18, True), Alignment(vertical='center'), border=False)
    ws.row_dimensions[1].height = 36
    put(ws, 2, 2, 4, '노란 칸만 채우면 서류 24가지에 저절로 들어갑니다. 모르는 칸은 비워 두세요 — 서류에서 그 자리만 빈칸으로 나옵니다.',
        f(10.5, color='333333'), Alignment(vertical='center', wrap_text=True), border=False)
    ws.row_dimensions[2].height = 30
    back = ws.cell(3, 2, '← 처음(목록)으로')
    back.hyperlink = Hyperlink(ref='B3', location="'처음'!A1", display='← 처음(목록)으로')
    back.font = Font(name=FONT, size=10, color='1F5FBF', underline='single')
    r = 5
    meta_inputs = []
    dvs = {}
    for sec, items in SECTIONS:
        put(ws, r, 2, 4, sec, f(12, True, 'FFFFFF'), Alignment(vertical='center', indent=1), SEC_FILL, border=False)
        ws.row_dimensions[r].height = 24
        r += 1
        for key, label, typ, ex, hlp, opts in items:
            put(ws, r, 2, 2, label, f(10.5, True), Alignment(vertical='center', indent=1, wrap_text=True), HEAD_FILL)
            cell = put(ws, r, 3, 3, None, f(11), Alignment(vertical='center', wrap_text=True, indent=1), INPUT_FILL)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if typ == 'date':
                cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
            elif typ == 'money':
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right', vertical='center', indent=1)
            elif typ == 'num':
                cell.number_format = 'General'
                cell.alignment = Alignment(horizontal='right', vertical='center', indent=1)
            from openpyxl.styles import Protection
            cell.protection = Protection(locked=False)
            h = put(ws, r, 4, 4, hlp, f(9.5, color='666666'), Alignment(vertical='center', wrap_text=True, indent=1), border=False)
            ws.row_dimensions[r].height = 22 if typ != 'text' or key not in ('공사명', '현장위치', '공사개요', '연장사유', '업체주소') else 30
            REF[key] = f"'입력'!$C${r}"
            CELL[key] = f'C{r}'
            if opts:
                sig = ','.join(opts)
                if sig not in dvs:
                    dv = DataValidation(type='list', formula1='"' + sig + '"', allow_blank=True, showErrorMessage=False)
                    ws.add_data_validation(dv); dvs[sig] = dv
                dvs[sig].add(f'C{r}')
            elif typ in ('date',):
                dv = dvs.get('__date')
                if not dv:
                    dv = DataValidation(type='date', operator='greaterThan', formula1='1', allow_blank=True,
                                        showErrorMessage=True, errorTitle='날짜', error='날짜로 넣어 주세요. 예: 2026-03-02')
                    ws.add_data_validation(dv); dvs['__date'] = dv
                dv.add(f'C{r}')
            elif typ in ('money', 'num'):
                dv = dvs.get('__num')
                if not dv:
                    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                                        showErrorMessage=True, errorTitle='숫자', error='숫자만 넣어 주세요 (쉼표·원 없이).')
                    ws.add_data_validation(dv); dvs['__num'] = dv
                dv.add(f'C{r}')
            meta_inputs.append(dict(key=key, label=label, type=typ, cell=f'C{r}', ex=ex, help=hlp, opts=opts,
                                    sec=sec, req=key in REQUIRED))
            r += 1
        r += 1
    # 자동 계산
    put(ws, r, 2, 4, '자동 계산 — 고치지 마세요 (위 칸을 고치면 따라 바뀝니다)', f(12, True, 'FFFFFF'),
        Alignment(vertical='center', indent=1), PatternFill('solid', fgColor='5A6B80'), border=False)
    ws.row_dimensions[r].height = 24
    r += 1
    for key, label, _, _ in CALCS:
        REF[key] = f"'입력'!$C${r}"
        CELL[key] = f'C{r}'
        r += 1
    r -= len(CALCS)
    for key, label, tpl, fmt in CALCS:
        put(ws, r, 2, 2, label, f(10.5), Alignment(vertical='center', indent=1), CALC_FILL)
        c = put(ws, r, 3, 3, F(fx(tpl)), f(11), Alignment(horizontal='right', vertical='center', indent=1), CALC_FILL, fmt=fmt)
        ws.row_dimensions[r].height = 20
        r += 1
    ws.freeze_panes = 'A5'
    ws.protection.sheet = True          # 암호 없음 — 검토 → 시트 보호 해제로 풀 수 있음
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    page(ws)
    ws.page_setup.fitToHeight = 0
    ws.print_area = f'A1:D{r}'
    return ws, meta_inputs


def make_calc(wb):
    """금액 → 한글 (숨김 시트)"""
    ws = wb.create_sheet('계산')
    ws['A1'] = '금액을 한글로 바꾸는 칸 (서류가 씁니다 — 지우지 마세요)'
    W = lambda d: f'CHOOSE({d}+1,"","일","이","삼","사","오","육","칠","팔","구")'
    hangul_ref = {}
    r = 3
    extra = {'증감액abs': f'ABS(N({REF["증감액"]}))'}
    for k in HANGUL:
        src = extra.get(k) or REF[k]
        ws.cell(r, 1, k)
        ws.cell(r, 2, f'=IF({src}="","",ROUND(ABS({src}),0))')
        ws.cell(r, 3, f'=IF(B{r}="",0,INT(B{r}/100000000))')
        ws.cell(r, 4, f'=IF(B{r}="",0,MOD(INT(B{r}/10000),10000))')
        ws.cell(r, 5, f'=IF(B{r}="",0,MOD(B{r},10000))')
        for gi, gc in enumerate('CDE'):
            g = f'{gc}{r}'
            d1, d2, d3, d4 = f'INT({g}/1000)', f'MOD(INT({g}/100),10)', f'MOD(INT({g}/10),10)', f'MOD({g},10)'
            fxs = (f'={W(d1)}&IF({d1}>0,"천","")&{W(d2)}&IF({d2}>0,"백","")&{W(d3)}&IF({d3}>0,"십","")&{W(d4)}')
            ws.cell(r, 6 + gi, fxs)
        ws.cell(r, 9, f'=IF(B{r}="","",IF(B{r}=0,"영",IF(C{r}>0,F{r}&"억","")&IF(D{r}>0,G{r}&"만","")&H{r}))')
        ws.cell(r, 10, f'=IF(B{r}="","","일금 "&I{r})')
        hangul_ref[k] = f"'계산'!$J${r}"
        r += 1
    ws.sheet_state = 'hidden'
    return ws, hangul_ref


def make_front(wb, docs_meta, meta_inputs):
    ws = wb.active
    ws.title = '처음'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = '1F3A5F'
    for c, w in zip('ABCDEFG', (1.5, 6, 34, 12, 10, 10, 44)):
        ws.column_dimensions[c].width = w
    put(ws, 1, 2, 7, 'K-건설맵 공사서류 원클릭', f(22, True, '1F3A5F'), Alignment(vertical='center'), border=False)
    ws.row_dimensions[1].height = 46
    put(ws, 2, 2, 7, '한 번 입력하면 착공부터 준공·하자까지 서류 24가지가 채워집니다. 매크로 없이 수식만 써서 '
                     '엑셀·한셀·구글 시트 어디서나 열립니다. 모든 공사 현장(관급·민간)에 쓸 수 있습니다.',
        f(10.5, color='333333'), Alignment(vertical='center', wrap_text=True), border=False)
    ws.row_dimensions[2].height = 34
    steps = [('①', '「입력」 시트의 노란 칸을 채웁니다.', "'입력'!A1", '입력 하러 가기 →'),
             ('②', '아래 목록에서 서류 이름을 누르면 그 서류로 갑니다. 서류 칸은 입력에서 저절로 채워집니다.', None, None),
             ('③', '인쇄하거나 PDF로 저장합니다 (한 서류가 A4 한 장).', None, None)]
    r = 4
    for n, s, loc, lab in steps:
        put(ws, r, 2, 2, n, f(13, True, '1F5FBF'), Alignment(horizontal='center', vertical='center'), border=False)
        put(ws, r, 3, 6, s, f(11), Alignment(vertical='center', wrap_text=True), border=False)
        if loc:
            c = ws.cell(r, 7, lab)
            c.hyperlink = Hyperlink(ref=f'G{r}', location=loc, display=lab)
            c.font = Font(name=FONT, size=11, bold=True, color='1F5FBF', underline='single')
            c.alignment = Alignment(vertical='center')
        ws.row_dimensions[r].height = 24
        r += 1
    # 빠진 칸 알림
    req = [f'({REF[k]}="")' for k in REQUIRED]
    r += 1
    put(ws, r, 2, 7, F('IF(' + '+'.join(req) + '=0,"✔ 꼭 필요한 칸은 다 채웠습니다.","⚠ 꼭 필요한 칸 중 "&(' + '+'.join(req) + ')&"개가 비어 있습니다 — 입력 시트를 확인하세요.")'),
        f(11, True, 'B25A00'), Alignment(vertical='center', indent=1), PatternFill('solid', fgColor='FFF4E5'), border=False)
    ws.row_dimensions[r].height = 26
    r += 2
    heads = ['번호', '서  류', '누가 내나', '언  제', '관급', '무엇인가']
    for j, h in enumerate(heads):
        put(ws, r, 2 + j, 2 + j, h, f(10.5, True, 'FFFFFF'), Alignment(horizontal='center', vertical='center'), SEC_FILL)
    ws.row_dimensions[r].height = 24
    r += 1
    for d in docs_meta:
        put(ws, r, 2, 2, d['no'], f(10.5), Alignment(horizontal='center', vertical='center'))
        c = put(ws, r, 3, 3, d['title'].replace(' ', '') if len(d['title']) > 0 else '', Font(name=FONT, size=11, bold=True, color='1F5FBF', underline='single'),
                Alignment(vertical='center', indent=1))
        c.value = d['name']
        c.hyperlink = Hyperlink(ref=f'C{r}', location=f"'{d['sheet']}'!A1", display=d['name'])
        put(ws, r, 4, 4, d['who'], f(10.5, color='2F6FB5' if d['who'] == '업체' else 'B7802F'), Alignment(horizontal='center', vertical='center'))
        put(ws, r, 5, 5, d['when'], f(10.5), Alignment(horizontal='center', vertical='center'))
        put(ws, r, 6, 6, '관급' if d['pub'] else '', f(10, color='777777'), Alignment(horizontal='center', vertical='center'))
        put(ws, r, 7, 7, d['desc'], f(9.5, color='444444'), Alignment(vertical='center', wrap_text=True, indent=1))
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    notes = [
        '• 보증금률·지체상금률·하자담보책임기간은 계약서(공고서)에 적힌 값을 넣으세요. 공종·계약마다 다릅니다.',
        '• 발주기관이 정한 서식이 있으면 그 서식을 씁니다. 이 파일은 정해진 서식이 없을 때 쓰는 기본 양식입니다.',
        '• 「관급」 표시는 관공서 발주 공사에서 주로 내는 서류입니다. 민간 공사는 필요한 것만 쓰세요.',
        '• 입력 시트는 수식이 지워지지 않게 보호해 두었습니다(암호 없음). 칸 모양을 바꾸려면 검토 → 시트 보호 해제.',
        '• 인쇄하면 머리글 오른쪽에 작은 K-건설맵 표시가 나옵니다. 지우려면 페이지 레이아웃 → 페이지 설정 → 머리글/바닥글에서 «(없음)».',
        '• 사이트(k-conmap.com/tools/wonclick)에서 칸을 채우고 받으면, 입력이 채워진 파일을 바로 받을 수 있습니다.',
    ]
    for s in notes:
        put(ws, r, 2, 7, s, f(10, color='444444'), Alignment(vertical='center', wrap_text=True), border=False)
        ws.row_dimensions[r].height = 30 if len(s) > 60 else 20
        r += 1
    page(ws)
    ws.page_setup.fitToHeight = 0
    ws.print_area = f'A1:G{r}'


def set_default_font(wb):
    """기본 글꼴 맑은 고딕 11 (열 너비 기준 — 엑셀과 인쇄 폭이 맞게)"""
    from openpyxl.styles import NamedStyle
    ns = wb._named_styles['Normal']
    ns.font = Font(name=FONT, size=11)
    wb._fonts[0] = Font(name=FONT, size=11)


def build(outdir, sample=None):
    REF.clear(); CELL.clear()
    wb = Workbook()
    set_default_font(wb)
    front = wb.active
    front.title = '처음'
    ws_in, meta_inputs = make_input(wb)
    calc_ws, hangul_ref = make_calc(wb)
    build_docs()
    number_docs()
    docs_meta = []
    for d in DOCS:
        make_sheet_for_doc(wb, d, hangul_ref)
        docs_meta.append(dict(no=d['no'], sheet=d['sheet'], name=d['name'], title=d['title'], who=d['who'],
                              when=d['when'], pub=d['pub'], desc=d['desc']))
    make_front(wb, docs_meta, meta_inputs)
    # 순서: 처음, 입력, 서류들…, 계산(숨김)
    order = ['처음', '입력'] + [d['sheet'] for d in DOCS] + ['계산']
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == '처음'
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = 'K-건설맵'
    wb.properties.title = 'K-건설맵 공사서류 원클릭'
    if sample:
        from datetime import date
        for k, v in sample.items():
            c = ws_in[CELL[k]]
            if isinstance(v, str) and len(v) == 10 and v[4] == '-' and v[7] == '-':
                y, m, dd = map(int, v.split('-')); c.value = date(y, m, dd)
            else:
                c.value = v
    os.makedirs(outdir, exist_ok=True)
    name = 'wonclick_sample.xlsx' if sample else 'wonclick.xlsx'
    wb.save(os.path.join(outdir, name))
    sheet_files = {ws.title: f'xl/worksheets/sheet{i + 1}.xml' for i, ws in enumerate(wb.worksheets)}
    meta = dict(version=1, file='/tools/files/k-conmap-wonclick.xlsx', inputSheet=sheet_files['입력'],
                sections=[s for s, _ in SECTIONS], inputs=meta_inputs, docs=docs_meta, required=REQUIRED)
    if not sample:
        json.dump(meta, open(os.path.join(outdir, 'wonclick_meta.json'), 'w'), ensure_ascii=False, indent=1)
    return meta


if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else '/tmp/wc/out'
    sample = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else None
    m = build(out)
    if sample:
        build(out, sample)
    print('ok', len(m['inputs']), 'inputs', len(m['docs']), 'docs')
